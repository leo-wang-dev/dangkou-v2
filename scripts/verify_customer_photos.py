"""Real photos through TgApi HTTP + CsBot + SQLite + export; model can be offline replay or real.

Default mode proves transport/state/export only. --real-model invokes configured vision/text models.
The local Telegram-compatible server never sends a message to a real Telegram account.
"""
import argparse
import hashlib
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import io
import json
from pathlib import Path
import re
import sys
import threading
import time
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import openpyxl
from fastapi import FastAPI
from fastapi.testclient import TestClient
from catalog import db, config, llm
from catalog.api import register_routes
from catalog.csbot import CsBot
from catalog.storage import LocalStorage
from catalog.tg import TgApi


class ReplayModel:
    """Visually reviewed reference responses. Not OCR and not a quality measurement."""
    def __init__(self, cases, images):
        self.answers={hashlib.sha256(images[c['id']]).hexdigest():c['items'] for c in cases}

    def chat_vision(self, prompt, data):
        return json.dumps(self.answers[hashlib.sha256(data).hexdigest()],ensure_ascii=False)

    def chat_text(self, system, messages, **kwargs):
        if '草稿编辑判定器' in system:
            text=messages[-1]['content'].split('客户消息：')[-1]
            if '1 颜色改成黑色' in text:
                return json.dumps({'action':'edit','index':1,'field':'颜色','value':'黑色'})
            if '起订量一箱起' in text:
                return json.dumps({'action':'add','index':1,'field':'起订量','value':'一箱起'})
            return '{"action":"none"}'
        return '<<PASS>>'


class RecordedModel:
    def __init__(self, out):
        self.out=out
        self.calls=0

    def _record(self, kind, invoke, prompt):
        self.calls+=1
        start=time.monotonic()
        response=invoke()
        (self.out/f'model-{self.calls:02d}-{kind}.json').write_text(json.dumps({
            'kind':kind,'model':llm.VISION_MODEL if kind=='vision' else llm.TEXT_MODEL,
            'prompt':prompt,'elapsed_sec':round(time.monotonic()-start,2),'response':response},ensure_ascii=False,indent=2))
        return response

    def chat_vision(self,prompt,data):
        return self._record('vision',lambda:llm.chat_vision(prompt,data),prompt)

    def chat_text(self,system,messages,**kwargs):
        return self._record('text',lambda:llm.chat_text(system,messages,**kwargs),{'system':system,'messages':messages})


def run(out, real_model=False):
    out=Path(out);out.mkdir(parents=True,exist_ok=True)
    if (out/'catalog.db').exists():
        raise RuntimeError('输出目录已有测试数据库，请使用新的 --out 目录，避免混入上次数据')
    dataset=json.loads((Path(__file__).resolve().parents[1]/'tests/fixtures/customer_photos.json').read_text())
    source=Path(__import__('os').environ.get('CS_SAMPLE_PHOTO_DIR',dataset['source_dir']))
    if real_model and not config.BAILIAN_API_KEY:
        raise RuntimeError('BLOCKED: 缺 BAILIAN_API_KEY；照片已齐，不能用参考数据回放冒充真实识图')
    images={c['id']:(source/c['file']).read_bytes() for c in dataset['cases']}
    received=[]; updates=[]; requests=[]
    class Handler(BaseHTTPRequestHandler):
        def log_message(self,*args):pass
        def do_POST(self):
            raw=self.rfile.read(int(self.headers.get('Content-Length',0)))
            method=self.path.rsplit('/',1)[-1]
            if method == 'sendDocument':
                from email.parser import BytesParser
                from email.policy import default
                multipart=BytesParser(policy=default).parsebytes(
                    ('Content-Type: '+self.headers['Content-Type']+'\r\nMIME-Version: 1.0\r\n\r\n').encode()+raw)
                parts={p.get_param('name',header='content-disposition'):p for p in multipart.iter_parts()}
                data=parts['document'].get_payload(decode=True)
                (out/'telegram-attachment.xlsx').write_bytes(data)
                sheet=openpyxl.load_workbook(io.BytesIO(data)).active
                assert sheet.max_row > 1 and sheet._images
                requests.append({'method':method,'bytes':len(data)})
                response=json.dumps({'ok':True,'result':{'message_id':900}}).encode()
                self.send_response(200);self.send_header('Content-Length',str(len(response)));self.end_headers();self.wfile.write(response)
                return
            body=json.loads(raw or '{}')
            requests.append({'method':method,'params':body})
            if method=='getUpdates': result=[u for u in updates if u['update_id']>=body.get('offset',0)]
            elif method=='getFile':result={'file_path':body['file_id']+'.jpg'}
            elif method=='sendMessage':
                received.append(body);result={'message_id':len(received)}
            else:raise AssertionError(method)
            data=json.dumps({'ok':True,'result':result}).encode()
            self.send_response(200);self.send_header('Content-Length',str(len(data)));self.end_headers();self.wfile.write(data)
        def do_GET(self):
            data=images[Path(self.path).stem]
            self.send_response(200);self.send_header('Content-Length',str(len(data)));self.end_headers();self.wfile.write(data)
    server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    conn=db.connect(str(out/'catalog.db'));db.init_db(conn)
    conn.execute("UPDATE shop_profile SET owner_tg_username='test_owner', owner_wechat='TEST-ONLY-WECHAT'");conn.commit()
    api=TgApi(token='LOCAL-TEST-NOT-A-REAL-TOKEN',base=f'http://127.0.0.1:{server.server_port}')
    model=RecordedModel(out) if real_model else ReplayModel(dataset['cases'],images)
    notices=[]
    photo_dir=(out/'photos').resolve()
    bot=CsBot(conn,api,llm=model,notifier=lambda text:notices.append(text),img_dir=str(photo_dir))
    app=FastAPI();app.state.conn=conn;app.state.token='local-test-only';app.state.storage=LocalStorage(str(out/'images'));app.state.callback=None
    register_routes(app)
    results=[]
    transient_failures=[]
    try:
        for i,case in enumerate(dataset['cases'],1):
            update={'update_id':i,'message':{'chat':{'id':9001,'type':'private'},'from':{'id':9001,'username':'fixture_buyer'},'photo':[{'file_id':case['id'],'width':1280,'height':1707}]}}
            updates.append(update)
            polled=api.poll(timeout=0)
            assert polled==[update]
            start=time.monotonic()
            for attempt in range(3):
                try:
                    bot.handle_update(polled[0])
                    break
                except ValueError as exc:
                    if '空内容' not in str(exc) or attempt==2:
                        raise
                    transient_failures.append({'photo':case['id'],'attempt':attempt+1,'error':'provider_empty_content'})
                    (out/'transient-failures.json').write_text(json.dumps(transient_failures,ensure_ascii=False,indent=2))
                    time.sleep(2*(attempt+1))
            api._offset=i+1
            elapsed=round(time.monotonic()-start,2)
            cid=conn.execute("SELECT id FROM cs_customer WHERE tg_id='9001'").fetchone()[0]
            rows=conn.execute('SELECT fields_json,photo FROM cs_note WHERE customer_id=? ORDER BY id DESC',(cid,)).fetchall()
            recent=[json.loads(r[0]) for r in rows if r['photo']==rows[0]['photo']]
            joined=json.dumps(recent,ensure_ascii=False).lower()
            checks={term:term.lower() in joined for term in case['required_terms']}
            count_ok=len(recent)==len(case['items'])
            if case['id']=='hair_oil':
                checks['unclear_price_not_invented']=all('模糊' in str(f.get('价格','')) for f in recent)
            if case['id']=='two_tubes':
                cropped=[f for f in recent if '260' in str(f.get('体积或尺寸',''))]
                checks['cropped_price_not_invented']=bool(cropped) and all(any(word in str(f.get('价格','')) for word in ('模糊','未拍到')) for f in cropped)
            (out/(case['id']+'-extraction.json')).write_text(json.dumps({'fields':recent,'checks':checks,'count_ok':count_ok},ensure_ascii=False,indent=2))
            assert all(checks.values()) and count_ok,(case['id'],recent)
            # Duplicate Telegram update must not duplicate notes or generate another reply.
            count=len(rows);sent=len(received);bot.handle_update(update)
            assert conn.execute('SELECT COUNT(*) FROM cs_note').fetchone()[0]==count
            assert len(received)==sent
            results.append({'photo':case['id'],'elapsed_sec':elapsed,'sha256':hashlib.sha256(images[case['id']]).hexdigest(),'term_checks':checks,'fields':recent})
        for uid,text in enumerate(['1 颜色改成黑色','第1条起订量一箱起','你们在哪条街','确认','出表','这些样品有货吗？找老板询价'],10):
            bot.handle_update({'update_id':uid,'message':{'chat':{'id':9001,'type':'private'},'from':{'id':9001},'text':text}})
        cid=conn.execute("SELECT id FROM cs_customer WHERE tg_id='9001'").fetchone()[0]
        fields=json.loads(conn.execute('SELECT fields_json FROM cs_note WHERE customer_id=? ORDER BY id LIMIT 1',(cid,)).fetchone()[0])
        assert fields['颜色']=='黑色' and fields['起订量']=='一箱起'
        token=conn.execute('SELECT token FROM cs_link WHERE customer_id=?',(cid,)).fetchone()[0]
        with patch.dict('os.environ',{'CATALOG_CS_PHOTOS':str(photo_dir)}), TestClient(app) as client:
            data=client.get(f'/cs/link/{token}').json()
            response=client.get(f'/cs/link/{token}/export.xlsx');assert response.status_code==200
            (out/'customer_purchase_list.xlsx').write_bytes(response.content)
            wb=openpyxl.load_workbook(io.BytesIO(response.content));ws=wb.active
            note_count=conn.execute("SELECT COUNT(*) FROM cs_note WHERE customer_id=? AND status='confirmed'",(cid,)).fetchone()[0]
            assert ws.max_row==note_count+1 and len(ws._images)==note_count
            assert client.get('/cs/link/not-a-real-link/export.xlsx').status_code in (403,404,410)
        assert notices and 'TEST-ONLY-WECHAT' in received[-1]['text']
        evidence={'mode':'real-model-local-TG' if real_model else 'reference-fixture-replay-local-TG',
                  'not_model_quality_proof':not real_model,'transient_failures':transient_failures,'photo_count':4,'confirmed_items':note_count,
                  'embedded_excel_images':len(ws._images),'http_methods':sorted({r['method'] for r in requests}),
                  'photos':results,'messages':received,'merchant_notifications':notices}
        (out/'result.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2))
        print(json.dumps({k:v for k,v in evidence.items() if k not in ('photos','messages','merchant_notifications')},ensure_ascii=False))
        return evidence
    finally:
        conn.close();server.shutdown();server.server_close();thread.join(timeout=2)


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--out',default='audit/round4/photo-flow');p.add_argument('--real-model',action='store_true');args=p.parse_args()
    try:run(args.out,args.real_model)
    except RuntimeError as exc:
        if str(exc).startswith('BLOCKED:'):
            print(exc);sys.exit(2)
        raise
