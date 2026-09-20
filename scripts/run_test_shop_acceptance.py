"""Run the real test-shop flow without contacting real customers or legacy WeChat."""
import argparse
import base64
import io
import json
import os
from pathlib import Path
import sys
import time

import openpyxl
import requests

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from catalog import db, merchant_onboarding as hub
from catalog.csbot import CsBot
from catalog.merchant_binding import credentials, root


class CaptureTelegram:
    def __init__(self, photos, out):
        self.photos = photos
        self.out = out
        self.messages = []
        self.documents = []

    def download_photo(self, sizes):
        return self.photos[max(sizes, key=lambda x: x.get("width", 0))["file_id"]]

    def send_message(self, chat_id, text):
        self.messages.append({"chat_id": chat_id, "text": text})
        return {"message_id": len(self.messages)}

    def send_document(self, chat_id, filename, content, caption=""):
        target = self.out / filename
        target.write_bytes(content)
        self.documents.append({"chat_id": chat_id, "filename": filename,
                               "bytes": len(content), "caption": caption})
        return {"message_id": 1000 + len(self.documents)}


def _request(session, method, url, token, **kwargs):
    response = session.request(method, url, headers={"X-Service-Token": token},
                               timeout=180, **kwargs)
    response.raise_for_status()
    return response


def _build_import(path, prefix):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ITEM.NO型号", "装箱尺寸", "装箱数量", "价格", "电压", "功率",
               "发热体", "材质", "频率", "备注"])
    ws.append([prefix + "-A", "50×40×35cm", "24", "19.8", "110-240V", "45W",
               "PTC", "ABS+陶瓷", "50/60Hz", "TEST ONLY"])
    ws.append([prefix + "-B", "51×41×36cm", "24", "20.8", "110-240V", "48W",
               "PTC", "ABS+陶瓷", "50/60Hz", "TEST ONLY"])
    wb.save(path)


def _imports(session, base, token, out):
    listed = _request(session, "GET", base + "/tickets", token).json()["tickets"]
    before = {x["id"] for x in listed}
    completed = {}
    for ticket in listed:
        if ticket["status"] not in ("approved", "rejected"):
            continue
        detail = _request(session, "GET", f"{base}/tickets/{ticket['id']}", token).json()
        source_key = detail["payload"].get("source_key")
        if source_key in ("acceptance-A", "acceptance-B"):
            completed[source_key] = {
                "ticket": ticket["id"],
                "doc_id": detail["payload"]["doc_id"],
                "approved": ticket["status"] == "approved",
            }
    if set(completed) == {"acceptance-A", "acceptance-B"}:
        decisions = [completed[key] for key in sorted(completed)]
        if sorted(value["approved"] for value in decisions) != [False, True]:
            raise AssertionError("existing concurrent import decisions are not one approval and one rejection")
        statuses = {
            str(value["doc_id"]): _request(
                session, "GET", f"{base}/import/{value['doc_id']}", token
            ).json()
            for value in decisions
        }
        if any(value["status"] != "ticketed" for value in statuses.values()):
            raise AssertionError({"existing_imports": statuses})
        return {"docs": statuses, "decisions": decisions, "reused": True}
    docs = []
    for suffix in ("A", "B"):
        path = out / f"concurrent-import-{suffix}.xlsx"
        _build_import(path, "TEST-IMPORT-" + suffix)
        result = _request(session, "POST", base + "/import", token,
                          json={"path": str(path), "category": "curler",
                                "source_key": "acceptance-" + suffix}).json()
        docs.append(result["doc_id"])
    deadline = time.monotonic() + 1200
    statuses = {}
    while time.monotonic() < deadline:
        statuses = {doc: _request(session, "GET", f"{base}/import/{doc}", token).json()
                    for doc in docs}
        if all(value["status"] in ("ticketed", "failed") for value in statuses.values()):
            break
        time.sleep(3)
    if any(value["status"] != "ticketed" for value in statuses.values()):
        raise AssertionError({"imports": statuses})
    tickets = [x for x in _request(session, "GET", base + "/tickets", token).json()["tickets"]
               if x["id"] not in before and x["status"] == "pending"]
    selected = []
    for ticket in tickets:
        detail = _request(session, "GET", f"{base}/tickets/{ticket['id']}", token).json()
        if detail["payload"].get("doc_id") in docs:
            selected.append(ticket)
    if len(selected) != 2:
        raise AssertionError("two concurrent imports did not produce two independent tickets")
    decisions = []
    for index, ticket in enumerate(sorted(selected, key=lambda x: x["id"])):
        result = session.post(f"{base}/tickets/{ticket['id']}/decision",
                              json={"token": ticket["token"], "approved": index == 0},
                              timeout=180)
        result.raise_for_status()
        decisions.append({"ticket": ticket["id"], "approved": index == 0,
                          "result": result.json()})
    return {"docs": statuses, "decisions": decisions}


def run(merchant_id, photos_dir, out, with_imports=False):
    out = Path(out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    c = hub.connect()
    try:
        merchant = c.execute("SELECT * FROM merchant WHERE id=?", (merchant_id,)).fetchone()
        if not merchant or merchant["state"] != "catalog_ready" or merchant["runtime_status"] != "running":
            raise RuntimeError("test shop catalog service is not running")
        private = json.loads(credentials(merchant_id).read_text())
        database, port = merchant["db_path"], merchant["port"]
    finally:
        c.close()
    base = f"http://127.0.0.1:{port}"
    session = requests.Session()
    session.trust_env = False
    token = private["api_token"]

    health = _request(session, "GET", base + "/health", token).json()
    initial_catalog = _request(session, "GET", base + "/cs/catalog", token).json()
    names = {item["name"] for item in initial_catalog["products"]}
    if not {"TEST-RZ-8226", "TEST-CURL-2026"}.issubset(names):
        raise AssertionError(names)
    probe = db.connect(database)
    try:
        exact_rel = probe.execute(
            "SELECT image_main FROM product_curler WHERE id='test-curl-2026'"
        ).fetchone()[0]
    finally:
        probe.close()
    exact_image = Path(database).parent / "images" / exact_rel
    exact_match = _request(
        session, "POST", base + "/cs/catalog/search", token,
        json={"fields": [{"型号或品名": "未知商品"}],
              "image_base64": base64.b64encode(exact_image.read_bytes()).decode()},
    ).json()["candidates"]
    if not exact_match or exact_match[0]["product_id"] != "test-curl-2026":
        raise AssertionError({"exact_image_match": exact_match})

    created = _request(session, "POST", base + "/products/razor/direct", token,
                       json={"changes": {"model_no": "TEST-CRUD-TEMP",
                                          "description": "贯通测试临时商品",
                                          "color": "蓝色", "cs_visible": "1"}}).json()
    temporary_id = created["id"]
    _request(session, "PATCH", f"{base}/products/razor/{temporary_id}/direct", token,
             json={"changes": {"color": "绿色"}})
    after_create = _request(session, "GET", base + "/cs/catalog", token).json()
    assert "TEST-CRUD-TEMP" in {item["name"] for item in after_create["products"]}
    _request(session, "DELETE", f"{base}/products/razor/{temporary_id}/direct", token)
    after_delete = _request(session, "GET", base + "/cs/catalog", token).json()
    assert "TEST-CRUD-TEMP" not in {item["name"] for item in after_delete["products"]}

    import_result = _imports(session, base, token, out) if with_imports else {"skipped": True}

    fixture = json.loads((Path(__file__).resolve().parents[1] / "tests/fixtures/customer_photos.json").read_text())
    source = Path(photos_dir)
    photos = {case["id"]: (source / case["file"]).read_bytes() for case in fixture["cases"]}
    api = CaptureTelegram(photos, out)
    conn = db.connect(database)
    customer_tg = "99000001"
    try:
        customer = conn.execute("SELECT id FROM cs_customer WHERE tg_id=?", (customer_tg,)).fetchone()
        if customer:
            cid = customer["id"]
            conn.execute("DELETE FROM cs_note WHERE customer_id=?", (cid,))
            conn.execute("DELETE FROM cs_conversation_log WHERE customer_id=?", (cid,))
            conn.execute("DELETE FROM cs_context WHERE customer_id=?", (cid,))
            conn.execute("DELETE FROM cs_photo_candidates WHERE customer_id=?", (cid,))
            conn.execute("DELETE FROM cs_link WHERE customer_id=?", (cid,))
        conn.execute("DELETE FROM cs_inbox WHERE update_id BETWEEN 990000000 AND 990000099")
        conn.execute("DELETE FROM cs_outbox WHERE recipient=?", (customer_tg,))
        conn.commit()
        os.environ.update(CATALOG_CS_API_URL=base, CATALOG_CS_SERVICE_TOKEN=token)
        bot = CsBot(conn, api, img_dir=str(root() / merchant_id / "acceptance-photos"))
        started = time.monotonic()
        unrelated_photo_checks = []
        for index, case in enumerate(fixture["cases"]):
            update = {"update_id": 990000000 + index,
                      "message": {"chat": {"id": int(customer_tg), "type": "private"},
                                  "from": {"id": int(customer_tg), "username": "test_buyer"},
                                  "photo": [{"file_id": case["id"], "width": 1280, "height": 1707}]}}
            for attempt in range(3):
                try:
                    bot.handle_update(update)
                    break
                except ValueError as exc:
                    if "空内容" not in str(exc) or attempt == 2:
                        raise
                    time.sleep(2 * (attempt + 1))
            customer = conn.execute(
                "SELECT id FROM cs_customer WHERE tg_id=?", (customer_tg,)
            ).fetchone()
            candidate_row = conn.execute(
                "SELECT candidates FROM cs_photo_candidates WHERE customer_id=?",
                (customer["id"],),
            ).fetchone()
            candidates = json.loads(candidate_row[0]) if candidate_row else []
            if candidates:
                raise AssertionError({"unrelated_photo": case["id"], "candidates": candidates})
            unrelated_photo_checks.append(case["id"])
        base_update = 990000010
        for offset, text in enumerate(("1 颜色改成黑色", "第1条起订量一箱起", "确认", "出表")):
            bot.handle_update({"update_id": base_update + offset,
                               "message": {"chat": {"id": int(customer_tg), "type": "private"},
                                           "from": {"id": int(customer_tg)}, "text": text}})
        cid = conn.execute("SELECT id FROM cs_customer WHERE tg_id=?", (customer_tg,)).fetchone()[0]
        notes = conn.execute("SELECT * FROM cs_note WHERE customer_id=? AND status='confirmed' ORDER BY id", (cid,)).fetchall()
        if len(notes) != 5:
            raise AssertionError(f"expected 5 confirmed photo items, got {len(notes)}")
        first = json.loads(notes[0]["fields_json"])
        assert first["颜色"] == "黑色" and first["起订量"] == "一箱起"
        assert api.documents
        workbook = openpyxl.load_workbook(out / api.documents[-1]["filename"])
        assert workbook.active.max_row == 6 and len(workbook.active._images) == 5

        cust = dict(conn.execute("SELECT * FROM cs_customer WHERE id=?", (cid,)).fetchone())
        catalog_reply = bot._on_text(cust, "看看商品")
        price_reply = bot._on_text(cust, "TEST-RZ-8226 100个多少钱")
        transfer_reply = bot._on_text(cust, "可以月结60天吗？")
        pass_reply = bot._on_text(cust, "可以月结15天吗？")
        owner_reply = bot._on_text(cust, "找老板")
        assert "TEST-RZ-8226" in catalog_reply and "TEST-CURL-2026" in catalog_reply
        assert "仅展示商家原文" in price_reply and "TEST_ONLY_WECHAT" in transfer_reply
        assert "TEST_ONLY_WECHAT" not in pass_reply and "@hehhh553" in owner_reply
        elapsed = round(time.monotonic() - started, 2)
    finally:
        conn.close()

    evidence = {
        "merchant_id": merchant_id, "shop_id": initial_catalog["shop_id"],
        "health": health, "catalog_names": sorted(names),
        "crud": {"create_update_visible": True, "delist_hidden": True},
        "imports": import_result, "photo_count": 4, "confirmed_items": 5,
        "excel_rows": workbook.active.max_row, "excel_images": len(workbook.active._images),
        "customer_flows": {"catalog": True, "merchant_quote_text": True,
                           "configured_handoff": True, "condition_not_overmatched": True,
                           "explicit_owner_contact": True,
                           "exact_catalog_image_hit": True,
                           "unrelated_photos_not_hit": len(unrelated_photo_checks) == 4},
        "telegram_transport": "captured local transport; no real second merchant bot was available",
        "elapsed_sec": elapsed,
    }
    (out / "result.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2))
    print(json.dumps(evidence, ensure_ascii=False))
    return evidence


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--merchant-id", required=True)
    parser.add_argument("--photos-dir", required=True)
    parser.add_argument("--out", required=True)
    parser.add_argument("--with-imports", action="store_true")
    args = parser.parse_args()
    run(args.merchant_id, args.photos_dir, args.out, args.with_imports)
