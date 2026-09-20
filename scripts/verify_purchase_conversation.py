"""Real model acceptance; isolated database and local Telegram transport, no customer messages."""
from pathlib import Path
import sys, json, tempfile, io
from types import SimpleNamespace
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from catalog import db, llm, config, merchant_policy
from catalog.csbot import CsBot
import openpyxl


def run(out):
    out=Path(out);out.mkdir(parents=True,exist_ok=True)
    if not config.BAILIAN_API_KEY:raise RuntimeError('BAILIAN_API_KEY missing')
    results=[]
    with tempfile.TemporaryDirectory() as temp:
        c=db.connect(str(Path(temp)/'shop.db'));db.init_db(c)
        c.execute("UPDATE shop_profile SET shop_name='验收档口',tg_bot_id='999',owner_wechat='TEST-BOSS' WHERE id=1")
        merchant_policy.apply(c,{'shop_name':'验收档口','logistics':'客户要求加急赶货时转人工'},1);c.commit()
        api=SimpleNamespace(sent=[],documents=[])
        api.send_message=lambda *a:api.sent.append(a)
        api.send_document=lambda *a:api.documents.append(a)
        bot=CsBot(c,api,llm=llm,img_dir=str(Path(temp)/'photos'))
        cases=[('new','帮我记录：杯子蓝色100个，盘子白色20件。',2),
               ('edit_handoff','第一条杯子数量改成200个，能加急赶货吗？',2),
               ('question','杯子有绿色的吗？',2),
               ('export','出表',2)]
        for uid,(name,text,count) in enumerate(cases,1):
            bot.handle_update({'update_id':uid,'message':{'from':{'id':100},'chat':{'id':100,'type':'private'},'text':text}})
            notes=[json.loads(r[0]) for r in c.execute('SELECT fields_json FROM cs_note ORDER BY id')]
            reply=api.sent[-1][1]
            passed=len(notes)==count
            if name=='new':passed=passed and notes[0].get('数量')=='100个' and notes[1].get('数量')=='20件'
            if name=='edit_handoff':passed=passed and notes[0].get('数量')=='200个' and 'TEST-BOSS' in reply
            if name=='export':
                passed=passed and bool(api.documents)
                if api.documents:
                    content=api.documents[-1][2];(out/'采购清单.xlsx').write_bytes(content)
                    wb=openpyxl.load_workbook(io.BytesIO(content));passed=passed and wb.active.max_row==3
            results.append({'case':name,'input':text,'notes':notes,'reply':reply,'passed':passed})
            print(name,passed,flush=True)
        fixture=json.loads((Path(__file__).resolve().parents[1]/'tests/fixtures/customer_photos.json').read_text())
        photo=Path(fixture['source_dir'])/fixture['cases'][0]['file']
        if not photo.is_file():raise RuntimeError('客户原图缺失')
        api.download_photo=lambda _:photo.read_bytes()
        bot.handle_update({'update_id':100,'message':{'from':{'id':200},'chat':{'id':200,'type':'private'},
            'photo':[{'file_id':'original'}], 'caption':'这款要100瓶，能加急赶货吗？'}})
        cust=bot._ensure_customer({'id':200})
        notes=[json.loads(r[0]) for r in c.execute('SELECT fields_json FROM cs_note WHERE customer_id=?',(cust['id'],))]
        reply=api.sent[-1][1]
        passed=len(notes)==1 and notes[0].get('数量')=='100瓶' and 'TEST-BOSS' in reply
        bot.handle_update({'update_id':101,'message':{'from':{'id':200},'chat':{'id':200,'type':'private'},'text':'出表'}})
        content=api.documents[-1][2];(out/'原图采购清单.xlsx').write_bytes(content)
        wb=openpyxl.load_workbook(io.BytesIO(content))
        passed=passed and wb.active.max_row==2 and len(wb.active._images)==1
        results.append({'case':'real_photo_caption','notes':notes,'reply':reply,'passed':passed})
        print('real_photo_caption',passed,flush=True)
        c.close()
    (out/'result.json').write_text(json.dumps(results,ensure_ascii=False,indent=2))
    return all(r['passed'] for r in results)

if __name__=='__main__':
    sys.exit(0 if run(sys.argv[1] if len(sys.argv)>1 else 'audit/round16/real-purchase') else 1)
