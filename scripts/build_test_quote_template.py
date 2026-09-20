"""Build a synthetic 14-column quotation fixture; never a production fallback."""
from pathlib import Path
from io import BytesIO
import base64
import argparse

import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill

PIXEL = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==')
HEADERS = ['ITEM NO.', 'PHOTO', 'DESCRIPTION', 'COLORS', 'PRICE', 'QUANTITY',
           'TOTAL AMOUNT', 'PCS/CTN', 'CTNS', 'G.W/CTN', 'N.W/CTN', 'MEAS', 'T.G.W', 'T-CBM']


def build(path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'TEST QUOTATION'
    wb.properties.title = 'SYNTHETIC TEST FIXTURE - NOT A MERCHANT QUOTATION'
    ws.merge_cells('A1:N2')
    ws['A1'] = 'TEST ONLY / 测试模板 — 非正式报价、不可付款'
    ws['A1'].font = Font(size=20, bold=True, color='C00000')
    ws['A4'] = 'Synthetic merchant / 测试档口'
    ws['A6'] = 'No bank account, contact details or payment QR codes'
    for c, label in enumerate(HEADERS, 1):
        cell = ws.cell(17, c, label)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='243B53')
        ws.column_dimensions[cell.column_letter].width = 16
    ws.column_dimensions['B'].width = 45
    ws.column_dimensions['C'].width = 30
    for r in range(18, 24):
        ws.row_dimensions[r].height = 100
        for c in range(1, 15):
            ws.cell(r, c).alignment = Alignment(vertical='center', wrap_text=True)
        for c in (5, 7):
            ws.cell(r, c).number_format = '"¥"0.00'
    ws.merge_cells('A24:N24')
    ws['A24'] = 'TEST ONLY: quantities rounded to full cartons when carton size is known.'
    ws['A26'] = 'SUMMARY'
    for r, label in [(27, 'TOTAL'), (28, 'DEPOSIT'), (29, 'BALANCE')]:
        ws.cell(r, 1, label)
        ws.cell(r, 7, 0)
    ws['A45'] = 'PLACEHOLDERS ONLY — NO PAYMENT CODES'
    # Five inert image anchors exercise logo/header/footer movement, not real merchant assets.
    for anchor in ('B4', 'F8', 'J8', 'B48', 'H48'):
        im = Image(BytesIO(PIXEL))
        im.width = im.height = 12
        ws.add_image(im, anchor)
    ws.print_options.horizontalCentered = True
    ws.print_area = 'A1:N52'
    wb.save(path)
    return path


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', nargs='?', default='audit/round4/test_quote_template.xlsx')
    print(build(parser.parse_args().output))
