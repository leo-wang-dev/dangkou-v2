"""Shared customer spreadsheet renderer for web and Telegram attachments."""
import io
import json
import os
import openpyxl
from .cs_supplier import normalize

INTERNAL_NOTE_FIELDS = frozenset({'商品编号', '商品类别'})
# 内部留档口径，对客户导出没有意义（档口归属依据 = bot_context/photo/...）。
HIDDEN_NOTE_FIELDS = frozenset({'商品编号', '商品类别', '档口归属依据'})
# 整列都是这些占位值时视为空列，直接不出现在导出里。
EMPTY_TOKENS = ('', '未拍到', '待补充', '模糊', '模糊（待确认）', '—', 'unknown', 'None')


def _column_useful(key, values):
    if key in HIDDEN_NOTE_FIELDS:
        return False
    return any(str(v).strip() not in EMPTY_TOKENS for v in values)


def render_notes(notes, include_status=False, lang='', conn=None, llm=None, texts=None):
    """客户采购清单 Excel。

    - 只保留有实际内容的列（全空/未拍到/待补充的列删掉，档口归属依据等内部字段不导出）。
    - lang 非空且非中文时表头、确认状态和文字内容按客户语言出（cs_i18n 翻译缓存）。
    - texts 可注入自定义翻译函数（HTTP 请求路径用它限流，超量回退中文）。
    """
    if texts is None and lang and lang != '中文' and conn is not None and llm is not None:
        from . import cs_i18n
        texts = lambda values: cs_i18n.translate_texts(conn, llm, lang, values)  # noqa: E731
    items = [normalize(json.loads(n['fields_json'])) for n in notes]
    # Old durable outbox snapshots may predate the public projection. Strip
    # binding metadata again at render time so retries cannot disclose it.
    for item in items:
        for key in INTERNAL_NOTE_FIELDS:
            item.pop(key, None)
    if include_status:
        for item, note in zip(items, notes):
            item['确认状态'] = '待确认' if note['status'] == 'draft' else '已确认'
    keys = []
    for f in items:                              # 动态字段列（保持出现顺序）
        for k in f:
            if k not in keys:
                keys.append(k)
    keys = [k for k in keys if _column_useful(k, [f.get(k, '') for f in items])]
    if texts is not None and lang and lang != '中文':
        headers = texts(['序号', *keys, '商品照片'])
        # 全表单元格一批翻译（走缓存），不要逐行打翻译请求。
        flat = [str(f.get(k, '')) for f in items for k in keys]
        flat = texts(flat)
        width = len(keys)
        rows = [flat[i * width:(i + 1) * width] for i in range(len(items))]
    else:
        headers = ['序号', *keys, '商品照片']
        rows = [[str(f.get(k, '')) for k in keys] for f in items]
    wb = openpyxl.Workbook()
    ws = wb.active
    from openpyxl.drawing.image import Image as XlImage
    ws.append(headers)
    for i, values in enumerate(rows, 1):
        ws.append([i, *values])
    # User fields, including headers, are literal text, never Excel formulas.
    for cells in ws:
        for cell in cells:
            if cell.data_type == 'f':
                cell.data_type = 's'
    for r_i, n in enumerate(notes, start=2):      # 商品图嵌入（有图且装了 pillow）
        if not (n['photo'] and os.path.exists(n['photo'])):
            continue
        try:
            img = XlImage(n['photo'])
            scale = min(120 / img.width, 120 / img.height)
            img.width, img.height = img.width * scale, img.height * scale
            ws.add_image(img, f'{openpyxl.utils.get_column_letter(len(keys)+2)}{r_i}')
            ws.row_dimensions[r_i].height = 96
        except Exception:                         # noqa: BLE001 pillow 缺失/图损坏 → 跳过
            pass
    for col in range(2, len(keys) + 2):           # 列宽
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 16
    ws.column_dimensions[openpyxl.utils.get_column_letter(len(keys)+2)].width = 20
    from openpyxl.styles import Alignment
    for cells in ws:
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
