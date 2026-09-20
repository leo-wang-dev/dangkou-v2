"""Discover category templates and row-aligned images from merchant workbooks."""
from __future__ import annotations

import hashlib
import json
import os
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook


def _text(value) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _label(value) -> str:
    return re.sub(r'\s+', ' ', _text(value)).strip()


def _category_key(name: str) -> str:
    normalized = name.strip().casefold()
    return 'cat_' + hashlib.sha256(normalized.encode('utf-8')).hexdigest()[:12]


def _role(label: str) -> tuple[str, str, str, bool]:
    compact = re.sub(r'\s+', '', label).casefold()
    if compact in {'序号', 'no.', 'no', '编号'}:
        return 'sequence', 'number', 'internal', False
    if any(word in compact for word in ('图片', '照片', 'image', 'picture', 'pic')):
        return 'image', 'image', 'public', True
    if any(word in compact for word in ('产品型号', '商品型号', 'item.no', 'itemno', 'sku', '货号')) or compact == '型号':
        return 'model', 'text', 'public', True
    if any(word in compact for word in ('成本', 'cost')):
        return 'cost', 'money', 'internal', False
    if any(word in compact for word in ('价格', '报价', '单价', '批发价', '出厂价', '拿货价', '售价', 'price')):
        return 'price', 'money', 'internal', False
    if any(word in compact for word in ('链接', '网址', 'url', 'link')):
        return 'note', 'text', 'internal', False
    if '库存' in compact:
        return 'stock', 'number', 'public', True
    if compact in {'备注', '说明', 'note', 'remark'}:
        return 'note', 'text', 'internal', False
    return 'spec', 'text', 'public', any(word in compact for word in ('颜色', '材质', '规格', '名称', '品名'))


def _field_key(label: str, role: str, used: set[str]) -> str:
    preferred = {'sequence': 'sequence', 'model': 'model', 'image': 'image',
                 'stock': 'stock', 'note': 'note'}
    base = preferred.get(role) or 'field_' + hashlib.sha256(label.casefold().encode('utf-8')).hexdigest()[:10]
    key = base
    number = 2
    while key in used:
        key = f'{base}_{number}'
        number += 1
    used.add(key)
    return key


def _header_row(ws) -> int | None:
    best = None
    for row in range(1, min(ws.max_row, 30) + 1):
        labels = [_label(ws.cell(row, col).value) for col in range(1, ws.max_column + 1)]
        nonempty = [value for value in labels if value]
        if len(nonempty) < 2:
            continue
        hints = sum(any(word in value.casefold() for word in
                        ('型号', '货号', 'sku', '图片', '颜色', '价格', '成本', '库存', '规格', '序号'))
                    for value in nonempty)
        score = hints * 1000 + len(nonempty) * 10 - row
        if best is None or score > best[0]:
            best = (score, row)
    return best[1] if best else None


def _merged_sources(ws) -> tuple[dict[tuple[int, int], tuple[int, int]], dict[tuple[int, int], range]]:
    values = {}
    image_rows = {}
    for merged in ws.merged_cells.ranges:
        top = (merged.min_row, merged.min_col)
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                values[(row, col)] = top
        for col in range(merged.min_col, merged.max_col + 1):
            image_rows[(merged.min_row, col)] = range(merged.min_row, merged.max_row + 1)
    return values, image_rows


def _cell_value(ws, sources, row: int, col: int):
    source = sources.get((row, col), (row, col))
    return ws.cell(*source).value


def _extract_images(ws, sheet_key: str, image_dir: Path | None):
    _, merged_image_rows = _merged_sources(ws)
    by_row: dict[int, list[dict]] = {}
    image_dir = Path(image_dir) if image_dir is not None else None
    if image_dir:
        image_dir.mkdir(parents=True, exist_ok=True)
    for index, image in enumerate(ws._images, 1):
        anchor = getattr(image.anchor, '_from', None)
        if anchor is None:
            continue
        row, col = anchor.row + 1, anchor.col + 1
        data = image._data()
        extension = str(getattr(image, 'format', '') or 'png').lower()
        if extension not in {'png', 'jpeg', 'jpg', 'webp', 'gif'}:
            extension = 'png'
        filename = f'{sheet_key}_r{row}_img{index}.{extension}'
        if image_dir:
            (image_dir / filename).write_bytes(data)
            value = filename
        else:
            value = filename
        item = {'path': value, 'sha256': hashlib.sha256(data).hexdigest(),
                'anchor_row': row, 'anchor_col': col}
        target_rows = merged_image_rows.get((row, col), range(row, row + 1))
        for target in target_rows:
            by_row.setdefault(target, []).append(item)
    return by_row, len(ws._images)


def discover_workbook(path, image_dir=None) -> list[dict]:
    """Return one template/product draft for every visible worksheet."""
    path = Path(path)
    max_file = int(os.environ.get('CATALOG_IMPORT_MAX_BYTES', 50 * 1024 * 1024))
    max_unpacked = int(os.environ.get('CATALOG_IMPORT_MAX_UNPACKED_BYTES', 250 * 1024 * 1024))
    if path.stat().st_size > max_file:
        raise ValueError('Excel 文件超过 50MB，请拆分后导入')
    try:
        with zipfile.ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > max_unpacked:
                raise ValueError('Excel 解压内容过大，请删除无关图片或拆分后导入')
    except zipfile.BadZipFile as exc:
        raise ValueError('Excel 文件损坏或不是有效的 xlsx') from exc
    workbook = load_workbook(path, data_only=False)
    if len(workbook.worksheets) > 20:
        raise ValueError('Excel Sheet 超过 20 个，请拆分后导入')
    drafts = []
    for ws in workbook.worksheets:
        if ws.sheet_state != 'visible':
            continue
        if ws.max_row > 20000 or ws.max_column > 200:
            raise ValueError(f'Sheet“{ws.title}”超过 20000 行或 200 列，请拆分后导入')
        if len(ws._images) > 5000:
            raise ValueError(f'Sheet“{ws.title}”图片超过 5000 张，请拆分后导入')
        header_row = _header_row(ws)
        if header_row is None:
            continue
        used: set[str] = set()
        fields = []
        for col in range(1, ws.max_column + 1):
            label = _label(ws.cell(header_row, col).value)
            if not label:
                continue
            role, field_type, visibility, searchable = _role(label)
            fields.append({'key': _field_key(label, role, used), 'label': label,
                           'type': field_type, 'required': False,
                           'visibility': visibility, 'searchable': searchable,
                           'role': role, 'source_column': col})
        if not fields:
            continue
        key = _category_key(ws.title)
        sources, _ = _merged_sources(ws)
        images_by_row, image_count = _extract_images(ws, key, Path(image_dir) if image_dir else None)
        rows = []
        for row_number in range(header_row + 1, ws.max_row + 1):
            direct = any(_text(ws.cell(row_number, field['source_column']).value) for field in fields)
            if not direct and row_number not in images_by_row:
                continue
            data = {}
            for field in fields:
                if field['role'] == 'image':
                    continue
                data[field['key']] = _text(_cell_value(ws, sources, row_number, field['source_column']))
            images = [item['path'] for item in images_by_row.get(row_number, [])]
            # Row order is presentation, not product identity. Keeping the row
            # number out prevents a sort/insert from attaching an old ID to a
            # different physical product.
            fingerprint_payload = {'data': data,
                                   'images': [item['sha256'] for item in images_by_row.get(row_number, [])]}
            rows.append({'_rid': f'{key}-r{row_number}', 'data': data,
                         'source_row': row_number, 'images': images,
                         'image_main': images[0] if images else '',
                         'row_fingerprint': hashlib.sha256(
                             json.dumps(fingerprint_payload, ensure_ascii=False, sort_keys=True).encode('utf-8')).hexdigest()})
        title_values = [_label(ws.cell(row, col).value) for row in range(1, header_row)
                        for col in range(1, ws.max_column + 1) if _label(ws.cell(row, col).value)]
        drafts.append({'key': key, 'name': ws.title.strip() or key,
                       'source_sheet': ws.title, 'title': title_values[0] if title_values else '',
                       'header_row': header_row, 'fields': fields, 'rows': rows,
                       'image_count': image_count})
    return drafts
