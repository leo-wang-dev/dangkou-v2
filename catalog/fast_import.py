"""Bounded, lossless import for sheets that exactly match an approved header.

Ambiguous layouts return None so the existing model parser remains responsible.
"""
from __future__ import annotations

from bisect import bisect_right
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from . import workbook_templates as wbtools


def parse_grouped(template: dict, path, output_dir, *, columns: dict,
                  header_row: int, sheet: str = '') -> dict | None:
    """Group adjacent identical model rows when a reviewed column map is known.

    Conflicting rows are withheld as failures, never resolved by first/last value.
    Every image and non-template cell is surfaced for approval review.
    """
    fields = template['fields']
    if (not 1 <= header_row <= 30 or set(columns) != {f['key'] for f in fields}
            or len(set(columns.values())) != len(columns)
            or any(type(c) is not int or not 1 <= c <= 200 for c in columns.values())):
        return None
    models = [f for f in fields if f['role'] == 'model']
    if len(models) != 1:
        return None
    wbtools.preflight_workbook(path)
    book = load_workbook(path, data_only=False, read_only=False)
    try:
        visible = [ws for ws in book.worksheets if ws.sheet_state == 'visible' and
                   (not sheet or ws.title == sheet) and (ws._cells or ws._images)]
        if len(visible) != 1:
            return None
        ws = visible[0]
        if header_row >= ws.max_row:
            return None
        # The AI mapping must agree with every printed header in the approved
        # template. Empty headers remain a human-review warning below.
        for field in fields:
            col = columns[field['key']]
            printed = [wbtools._label(ws.cell(r, col).value)
                       for r in range(1, header_row + 1)]
            printed = [label for label in printed if label]
            if printed and field['label'] not in printed:
                return None
        model_col = columns[models[0]['key']]
        image_cols = {columns[f['key']] for f in fields if f['role'] == 'image'}
        content_rows = {r for (r, _), cell in ws._cells.items() if wbtools._text(cell.value)}
        image_rows = {im.anchor._from.row + 1 for im in ws._images
                      if getattr(im.anchor, '_from', None) is not None}
        groups = []
        seen_models = set()
        for row in sorted((content_rows | image_rows) - set(range(1, header_row + 1))):
            model = wbtools._text(ws.cell(row, model_col).value)
            if not model:
                return None  # A continuation without a key needs semantic parsing.
            if groups and groups[-1][0] == model:
                groups[-1][1].append(row)
            else:
                if model in seen_models:
                    return None  # Nonadjacent recurrence may be a distinct variant.
                seen_models.add(model)
                groups.append((model, [row]))
        if not groups or all(len(rows) == 1 for _, rows in groups):
            return None
        row_group = {row: index for index, (_, rows) in enumerate(groups) for row in rows}
        failures = []
        for field in fields:
            col = columns[field['key']]
            if (field['role'] != 'image' and
                    not any(wbtools._label(ws.cell(r, col).value)
                            for r in range(1, header_row + 1))):
                failures.append({'source_sheet': ws.title, 'source_rows': [header_row],
                                 'reason': f'{field["label"]}映射到无表头的{get_column_letter(col)}列，请人工核对'})
        for (row, col), cell in ws._cells.items():
            if row > header_row and col not in columns.values() and wbtools._text(cell.value):
                failures.append({'source_sheet': ws.title, 'source_rows': [row],
                                 'reason': f'{get_column_letter(col)}列有未映射内容：{wbtools._text(cell.value)[:80]}，请人工核对'})
        out = Path(output_dir)
        out.mkdir(parents=True, exist_ok=True)
        group_images = [[] for _ in groups]
        for index, image in enumerate(ws._images, 1):
            anchor = getattr(image.anchor, '_from', None)
            if anchor is None or anchor.row + 1 not in row_group:
                return None
            row, col = anchor.row + 1, anchor.col + 1
            extension = str(getattr(image, 'format', '') or 'png').lower()
            if extension not in {'png', 'jpeg', 'jpg', 'webp', 'gif'}:
                extension = 'png'
            name = f'r{row}_c{col}_i{index}.{extension}'
            (out / name).write_bytes(image._data())
            group_images[row_group[row]].append(name)
            if col not in image_cols:
                failures.append({'source_sheet': ws.title, 'source_rows': [row],
                                 'reason': f'图片锚点位于{get_column_letter(col)}列而非模板图片列，请人工核对归属'})
        products = []
        for index, (model, rows) in enumerate(groups):
            product = {'source_sheet': ws.title, 'source_rows': rows,
                       'images': group_images[index], 'image_count': len(group_images[index]),
                       'image_main': group_images[index][0] if group_images[index] else ''}
            conflicts = []
            for field in fields:
                if field['role'] == 'image':
                    continue
                values = list(dict.fromkeys(wbtools._text(ws.cell(row, columns[field['key']]).value)
                                            for row in rows))
                values = [value for value in values if value]
                if len(values) > 1:
                    conflicts.append(f'{field["label"]}有不同原值：{", ".join(values[:4])}')
                else:
                    product[field['key']] = values[0] if values else ''
            if conflicts:
                failures.append({'source_sheet': ws.title, 'source_rows': rows,
                                 'reason': '同型号多行冲突，未自动建商品：' + '；'.join(conflicts)})
            else:
                products.append(product)
        return {'products': products, 'failures': failures, 'vendor': None}
    finally:
        book.close()


def parse_structured(template: dict, path, output_dir, *, sheet: str = '') -> dict | None:
    wbtools.preflight_workbook(path)
    book = load_workbook(path, data_only=False, read_only=False)
    try:
        visible = [ws for ws in book.worksheets if ws.sheet_state == 'visible' and
                   (not sheet or ws.title == sheet) and
                   (any(wbtools._text(cell.value) for cell in ws._cells.values()) or ws._images)]
        if len(visible) != 1:
            return None
        ws = visible[0]
        fields = template['fields']
        expected = [wbtools._label(f['label']) for f in fields]
        if len(expected) != len(set(expected)) or not expected:
            return None
        matches = []
        for row in range(1, min(ws.max_row, 30) + 1):
            found = [(col, wbtools._label(ws.cell(row, col).value))
                     for col in range(1, min(ws.max_column, 200) + 1)]
            found = [(col, label) for col, label in found if label]
            if [label for _, label in found] == expected:
                matches.append((row, [col for col, _ in found], False, False))
            elif (len(found) == len(expected) and
                  sum(field['role'] == 'model' and field['label'] == '型号' and
                      label.casefold() in {'picture', 'image', '图片', '照片'}
                      for (_, label), field in zip(found, fields)) == 1 and
                  all(label == field['label'] or
                      (field['role'] == 'model' and field['label'] == '型号' and
                       label.casefold() in {'picture', 'image', '图片', '照片'})
                      for (_, label), field in zip(found, fields))):
                matches.append((row, [col for col, _ in found], False, True))
            elif (len(found) == len(expected) - 1 and [label for _, label in found] == expected[:-1]
                  and fields[-1]['role'] != 'image' and found
                  and not wbtools._text(ws.cell(row, found[-1][0] + 1).value)
                  and any(wbtools._text(ws.cell(r, found[-1][0] + 1).value)
                          for r in range(row + 1, ws.max_row + 1))):
                matches.append((row, [col for col, _ in found] + [found[-1][0] + 1], True, False))
        if len(matches) != 1:
            return None
        header, columns, unlabeled_tail, picture_header = matches[0]
        key_col = next((columns[i] for i, f in enumerate(fields)
                        if f['role'] != 'image' and any(term in f['label'] for term in ('名称', '品名'))), None)
        model_col = next((columns[i] for i, f in enumerate(fields) if f['role'] == 'model'), None)
        if key_col is None and model_col is None:
            return None
        model_only = key_col is None
        picture_model = (model_only and model_col is not None and
                         (picture_header or re.search(r'picture|image|图片|照片',
                          next(f['label'] for f in fields if f['role'] == 'model'), re.I)))
        source_map, _ = wbtools._merged_sources(ws)
        companions = {}
        if picture_model:
            for field, col in zip(fields, columns):
                if ('packing' in field['label'].casefold() or '包装' in field['label']):
                    neighbor = col + 1
                    if (neighbor not in columns and
                            not wbtools._text(ws.cell(header, neighbor).value)):
                        companions[col] = neighbor
        unmapped_rows = {r for (r, c), cell in ws._cells.items()
                         if r > header and c not in columns and c not in companions.values()
                         and wbtools._text(cell.value)}
        if unmapped_rows:
            return None
        starts = []
        footer_start = None
        supplementary = []
        for row in range(header + 1, ws.max_row + 1):
            direct = {col: wbtools._text(ws.cell(row, col).value) for col in columns}
            if picture_model:
                model = direct[model_col]
                peers = any(value for col, value in direct.items() if col != model_col)
                if (model and peers and
                        re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9._/+-]{0,39}', model)):
                    if footer_start is not None:
                        return None
                    starts.append(row)
                elif model and not peers and model.casefold() in {'remark', 'remarks', 'terms', '条款', '备注'} and starts:
                    footer_start = row
                elif model and starts and footer_start is None:
                    supplementary.append((row, model))
                elif model and not starts:
                    return None
                continue
            if direct.get(key_col, ''):
                starts.append(row)
            elif direct.get(model_col, ''):
                starts.append(row)
            elif model_only and any(direct[col] for field, col in zip(fields, columns)
                                    if field['role'] != 'image'):
                return None  # Unknown product boundary: let the agent decide.
            elif any(direct.values()) and not starts:
                return None
        if not starts:
            return None
        meaningful_last = max((r for (r, _), cell in ws._cells.items()
                               if wbtools._text(cell.value)), default=header)
        anchors = [getattr(image.anchor, '_from', None) for image in ws._images]
        meaningful_last = max([meaningful_last,
                               *(anchor.row + 1 for anchor in anchors if anchor is not None)])
        for merged in ws.merged_cells.ranges:
            if (wbtools._text(ws.cell(merged.min_row, merged.min_col).value)
                    or any(anchor is not None and merged.min_row <= anchor.row + 1 <= merged.max_row
                           and merged.min_col <= anchor.col + 1 <= merged.max_col for anchor in anchors)):
                meaningful_last = max(meaningful_last, merged.max_row)
        if footer_start is not None:
            meaningful_last = footer_start - 1
        products = []
        inherited_model = ''
        for index, start in enumerate(starts):
            end = starts[index + 1] if index + 1 < len(starts) else meaningful_last + 1
            item = {'source_sheet': ws.title, 'source_rows': list(range(start, end)), 'images': []}
            for field, col in zip(fields, columns):
                if field['role'] == 'image':
                    continue
                values = []
                for row in range(start, end):
                    raw = (wbtools._cell_value(ws, source_map, row, col) if row == start
                           else ws.cell(row, col).value)
                    value = wbtools._text(raw)
                    if picture_model and field['role'] == 'model' and row != start:
                        continue
                    if col in companions:
                        companion = wbtools._text(ws.cell(row, companions[col]).value)
                        if companion:
                            value = f'{value} {companion}'.strip()
                    if value and (not values or value != values[-1]):
                        values.append(value)
                item[field['key']] = '\n'.join(values)
                if field['role'] == 'model':
                    if item[field['key']]:
                        inherited_model = item[field['key']]
                    else:
                        item[field['key']] = inherited_model
            products.append(item)
        out = Path(output_dir)
        out.mkdir(parents=True, exist_ok=True)
        failures = []
        for row, value in supplementary:
            failures.append({'source_sheet': ws.title, 'source_rows': [row],
                             'reason': f'{get_column_letter(model_col)}列非型号文字：{value[:80]}，请核对'})
        if footer_start is not None:
            failures.append({'source_sheet': ws.title,
                             'source_rows': list(range(footer_start, ws.max_row + 1)),
                             'reason': '表尾交易条款不作为商品，请人工核对'})
        if unlabeled_tail:
            failures.append({'source_sheet': ws.title, 'source_rows': [header],
                             'reason': f'末列原表头为空，按已审核模板映射为“{fields[-1]["label"]}”，请核对'})
        for field, col in zip(fields, columns):
            if field['role'] != 'image':
                continue
            for row in range(header + 1, ws.max_row + 1):
                value = wbtools._text(ws.cell(row, col).value)
                if not value:
                    continue
                if value != '#NAME?' and 'DISPIMG' not in value.upper():
                    return None
                position = bisect_right(starts, row) - 1
                if position < 0:
                    return None
                failures.append({'source_sheet': ws.title, 'source_rows': [row],
                                 'reason': f'{field["label"]}含未解析的单元格图片引用，请核对该商品图片'})
        for index, image in enumerate(ws._images, 1):
            anchor = getattr(image.anchor, '_from', None)
            if anchor is None:
                return None
            row = anchor.row + 1
            if footer_start is not None and row >= footer_start:
                failures.append({'source_sheet': ws.title, 'source_rows': [row],
                                 'reason': '表尾图片未归属商品，请核对'})
                continue
            position = bisect_right(starts, row) - 1
            if position < 0:
                failures.append({'source_sheet': ws.title, 'source_rows': [row],
                                 'reason': '表头前或商品行前的图片未归属商品，请核对'})
                continue
            image_columns = {col for field, col in zip(fields, columns)
                             if field['role'] == 'image' or (picture_model and col == model_col)}
            if anchor.col + 1 not in image_columns:
                failures.append({'source_sheet': ws.title, 'source_rows': [row],
                                 'reason': '图片锚点不在模板图片列，已按同一商品行归属，请核对'})
            extension = str(getattr(image, 'format', '') or 'png').lower()
            if extension not in {'png', 'jpeg', 'jpg', 'webp', 'gif'}:
                extension = 'png'
            name = f'r{row}_c{anchor.col + 1}_i{index}.{extension}'
            (out / name).write_bytes(image._data())
            products[position]['images'].append(name)
        for item in products:
            item['image_count'] = len(item['images'])
            item['image_main'] = item['images'][0] if item['images'] else ''
        return {'products': products, 'failures': failures, 'vendor': None}
    finally:
        book.close()
