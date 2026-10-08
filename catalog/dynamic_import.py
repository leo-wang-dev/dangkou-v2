"""Build and apply approval payloads for Sheet-defined category imports."""
from __future__ import annotations

import json
import hashlib
import os
import secrets
import shutil
import re
from pathlib import Path
import time
from concurrent.futures import ThreadPoolExecutor

from openpyxl import load_workbook

from . import agent, ai_extract, dynamic_catalog, inner_code, workbook_templates, fast_import


def _clean_template(draft: dict) -> dict:
    return {'key': draft['key'], 'name': draft['name'],
            'source_sheet': draft['source_sheet'], 'storage': 'dynamic',
            'fields': [{key: value for key, value in field.items() if key != 'source_column'}
                       for field in draft['fields']]}


def _field_signature(fields: list[dict]) -> list[tuple]:
    return [(field['label'].casefold(), field['type'], field['visibility'], field['role'])
            for field in fields]


def _source_rows(conn, category_key: str, source_key: str, source_sheet: str) -> list[dict]:
    # Identity spans reuploaded filenames; snapshot this same scope at approval.
    return dynamic_catalog.list_products(conn, category_key)


def _source_snapshot(rows: list[dict]) -> str:
    values = [{key: row.get(key) for key in (
        'id', 'inner_code', 'data', 'status', 'images', 'source_row',
        'row_fingerprint', 'cs_visible', 'supplier')} for row in sorted(rows, key=lambda value: value['id'])]
    return hashlib.sha256(json.dumps(values, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def _model_identity(row: dict, model_keys: list[str]):
    data = row.get('data') or {}
    values = tuple(str(data.get(key) or '').strip().casefold() for key in model_keys)
    return (row.get('supplier', ''), *values) if any(values) else None


def _classify_rows(existing: list[dict], incoming: list[dict], model_keys: list[str], variant_keys=None) -> dict:
    unused = list(existing)
    variant_keys = list(variant_keys) if variant_keys is not None else sorted(
        set().union(*(row.get('data', {}).keys() for row in incoming))
        - set(model_keys) - {'price', 'cost', 'stock', 'note'})
    seen = set()
    for index, row in enumerate(incoming):
        identity = _model_identity(row, model_keys)
        if identity:
            signature = (identity, tuple(str(row.get('data', {}).get(key) or '').strip() for key in variant_keys))
            if signature in seen:
                raise ValueError(f"供应商 {row.get('supplier', '')} 型号 {[row.get('data', {}).get(key, '') for key in model_keys]} 重复规格，无法确定更新对象（行 {index + 1}）")
            seen.add(signature)
    new, update = [], []
    matches: dict[int, dict] = {}
    # Reserve every exact content match first. This disambiguates duplicate
    # models when just one colour/spec row changed.
    for index, draft in enumerate(incoming):
        exact = [row for row in unused if draft.get('row_fingerprint')
                 and row['row_fingerprint'] == draft['row_fingerprint']
                 and row.get('data', {}) == draft.get('data', {})
                 and row.get('supplier', '') == draft.get('supplier', '')]
        if len(exact) > 1:
            raise ValueError(f"供应商 {draft.get('supplier', '')} 型号 {[draft.get('data', {}).get(key, '') for key in model_keys]} 有多个相同商品，无法确定更新对象（行 {index + 1}）")
        match = exact[0] if exact else None
        if match is not None:
            unused.remove(match)
            matches[index] = match
    for index, draft in enumerate(incoming):
        match = matches.get(index)
        if match is None:
            identity = _model_identity(draft, model_keys)
            candidates = [row for row in unused if identity and _model_identity(row, model_keys) == identity]
            keys = variant_keys
            compatible = [row for row in candidates if all(
                not str(draft.get('data', {}).get(key) or '').strip()
                or str(row.get('data', {}).get(key) or '').strip() == str(draft['data'][key]).strip()
                for key in keys)]
            if len(compatible) > 1:
                raise ValueError(f"供应商 {draft.get('supplier', '')} 型号 {[draft.get('data', {}).get(key, '') for key in model_keys]} 规格不明确，无法确定更新对象（行 {index + 1}）")
            candidates = [row for row in candidates if all(
                str(row.get('data', {}).get(key) or '').strip() == str(draft.get('data', {}).get(key) or '').strip()
                for key in keys)]
            if compatible and not candidates:
                raise ValueError(f"供应商 {draft.get('supplier', '')} 型号 {[draft.get('data', {}).get(key, '') for key in model_keys]} 缺少规格，无法确定更新对象（行 {index + 1}）")
            match = candidates[0] if candidates else None
        if match is None:
            new.append(draft)
            continue
        if match in unused:
            unused.remove(match)
        if match['row_fingerprint'] != draft['row_fingerprint'] or match.get('data') != draft.get('data'):
            update.append([match, draft])
    return {'new': new, 'update': update, 'delist': []}


class ParsedRows(list):
    def __init__(self, rows=(), *, failures=(), coverage=None):
        super().__init__(rows)
        self.failures = list(failures)
        self.coverage = coverage or {}


def _classify_template_rows(existing, incoming, template):
    failures = list(getattr(incoming, 'failures', []))
    accepted = []
    # Validate each proposal against the whole accepted batch. A bad row does
    # not erase good rows, and identity collisions remain explicit failures.
    for row in incoming:
        row = {**row, 'supplier': row.get('supplier', template.get('supplier', ''))}
        try:
            _classify_rows(existing, [*accepted, row],
                [f['key'] for f in template['fields'] if f['role'] == 'model'],
                [f['key'] for f in template['fields'] if f['role'] == 'spec'])
            accepted.append(row)
        except ValueError as exc:
            failures.append({'source_sheet': row.get('source_sheet', ''),
                'source_rows': row.get('source_rows', []), 'reason': str(exc)})
    drafts = _classify_rows(existing, accepted,
        [f['key'] for f in template['fields'] if f['role'] == 'model'],
        [f['key'] for f in template['fields'] if f['role'] == 'spec'])
    drafts['_accepted'] = len(accepted)
    drafts['_failures'] = failures
    drafts['_coverage'] = getattr(incoming, 'coverage', {})
    return drafts


def _product_metadata(payload):
    accepted = 0
    for section in payload.get('sheets', []):
        drafts = section.get('drafts', {})
        accepted += drafts.pop('_accepted', len(drafts.get('new', [])) + len(drafts.get('update', [])))
        section['failures'] = drafts.pop('_failures', [])
        section['coverage'] = drafts.pop('_coverage', {})
    if not accepted:
        reasons = [f['reason'] for s in payload.get('sheets', []) for f in s.get('failures', [])]
        raise ValueError('没有有效商品可审批：' + '; '.join(reasons)[:300])
    return payload


def _file_sha256(path: str) -> str:
    try:
        digest = hashlib.sha256()
        with open(path, 'rb') as handle:
            for chunk in iter(lambda: handle.read(1048576), b''):
                digest.update(chunk)
        return digest.hexdigest()
    except OSError:
        return ''


_TEMPLATE_ROLES = {'model', 'image', 'sequence', 'price', 'cost', 'stock', 'note', 'spec'}


def _agent_template_sheets(xlsx_path: str, work_dir) -> list[dict] | None:
    """模板阶段表头发现（慢路，子代理）：多级表头、无标头图片列靠语义找齐。

    把 agent 的 columns 转成 discover_workbook 的 discovered 形状
    （label/key 用 _field_key，图片列 role=image type=image，image_count=0）。
    失败（异常或 0 个有效 Sheet）返回 None——调用方回落 discover_workbook。
    """
    try:
        result = agent.parse_dynamic_template(xlsx_path, work_dir)
    except Exception as exc:  # noqa: BLE001
        print(f'[dynamic_import] 子代理模板发现失败，回落代码表头发现：{exc}', flush=True)
        return None
    return _convert_template_sheets(result.get('sheets') or [])


def _qwen_template_sheets(xlsx_path: str) -> list[dict] | None:
    """模板阶段表头发现（快路，qwen 关思考直调，约 5-8 秒/份）。

    代码抽证据（前几行网格+每列图片锚点+合并区）交给 qwen3.8-max 判断，
    与子代理共用 columns→discovered 转换；失败回落子代理再回落代码。
    """
    try:
        evidence = workbook_templates.extract_header_evidence(xlsx_path)
    except Exception as exc:  # noqa: BLE001
        print(f'[dynamic_import] 表头证据抽取失败：{exc}', flush=True)
        return None
    data = ai_extract.discover_headers(evidence)
    if data is None:
        return None
    return _convert_template_sheets(data)


def prefetch_xls_template(xls_path: str, source_key: str) -> dict | None:
    """Start semantic header discovery while the XLS converter is running."""
    try:
        evidence = workbook_templates.extract_xls_header_evidence(xls_path)
        if not evidence:
            return None
        with ThreadPoolExecutor(max_workers=2) as pool:
            supplier = pool.submit(ai_extract.guess_supplier, source_key, [x['title'] for x in evidence])
            headers = pool.submit(ai_extract.discover_headers, evidence)
            result = _convert_template_sheets(headers.result() or [])
            supplier_guess = supplier.result()
        return {'evidence': evidence, 'discovered': result, 'supplier_guess': supplier_guess}
    except Exception as exc:  # Conversion remains the authoritative path.
        print(f'[dynamic_import] XLS 并行表头预读失败，改用转换后识别：{exc}', flush=True)
        return None


def _safe_prefetched_xls(prefetched: dict | None, converted_evidence: list[dict]) -> bool:
    """A drawing in an unlabelled column needs the full converted evidence."""
    if not prefetched or not prefetched.get('discovered'):
        return False
    early = prefetched['evidence']
    if [x['title'] for x in early] != [x['title'] for x in converted_evidence]:
        return False
    for sheet in converted_evidence:
        source = next(x for x in early if x['title'] == sheet['title'])
        # Only reuse the AI answer when both readers saw the same grid.  An
        # image-only column still needs the converted drawing anchors.
        for row_index, row in enumerate(sheet['grid']):
            before = source['grid'][row_index] if row_index < len(source['grid']) else []
            for col, value in enumerate(row):
                if value.strip() != (before[col].strip() if col < len(before) else ''):
                    return False
        for raw_col in sheet['图片锚点列']:
            col = int(raw_col)
            if not any(col <= len(row) and row[col - 1].strip() for row in sheet['grid']):
                return False
    return True


def _convert_template_sheets(items: list) -> list[dict] | None:
    """子代理/qwen 共用：columns 输出 → discovered 形状；0 个有效 Sheet 返回 None。"""
    discovered = []
    for item in items:
        if not isinstance(item, dict):
            continue
        title = workbook_templates._label(item.get('title'))
        columns = item.get('columns')
        # title 是商品阶段按 Sheet 名回查模板的钥匙，空标题无法回查。
        if not title or not isinstance(columns, list):
            continue
        used: set[str] = set()
        fields = []
        for spec in columns:
            if not isinstance(spec, dict):
                continue
            try:
                col = int(spec.get('col') or 0)
            except (TypeError, ValueError):
                continue
            if col < 1 or col > 200:
                continue
            label = workbook_templates._label(spec.get('label'))
            role = spec.get('role')
            fb_role, fb_type, fb_visibility, fb_searchable = workbook_templates._role(label)
            if role not in _TEMPLATE_ROLES:
                role = fb_role
            if role == 'image':
                label = label or '图片'
            elif not label:
                # 无表头又非图片角色=空列，跳过
                continue
            field_type = spec.get('type')
            if role == 'image' or field_type == 'image':
                field_type = 'image'
            elif field_type not in ('text', 'number', 'money'):
                field_type = fb_type
            visibility = spec.get('visibility')
            if visibility not in ('public', 'internal'):
                visibility = fb_visibility
            searchable = spec.get('searchable')
            if not isinstance(searchable, bool):
                searchable = fb_searchable
            fields.append({'key': workbook_templates._field_key(label, role, used),
                           'label': label, 'type': field_type, 'required': False,
                           'visibility': visibility, 'searchable': searchable,
                           'role': role, 'source_column': col})
        if not fields:
            continue
        try:
            header_row = max(1, int(item.get('header_row') or 1))
        except (TypeError, ValueError):
            header_row = 1
        key = workbook_templates._category_key(title)
        discovered.append({'key': key, 'name': title, 'source_sheet': title,
                           'title': title, 'header_row': header_row,
                           'fields': fields, 'rows': [], 'image_count': 0})
    if not discovered:
        print('[dynamic_import] 模板发现 0 个有效 Sheet，回落下一条路径', flush=True)
        return None
    return discovered


def _products_into_template(conn, xlsx_path: str, work_dir, template: dict, *,
                            source_key: str, doc_id: int | None) -> dict:
    """手工分类直灌：整簿各 Sheet 全部按该分类已审批模板解析成商品草稿。"""
    incoming = _agent_rows(template, xlsx_path, work_dir)
    if incoming is None:
        raise ValueError('解析服务暂不可用，请稍后重试导入')
    source_rows = _source_rows(conn, template['key'], source_key, template['source_sheet'])
    existing = [row for row in source_rows if row['status'] != 'delisted']
    section = {'template': template, 'template_action': 'reuse',
               'expected_version': template['version'], 'title': template['name'],
               'header_row': None, 'image_count': 0,
               'source_sheet': template['source_sheet'],
               'source_snapshot': _source_snapshot(source_rows),
               'drafts': _classify_template_rows(existing, incoming, template)}
    return {'kind': 'template_import', 'phase': 'products', 'doc_id': doc_id,
            'filename': os.path.basename(xlsx_path),
            'template_doc_id': None, 'source_key': source_key,
            'mode': 'existing', 'category_key': template['key'],
            'work_dir': str(work_dir), 'sheets': [section]}


def _agent_rows(template: dict, xlsx_path: str, work_dir, sheet: str = '') -> list[dict] | None:
    """Parse products with source coverage and reusable checkpoints."""
    identity = hashlib.sha256(json.dumps({'template': template, 'sheet': sheet,
        'source': _file_sha256(xlsx_path)}, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    root = Path(work_dir); root.mkdir(parents=True, exist_ok=True)
    checkpoint = root / (identity + '.checkpoint.json')
    candidates = [checkpoint]
    if root.name.startswith('attempt-'):
        candidates += sorted(root.parent.glob('attempt-*/' + checkpoint.name))
    for cached in candidates:
        if not cached.is_file():
            continue
        try:
            saved = json.loads(cached.read_text())
            for row in saved['rows']:
                hashes = []
                for name in row['images']:
                    source = (cached.parent / name).resolve()
                    target = (root / name).resolve()
                    if not source.is_relative_to(cached.parent.resolve()) or not target.is_relative_to(root.resolve()):
                        raise ValueError('invalid checkpoint image path')
                    hashes.append(_file_sha256(str(source)))
                    if cached != checkpoint:
                        target.parent.mkdir(parents=True, exist_ok=True)
                        shutil.copyfile(source, target)
                fingerprint = hashlib.sha256(json.dumps({'data': row['data'], 'images': hashes}, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
                if fingerprint != row['row_fingerprint']:
                    raise ValueError('checkpoint image content changed')
            return ParsedRows(saved['rows'], failures=saved['failures'], coverage=saved['coverage'])
        except (OSError, ValueError, KeyError, TypeError):
            continue  # Reparse when a checkpoint or its artifacts are incomplete.
    # Structural rows are evidence, never a fallback source of product data.
    manifest = workbook_templates.discover_workbook(xlsx_path, include_images=False)
    evidence = {d['source_sheet']: set(d.get('observed_rows', [r['source_row'] for r in d['rows']]))
                for d in manifest if not sheet or d['source_sheet'] == sheet}
    source_bounds = {d['source_sheet']: d['source_max_row'] for d in manifest if d['source_sheet'] in evidence}
    started = time.monotonic()
    parse_dir = root / identity[:16]; parse_dir.mkdir(exist_ok=True)
    try:
        result = fast_import.parse_structured(template, xlsx_path, parse_dir, sheet=sheet)
        if result is None:
            for _ in range(3):
                mapping = ai_extract.map_approved_fields(xlsx_path, template['fields'], sheet)
                if not mapping:
                    continue
                result = fast_import.parse_grouped(
                    template, xlsx_path, parse_dir, columns=mapping['columns'],
                    header_row=mapping['header_row'], sheet=sheet)
                if result is not None:
                    break
        if result is None:
            result = agent.parse_dynamic(template, xlsx_path, parse_dir, sheet=sheet)
    except Exception as exc:
        return ParsedRows(failures=[{'source_sheet': sheet, 'source_rows': sorted(set().union(*evidence.values())) if evidence else [],
            'reason': '解析服务暂不可用，请重试：' + str(exc)[:160]}], coverage={'uncertain': True})
    failures = []
    for failure in result.get('failures') or []:
        if isinstance(failure, dict):
            refs = failure.get('source_rows') or []
            reason = str(failure.get('reason') or '模型未能解析此区域')[:300]
            if not isinstance(refs, list) or any(type(r) is not int or not 1 <= r <= 20000 for r in refs):
                refs, reason = [], reason + '（失败位置格式无效，请核对来源区域）'
            failures.append({'source_sheet': str(failure.get('source_sheet') or sheet),
                'source_rows': refs, 'reason': reason})
    rows, represented = [], set()
    uncertain = False
    for index, product in enumerate(result.get('products') or []):
        if not isinstance(product, dict):
            failures.append({'source_sheet': sheet, 'source_rows': [], 'reason': f'第 {index+1} 个结果不是商品对象'})
            continue
        source_sheet = str(product.get('source_sheet') or sheet or (next(iter(evidence)) if len(evidence) == 1 else ''))
        refs = product.get('source_rows') or ([product['source_row']] if product.get('source_row') else [])
        reason = ''
        if not isinstance(refs, list) or any(type(r) is not int or r < 1 for r in refs):
            reason, refs = '来源行号格式无效', []
        if refs and (source_sheet not in source_bounds or any(r > source_bounds[source_sheet] for r in refs)):
            reason = '来源行不在已观察到的数据区域，请核对表头或商品分组'
        if not refs:
            uncertain = True
        # 模型偶发把 field_xxxx…key 抄短一两位（实发：glm-5.3 把 field_0f4371a7a9 写成
        # field_0f4371a7，价格整列被静默丢弃）。非模板键若恰好是唯一模板键的前缀
        # （≥8 字符，label 撞名排除），按前缀回映射；无法唯一映射的键打印告警，不再无声蒸发。
        template_keys = {f['key'] for f in template['fields'] if f.get('role') != 'image'}
        for extra in set(product) - template_keys - {'source_sheet', 'source_rows', 'source_row',
                                                     'supplier', 'image_main', 'images', 'image_count'}:
            candidates = [key for key in template_keys if key.startswith(extra) and len(extra) >= 8]
            if len(candidates) == 1:
                product.setdefault(candidates[0], product[extra])
            else:
                print(f'[dynamic_import] agent 输出含未知字段键 {extra!r}（值首段：{str(product[extra])[:24]!r}），已丢弃', flush=True)
        data = {f['key']: '' if product.get(f['key'], product.get(f['label'])) is None else str(product.get(f['key'], product.get(f['label'])))
                for f in template['fields'] if f.get('role') != 'image'}
        if not any(v.strip() for v in data.values()):
            reason = reason or '商品字段全部为空'
        images = product.get('images') or []
        if not isinstance(images, list):
            images, reason = [], '图片列表格式无效'
        images = [str(v) for v in images if v]
        if product.get('image_main') and product['image_main'] not in images:
            images.insert(0, str(product['image_main']))
        hashes = []
        for name in images:
            image = (parse_dir / name).resolve()
            if not image.is_relative_to(parse_dir.resolve()) or not image.is_file():
                reason = '图片文件缺失或路径无效：' + name[:80]
                break
            try:
                from PIL import Image
                with Image.open(image) as decoded:
                    decoded.verify()
                hashes.append(_file_sha256(str(image)))
            except Exception:
                reason = '图片文件无法读取：' + name[:80]
        if product.get('image_count') is not None and product['image_count'] != len(images):
            reason = '图片数量与解析结果不一致'
        if reason:
            failures.append({'source_sheet': source_sheet, 'source_rows': refs, 'reason': reason})
            continue
        represented.update((source_sheet, r) for r in refs)
        images = [str(Path(identity[:16]) / name) for name in images]
        rows.append({'_rid': identity[:16] + '-' + str(index), 'data': data,
            'supplier': str(product.get('supplier') or product.get('供应商') or result.get('vendor') or template.get('supplier') or '').strip()[:40],
            'images': images, 'image_main': images[0] if images else '',
            'source_sheet': source_sheet, 'source_rows': refs, 'source_row': refs[0] if refs else None,
            'row_fingerprint': hashlib.sha256(json.dumps({'data': data, 'images': hashes}, ensure_ascii=False, sort_keys=True).encode()).hexdigest()})
    failed_refs = {(f['source_sheet'], r) for f in failures for r in f['source_rows'] if type(r) is int}
    for name, candidates in evidence.items():
        missing = sorted(r for r in candidates if (name, r) not in represented | failed_refs)
        if missing:
            failures.append({'source_sheet': name, 'source_rows': missing,
                'reason': '来源覆盖未确认：这些物理行未被商品或失败区域引用，可能包含跨行规格/图片；请核对，不能认定为已解析商品'})
    coverage = {'candidate_rows': sum(map(len, evidence.values())), 'represented_rows': len(represented),
        'logical_products': len(rows), 'uncertain': uncertain or bool(failures),
        'note': '物理候选行不是逻辑商品数量', 'model_seconds': round(time.monotonic()-started, 3),
        'source_rows': {name: sorted(values) for name, values in evidence.items()}}
    if not rows:
        return ParsedRows(failures=failures, coverage=coverage)
    saved = {'rows': rows, 'failures': failures, 'coverage': coverage}
    # Checkpoint even partial successes. A crash resumes these results; an explicit
    # re-upload creates a new job/workspace and retries the whole sheet/document.
    temporary = checkpoint.with_suffix('.tmp'); temporary.write_text(json.dumps(saved, ensure_ascii=False)); temporary.replace(checkpoint)
    return ParsedRows(rows, failures=failures, coverage=coverage)


def manual_template_payload(name: str, fields_in: list[dict]) -> dict:
    """手工建分类（不经 Excel）：字段 key 生成与 Excel 导入同源，身份稳定。"""
    used: set[str] = set()
    fields = []
    for item in fields_in:
        role, field_type, visibility, searchable = workbook_templates._role(item['label'])
        fields.append({'key': workbook_templates._field_key(item['label'], role, used),
                       'label': item['label'], 'type': field_type, 'required': False,
                       'visibility': item.get('visibility') or visibility,
                       'searchable': searchable, 'role': role})
    key = workbook_templates._category_key(name)
    template = {'key': key, 'name': name, 'source_sheet': name,
                'storage': 'dynamic', 'fields': fields}
    return {'kind': 'template_import', 'phase': 'template', 'manual': True,
            'doc_id': None, 'source_key': f'manual:{key}', 'mode': 'new',
            'category_key': None, 'work_dir': '',
            'sheets': [{'template': template, 'template_action': 'create',
                        'expected_version': 0, 'title': name, 'header_row': None,
                        'image_count': 0, 'source_sheet': name,
                        'source_discovered_sheet': '', 'source_snapshot': '',
                        'drafts': {'new': [], 'update': [], 'delist': []}}]}


def _pending_claim_conflict(conn, key: str, signature) -> bool:
    """同名 Sheet 的在途工单已认领该分类 key 且表头不同——
    同 key 不同表头无法合并，后到的文件必须开独立分类，避免生成注定批不过的工单。
    表头相同的在途认领不算冲突：两张工单先后批准会按增量并入同一分类。"""
    rows = conn.execute("SELECT payload FROM approval_ticket "
                        "WHERE status='pending' AND ticket_type='template_import'").fetchall()
    for row in rows:
        try:
            payload = json.loads(row['payload'])
        except (TypeError, ValueError):
            continue
        for section in payload.get('sheets') or []:
            if section.get('template', {}).get('key') != key:
                continue
            if _field_signature(section['template']['fields']) != signature:
                return True
    return False


def _unique_new_template(conn, template: dict, *, source_key: str, doc_id: int | None) -> dict:
    """Force mode=new to create a distinct category even when the sheet name exists."""
    candidate = dict(template)
    base_key = template['key']
    base_name = template['name']
    signature = _field_signature(template['fields'])
    suffix = 0
    while True:
        try:
            current = dynamic_catalog.get_template(conn, candidate['key'])
        except KeyError:
            current = None
        if current is None and not _pending_claim_conflict(conn, candidate['key'], signature):
            return candidate
        suffix += 1
        seed = f'{base_key}\0new\0{source_key}\0{doc_id or ""}\0{suffix}'
        candidate['key'] = 'cat_' + hashlib.sha256(seed.encode('utf-8')).hexdigest()[:12]
        candidate['name'] = f'{base_name} ({suffix})'


def _template_sections(conn, discovered_sheets: list[dict], *, source_key: str,
                       doc_id: int | None, mode: str | None,
                       category_key: str | None) -> list[dict]:
    """Build schema-only sections.  No product rows or embedded images are attached."""
    if mode == 'existing':
        try:
            target = dynamic_catalog.get_template(conn, category_key)
        except KeyError as exc:
            raise ValueError('未知目标分类，请先从商品管理中选择已有分类') from exc
        if target['storage'] != 'dynamic':
            raise ValueError(f'分类 {target["name"]} 是旧版固定分类，已下线；请新建动态分类导入')
        discovered = discovered_sheets[0] if discovered_sheets else {}
        return [{'template': target, 'template_action': 'reuse',
                 'expected_version': target['version'],
                 'title': discovered.get('title', ''),
                 'header_row': discovered.get('header_row'), 'image_count': 0,
                 'source_sheet': target['source_sheet'], 'source_discovered_sheet': discovered.get('source_sheet', ''),
                 'source_snapshot': '',
                 'drafts': {'new': [], 'update': [], 'delist': []}}]

    sections = []
    for discovered in discovered_sheets:
        template = _clean_template(discovered)
        if mode == 'new':
            template = _unique_new_template(conn, template, source_key=source_key, doc_id=doc_id)
        try:
            current = dynamic_catalog.get_template(conn, template['key'])
        except KeyError:
            current = None
        if mode == 'new':
            current = None
        if current is None:
            action, expected = 'create', 0
        else:
            if current['storage'] != 'dynamic':
                raise ValueError(f'分类 {current["name"]} 是旧版固定分类，已下线；请新建动态分类导入')
            expected = current['version']
            action = 'reuse' if _field_signature(current['fields']) == _field_signature(template['fields']) else 'update'
        sections.append({'template': template, 'template_action': action,
                         'expected_version': expected, 'title': discovered['title'],
                         'header_row': discovered['header_row'], 'image_count': 0,
                         'source_sheet': discovered['source_sheet'],
                         'source_snapshot': '',
                         'drafts': {'new': [], 'update': [], 'delist': []}})
    return sections


def _attach_image_header_hints(sections: list[dict], discovered: list[dict], xlsx_path: str) -> None:
    """Expose image columns without header text for human template review."""
    try:
        evidence = {item['title']: item for item in workbook_templates.extract_header_evidence(xlsx_path)}
    except (OSError, ValueError):
        return
    by_sheet = {item['source_sheet']: item for item in discovered}
    for section in sections:
        source_sheet = section.get('source_discovered_sheet') or section.get('source_sheet')
        item = evidence.get(source_sheet)
        draft = by_sheet.get(source_sheet)
        if not item or not draft:
            continue
        header = max(1, int(draft.get('header_row') or 1))
        hints = []
        for raw_col, count in sorted(item.get('图片锚点列', {}).items()):
            col = int(raw_col)
            # A title or a printed header already names this column.
            if any(col <= len(row) and row[col - 1].strip()
                   for row in item['grid'][:header]):
                continue
            mapped = any(field.get('role') == 'image' and field.get('source_column') == col
                         for field in draft.get('fields', []))
            hints.append({'kind': 'unlabeled_image_column', 'column': col,
                          'image_count': count, 'mapped': mapped})
        if hints:
            section['review_hints'] = hints


def _merge_adjacent_blank_subcolumns(discovered: list[dict], xlsx_path: str) -> None:
    """Collapse an AI duplicate when its second column has no printed header."""
    if not any(left.get('label', '').casefold() == right.get('label', '').casefold()
               for item in discovered for left, right in zip(item.get('fields', []), item.get('fields', [])[1:])):
        return
    book = load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        for item in discovered:
            if item.get('source_sheet') not in book.sheetnames:
                continue
            ws = book[item['source_sheet']]
            fields = item.get('fields') or []
            index = 0
            while index + 1 < len(fields):
                left, right = fields[index:index + 2]
                if (left.get('label', '').casefold() != right.get('label', '').casefold()
                        or left.get('role') not in {'spec', 'note'}
                        or right.get('role') not in {'spec', 'note'}):
                    index += 1
                    continue
                hits = [(cell.row, cell.column) for row in ws.iter_rows(
                    min_row=1, max_row=min(ws.max_row, 30), max_col=min(ws.max_column, 200))
                    for cell in row if workbook_templates._label(cell.value).casefold()
                    == left['label'].casefold()]
                if len(hits) != 1:
                    index += 1
                    continue
                header, col = hits[0]
                if ((type(left.get('source_column')) is int and left['source_column'] != col)
                        or (type(right.get('source_column')) is int and right['source_column'] != col + 1)
                        or workbook_templates._text(ws.cell(header, col + 1).value)
                        or not any(workbook_templates._text(row[0].value) for row in ws.iter_rows(
                            min_row=header + 1, max_row=ws.max_row,
                            min_col=col + 1, max_col=col + 1))):
                    index += 1
                    continue
                left['source_column'] = col
                fields.pop(index + 1)
    finally:
        book.close()


def _recover_picture_models(discovered: list[dict], xlsx_path: str) -> None:
    """Use source cells to resolve an image-labeled column containing SKUs."""
    if not any(field.get('role') == 'image' and
               re.search(r'picture|image|图片|照片', field.get('label', ''), re.I)
               for item in discovered for field in item.get('fields', [])):
        return
    book = load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        for item in discovered:
            fields = item.get('fields') or []
            if item.get('source_sheet') not in book.sheetnames:
                continue
            ws = book[item['source_sheet']]
            printed = {}
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30),
                                    max_col=min(ws.max_column, 200)):
                for cell in row:
                    label = workbook_templates._label(cell.value).casefold()
                    if label:
                        printed.setdefault(label, []).append((cell.row, cell.column))
            header_rows = {}
            for field in fields:
                hits = printed.get(field.get('label', '').casefold(), [])
                if len(hits) == 1:
                    header_rows[id(field)], field['source_column'] = hits[0]
            candidates = [field for field in fields if field.get('role') == 'image'
                          and re.search(r'picture|image|图片|照片', field.get('label', ''), re.I)
                          and type(field.get('source_column')) is int]
            if len(candidates) != 1:
                continue
            field = candidates[0]
            first = max(1, int(item.get('header_row') or 1),
                        header_rows.get(id(field), 1)) + 1
            peer_cols = [other['source_column'] for other in fields if other is not field
                         and type(other.get('source_column')) is int]
            if not peer_cols:
                continue
            values = []
            sku_rows = []
            for row in ws.iter_rows(min_row=first, max_row=min(ws.max_row, first + 63)):
                value = workbook_templates._text(row[field['source_column'] - 1].value)
                if value and any(workbook_templates._text(row[col - 1].value) for col in peer_cols):
                    values.append(value)
                    if re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9._/+-]{0,39}', value):
                        sku_rows.append(row[0].row)
            # A picture formula or caption is not a product identity. Require
            # repeated row-aligned, short SKU tokens before changing the schema.
            sku_values = [value for value in values if re.fullmatch(
                r'[A-Za-z0-9][A-Za-z0-9._/+-]{0,39}', value)]
            if (len(set(sku_values)) < 2 or len(sku_values) * 3 < len(values) * 2):
                continue
            models = [other for other in fields if other is not field and other.get('role') == 'model']
            if len(models) > 1:
                continue
            if models:
                current = models[0]
                col = current.get('source_column')
                if type(col) is not int:
                    continue
                model_values = [workbook_templates._text(ws.cell(row, col).value)
                                for row in sku_rows]
                if sum(bool(re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9._/+-]{0,39}', value))
                       for value in model_values) * 2 >= len(sku_rows):
                    continue  # A plausible existing model must keep the AI's role.
                current['role'] = 'spec'
                current['key'] = workbook_templates._field_key(current['label'], 'spec',
                    {other['key'] for other in fields if other is not current})
            field.update(key=workbook_templates._field_key('型号', 'model',
                         {other['key'] for other in fields if other is not field}),
                         label='型号', role='model', type='text', searchable=True)
    finally:
        book.close()


def build_template_payload(conn, xlsx_path, work_dir, *, source_key: str,
                           doc_id: int | None = None, mode: str | None = None,
                           category_key: str | None = None,
                           prefetched_xls: dict | None = None) -> dict:
    """Create a small template-only approval payload.

    The workbook is opened once for schema discovery, but rows and images are
    intentionally skipped.  Product extraction happens only after this ticket
    is approved and the merchant uploads the workbook again.
    """
    if mode not in (None, 'new', 'existing'):
        raise ValueError('导入 mode 只能是 new 或 existing')
    if mode == 'existing' and not category_key:
        raise ValueError('并入已有分类时必须提供 category_key')
    if mode != 'existing' and category_key:
        raise ValueError('只有 existing 模式可以提供 category_key')
    workbook_templates.preflight_workbook(xlsx_path)
    # 表头发现三级链：qwen 快路（5-8 秒，证据=网格+图片锚点+合并区）→ 子代理
    # （约 4 分钟，稳）→ 代码按行猜（最后兜底）。多级表头/无标头图片列靠语义找齐。
    # Filename/sheet-name supplier inference does not depend on field discovery.
    # Run the two model calls together instead of adding their network latencies.
    book = load_workbook(xlsx_path, read_only=True)
    try:
        visible_sheets = [sheet.title for sheet in book.worksheets if sheet.sheet_state == 'visible']
    finally:
        book.close()
    converted_evidence = workbook_templates.extract_header_evidence(xlsx_path) if prefetched_xls else []
    use_prefetched = _safe_prefetched_xls(prefetched_xls, converted_evidence)
    with ThreadPoolExecutor(max_workers=1) as supplier_pool:
        if use_prefetched:
            supplier_future = supplier_pool.submit(lambda: prefetched_xls['supplier_guess'])
        else:
            supplier_future = supplier_pool.submit(
                ai_extract.guess_supplier, source_key or os.path.basename(xlsx_path), visible_sheets)
        discovered = prefetched_xls['discovered'] if use_prefetched else _qwen_template_sheets(xlsx_path)
        if discovered is not None:
            print(f'[dynamic_import] 模板阶段表头发现：qwen 快路（{len(discovered)} 个 Sheet）', flush=True)
        else:
            discovered = _agent_template_sheets(xlsx_path, work_dir)
            if discovered is not None:
                print(f'[dynamic_import] 模板阶段表头发现：子代理（{len(discovered)} 个 Sheet）', flush=True)
            else:
                discovered = workbook_templates.discover_workbook(
                    xlsx_path, None, include_rows=False, include_images=False)
                # 代码按行猜的结构不可靠，属性再由 qwen 修正。
                ai_extract.apply_field_attributes(discovered, ai_extract.infer_field_attributes(discovered))
                print('[dynamic_import] 模板阶段表头发现：代码 discover_workbook（兜底）', flush=True)
        _merge_adjacent_blank_subcolumns(discovered, xlsx_path)
        _recover_picture_models(discovered, xlsx_path)
        supplier_guess = supplier_future.result()
    sections = _template_sections(conn, discovered, source_key=source_key,
                                  doc_id=doc_id, mode=mode, category_key=category_key)
    _attach_image_header_hints(sections, discovered, xlsx_path)
    if not sections:
        raise ValueError('Excel 中没有识别到有效 Sheet 和表头')
    if supplier_guess:
        for section in sections:
            section['template'].setdefault('supplier', supplier_guess)
    return {'kind': 'template_import', 'phase': 'template', 'doc_id': doc_id,
            'filename': os.path.basename(xlsx_path),
            'source_key': source_key, 'mode': mode, 'category_key': category_key,
            'work_dir': str(work_dir), 'supplier_guess': supplier_guess, 'sheets': sections}


def _template_session(conn, template_doc_id: int) -> tuple[dict, list[dict]]:
    row = conn.execute('SELECT * FROM import_doc WHERE id=?', (template_doc_id,)).fetchone()
    if row is None or row['phase'] != 'template' or row['status'] != 'template_approved':
        raise ValueError('模板工单尚未审批通过，请先审批模板后再上传商品 Excel')
    try:
        mapping = json.loads(row['template_keys_json'] or '[]')
    except (TypeError, ValueError) as exc:
        raise ValueError('模板工单缺少有效的分类映射，请重新导入模板') from exc
    if not isinstance(mapping, list) or not mapping:
        raise ValueError('模板工单没有可用分类，请重新导入模板')
    return dict(row), mapping


def build_product_payload(conn, xlsx_path, work_dir, *, source_key: str,
                          template_doc_id: int, doc_id: int | None = None,
                          mode: str | None = None,
                          category_key: str | None = None) -> dict:
    """Parse a second upload against an already approved template session."""
    # 手工建分类直灌：无 template_doc_id 时按 category_key 用已审批模板
    # （商家页面对话建好分类字段后，Excel 只灌商品数据——复杂表的人工兜底）。
    if template_doc_id is None:
        if not category_key:
            raise ValueError('商品导入必须提供 templateDocId（模板工单）或 categoryKey（已有分类）')
        template = dynamic_catalog.get_template(conn, category_key)
        if template['storage'] != 'dynamic':
            raise ValueError('目标分类不是动态分类，不能直灌商品')
        return _products_into_template(conn, xlsx_path, work_dir, template,
                                       source_key=source_key, doc_id=doc_id)
    session, mapping = _template_session(conn, template_doc_id)
    # SQLite stores omitted mode as an empty string.  Treat that as the same
    # value as an omitted API argument so legacy callers can complete the
    # second phase without being rejected for a representation difference.
    if (mode or '') != (session.get('mode') or ''):
        raise ValueError('商品导入方式必须与模板导入一致，请使用模板工单返回的 mode')
    if (category_key or '') != (session.get('category_key') or ''):
        raise ValueError('商品导入目标分类必须与模板工单一致')
    discovered = workbook_templates.discover_workbook(xlsx_path, include_images=False)
    sections = []
    if mode == 'existing':
        # The existing-category contract accepts a workbook with several
        # sheets.  The template phase records one approved target schema;
        # the product phase maps every re-uploaded sheet into that schema so
        # no sheet silently disappears during the split.
        item = mapping[0]
        key = item.get('category_key') or category_key
        try:
            template = dynamic_catalog.get_template(conn, key)
        except KeyError as exc:
            raise ValueError('模板对应的分类已不存在，请重新导入模板') from exc
        expected = int(item.get('version') or 0)
        if template['version'] != expected:
            raise ValueError('分类模板已经更新，请重新导入并审批模板后再上传商品 Excel')
        if not discovered:
            raise ValueError('第二次上传没有识别到有效 Sheet，请上传同一份商品 Excel')
        incoming = _agent_rows(template, xlsx_path, work_dir)
        if incoming is None:
            raise ValueError('解析服务暂不可用，请稍后重试导入')
        source_rows = _source_rows(conn, template['key'], source_key, template['source_sheet'])
        existing = [row for row in source_rows if row['status'] != 'delisted']
        sections.append({'template': template, 'template_action': 'reuse',
                         'expected_version': template['version'],
                         'title': discovered[0]['title'],
                         'header_row': discovered[0]['header_row'],
                         'image_count': sum(item['image_count'] for item in discovered),
                         'source_sheet': template['source_sheet'],
                         'source_snapshot': _source_snapshot(source_rows),
                         'drafts': _classify_template_rows(existing, incoming, template)})
        return {'kind': 'template_import', 'phase': 'products', 'doc_id': doc_id,
                'filename': os.path.basename(xlsx_path),
                'template_doc_id': template_doc_id, 'source_key': source_key,
                'mode': mode, 'category_key': category_key, 'work_dir': str(work_dir),
                'sheets': sections}
    by_sheet = {item['source_sheet']: item for item in discovered}
    for item in mapping:
        source_sheet = item.get('source_sheet') or ''
        found = by_sheet.get(source_sheet)
        if found is None and mode == 'existing' and discovered:
            found = discovered[0]
        if found is None:
            raise ValueError(f'第二次上传缺少模板 Sheet“{source_sheet}”，请上传同一份商品 Excel')
        key = item.get('category_key') or category_key
        try:
            template = dynamic_catalog.get_template(conn, key)
        except KeyError as exc:
            raise ValueError('模板对应的分类已不存在，请重新导入模板') from exc
        expected = int(item.get('version') or 0)
        if template['version'] != expected:
            raise ValueError('分类模板已经更新，请重新导入并审批模板后再上传商品 Excel')
        incoming = _agent_rows(template, xlsx_path, work_dir, sheet=found.get('source_sheet') or '')
        if incoming is None:
            raise ValueError('解析服务暂不可用，请稍后重试导入')
        source_rows = _source_rows(conn, template['key'], source_key, template['source_sheet'])
        existing = [row for row in source_rows if row['status'] != 'delisted']
        sections.append({'template': template, 'template_action': 'reuse',
                         'expected_version': template['version'], 'title': found['title'],
                         'header_row': found['header_row'], 'image_count': found['image_count'],
                         'source_sheet': template['source_sheet'],
                         'source_snapshot': _source_snapshot(source_rows),
                         'drafts': _classify_template_rows(existing, incoming, template)})
    return {'kind': 'template_import', 'phase': 'products', 'doc_id': doc_id,
                'filename': os.path.basename(xlsx_path),
            'template_doc_id': template_doc_id, 'source_key': source_key,
            'mode': mode, 'category_key': category_key, 'work_dir': str(work_dir),
            'sheets': sections}


def build_ticket_payload(conn, xlsx_path, work_dir, *, source_key: str,
                         doc_id: int | None = None, mode: str | None = None,
                         category_key: str | None = None) -> dict:
    if mode not in (None, 'new', 'existing'):
        raise ValueError('导入 mode 只能是 new 或 existing')
    if mode == 'existing' and not category_key:
        raise ValueError('并入已有分类时必须提供 category_key')
    if mode != 'existing' and category_key:
        raise ValueError('只有 existing 模式可以提供 category_key')

    discovered_sheets = workbook_templates.discover_workbook(xlsx_path, work_dir)
    if mode == 'existing':
        try:
            target = dynamic_catalog.get_template(conn, category_key)
        except KeyError as exc:
            raise ValueError('未知目标分类，请先从商品管理中选择已有分类') from exc
        if target['storage'] != 'dynamic':
            raise ValueError(f'分类 {target["name"]} 是旧版固定分类，已下线；请新建动态分类导入')
        incoming = ParsedRows()
        for discovered in discovered_sheets:
            parsed = _agent_rows(target, xlsx_path, work_dir, sheet=discovered.get('source_sheet') or '')
            incoming.extend(parsed)
            incoming.failures.extend(parsed.failures)
            incoming.coverage.setdefault('sheets', []).append(parsed.coverage)
        source_rows = _source_rows(conn, target['key'], source_key, target['source_sheet'])
        existing = [row for row in source_rows if row['status'] != 'delisted']
        section = {'template': target, 'template_action': 'reuse',
                   'expected_version': target['version'], 'title': '',
                   'header_row': None, 'image_count': sum(s['image_count'] for s in discovered_sheets),
                   'source_sheet': target['source_sheet'],
                   'source_snapshot': _source_snapshot(source_rows),
                   'drafts': _classify_template_rows(existing, incoming, target)}
        if discovered_sheets:
            section['title'] = discovered_sheets[0]['title']
            section['header_row'] = discovered_sheets[0]['header_row']
        return {'kind': 'template_import', 'doc_id': doc_id, 'source_key': source_key,
                'mode': mode, 'category_key': category_key, 'work_dir': str(work_dir),
                'sheets': [section]}

    sheets = []
    for discovered in discovered_sheets:
        template = _clean_template(discovered)
        if mode == 'new':
            template = _unique_new_template(conn, template, source_key=source_key, doc_id=doc_id)
        try:
            current = dynamic_catalog.get_template(conn, template['key'])
        except KeyError:
            current = None
        if mode == 'new':
            current = None
        if current is None:
            action, expected, existing = 'create', 0, []
            source_rows = []
        else:
            if current['storage'] != 'dynamic':
                raise ValueError(f'分类 {current["name"]} 是旧版固定分类，已下线；请新建动态分类导入')
            expected = current['version']
            action = 'reuse' if _field_signature(current['fields']) == _field_signature(template['fields']) else 'update'
            source_rows = _source_rows(conn, current['key'], source_key, discovered['source_sheet'])
            existing = [row for row in source_rows if row['status'] != 'delisted']
        incoming = _agent_rows(template, xlsx_path, work_dir, sheet=discovered.get('source_sheet') or '')
        if incoming is None:
            raise ValueError('解析服务暂不可用，请稍后重试导入')
        sheets.append({'template': template, 'template_action': action,
                       'expected_version': expected, 'title': discovered['title'],
                       'header_row': discovered['header_row'], 'image_count': discovered['image_count'],
                       'source_sheet': discovered['source_sheet'],
                       'source_snapshot': _source_snapshot(source_rows),
                       'drafts': _classify_template_rows(existing, incoming, template)})
    if not sheets:
        raise ValueError('Excel 中没有识别到有效 Sheet 和表头')
    return {'kind': 'template_import', 'doc_id': doc_id, 'source_key': source_key,
            'mode': mode, 'category_key': category_key,
            'work_dir': str(work_dir), 'sheets': sheets}


def _edited(draft: dict, edits: dict, *aliases) -> dict:
    keys = [draft.get('_rid'), *aliases]
    change = next((edits.get(key) or edits.get(str(key)) for key in keys
                   if key is not None and (edits.get(key) or edits.get(str(key)))), {})
    if not isinstance(change, dict):
        return draft
    result = {**draft}
    data_changes = change.get('data') if isinstance(change.get('data'), dict) else {
        key: value for key, value in change.items() if not key.startswith('__')}
    result['data'] = {**draft.get('data', {}), **data_changes}
    if '__supplier' in change:
        result['supplier'] = str(change['__supplier'] or '').strip()[:40]
    if isinstance(change.get('__images'), list):
        result['images'] = change['__images']
        result['image_main'] = change['__images'][0] if change['__images'] else ''
    return result


def _apply_template_ticket(conn, payload: dict, decisions: dict | None = None) -> dict:
    decisions = decisions or {}
    template_edits = decisions.get('templates') or {}
    mapping = []
    for section in payload.get('sheets') or []:
        template_override = template_edits.get(section['template']['key'])
        template = {**section['template'], **(template_override or {})}
        try:
            current = None
            try:
                current = dynamic_catalog.get_template(conn, template['key'])
            except KeyError:
                pass
            if section['template_action'] in {'create', 'update'}:
                approved = dynamic_catalog.approve_template(
                    conn, template, expected_version=section['expected_version'])
            else:
                if current is None or current['version'] != section['expected_version']:
                    raise ValueError('此工单已过期（分类模板已更新），请直接驳回；如需入库请重新发送 Excel')
                changed = template_override and (
                    current['name'] != template['name']
                    or current.get('supplier', '') != template.get('supplier', '')
                    or current.get('source_sheet', '') != template.get('source_sheet', '')
                    or _field_signature(current['fields']) != _field_signature(template['fields']))
                approved = dynamic_catalog.approve_template(
                    conn, template, expected_version=section['expected_version']) if changed else current
        except ValueError as exc:
            from .tickets import TicketConflict
            raise TicketConflict(str(exc)) from exc
        mapping.append({'source_sheet': section.get('source_discovered_sheet') or section.get('source_sheet'),
                        'category_key': approved['key'], 'version': approved['version']})
    doc_id = payload.get('doc_id')
    if doc_id:
        conn.execute("UPDATE import_doc SET status='template_approved', template_keys_json=?, stats_json=? WHERE id=?",
                     (json.dumps(mapping, ensure_ascii=False),
                      json.dumps({'phase': 'template', 'template_doc_id': doc_id,
                                  'categories': [item['category_key'] for item in mapping],
                                  'next_action': 'upload_products'}, ensure_ascii=False), doc_id))
    return {'phase': 'template', 'template_approved': len(mapping),
            'template_doc_id': doc_id, 'template_keys': mapping,
            'next_action': 'upload_products'}


def _apply_product_payload(conn, payload: dict, decisions: dict | None = None) -> dict:
    return _apply_ticket_payload(conn, payload, decisions, product_only=True)


def _apply_ticket_payload(conn, payload: dict, decisions: dict | None = None,
                          *, product_only: bool = False) -> dict:
    decisions = decisions or {}
    rejected = {str(value) for value in decisions.get('reject', [])}
    edits = decisions.get('edits') or {}
    template_edits = decisions.get('templates') or {}
    created = updated = delisted = 0
    created_rows = []
    source = payload.get('source_key', '')
    for section in payload['sheets']:
        template = section['template']
        sheet_name = section.get('source_sheet') or template.get('source_sheet') or template['name']
        current_rows = []
        try:
            current_rows = _source_rows(conn, template['key'], source, sheet_name)
        except KeyError:
            pass
        if section.get('source_snapshot') != _source_snapshot(current_rows):
            from .tickets import TicketConflict
            raise TicketConflict('商品数据已变化，请重新导入并审批，避免覆盖最新修改')
    for section in payload['sheets']:
        template_override = {} if product_only else template_edits.get(section['template']['key'])
        template = {**section['template'], **(template_override or {})}
        try:
            current = None
            try:
                current = dynamic_catalog.get_template(conn, template['key'])
            except KeyError:
                pass
            if product_only:
                if current is None or current['version'] != section['expected_version']:
                    raise ValueError('分类模板已经更新，请重新导入模板并重新上传商品 Excel')
                approved = current
            elif (current is not None and not template_override
                    and current['version'] != section['expected_version']
                    and _field_signature(current['fields']) == _field_signature(template['fields'])):
                # 同名 Sheet 先后导入：模板已被同表头的其他工单创建/推进。
                # 表头一致即可安全按增量并入当前版本——商家要的“增量保留”，
                # 不再撞版本锁（真正表头变更的工单仍走下方过期报错）。
                approved = current
            elif section['template_action'] in {'create', 'update'}:
                approved = dynamic_catalog.approve_template(
                    conn, template, expected_version=section['expected_version'])
            else:
                approved = dynamic_catalog.get_template(conn, template['key'])
                if approved['version'] != section['expected_version']:
                    raise ValueError('此工单已过期（分类模板已更新），请直接驳回；如需入库请重新发送 Excel')
                changed = template_override and (
                    approved['name'] != template['name']
                    or approved.get('supplier', '') != template.get('supplier', '')
                    or approved.get('source_sheet', '') != template.get('source_sheet', '')
                    or _field_signature(approved['fields']) != _field_signature(template['fields']))
                if changed:
                    approved = dynamic_catalog.approve_template(
                        conn, template, expected_version=section['expected_version'])
        except ValueError as exc:
            from .tickets import TicketConflict
            raise TicketConflict(str(exc)) from exc
        allowed = {field['key'] for field in approved['fields']}
        sheet_name = section.get('source_sheet') or template.get('source_sheet') or template['name']
        for index, draft in enumerate(section['drafts'].get('new', [])):
            # 行钥匙与审批页/草稿编辑同源：_rid 优先，否则按段内序号 n{index}
            key = draft.get('_rid') or f'n{index}'
            if str(key) in rejected:
                continue
            draft = _edited(draft, edits, key)
            product_id = secrets.token_hex(8)
            row = {'id': product_id, 'inner_code': inner_code.gen(), 'cs_visible': 1,
                   'supplier': draft.get('supplier', approved.get('supplier', '')),
                   'data': {key: value for key, value in draft.get('data', {}).items() if key in allowed},
                   'images': draft.get('images') or [], 'source_row': draft.get('source_row'),
                   'row_fingerprint': draft.get('row_fingerprint', '')}
            outcome = dynamic_catalog.upsert_approved_products(
                conn, approved['key'], [row], source_key=source, source_sheet=sheet_name,
                source_doc=payload.get('doc_id'))
            created += outcome['created']
            created_rows.append({'id': product_id, '_category': approved['key'],
                                 '_table': 'product_dynamic', 'images': row['images'],
                                 'image_main': row['images'][0] if row['images'] else ''})
        for index, pair in enumerate(section['drafts'].get('update', [])):
            old, incoming = pair if isinstance(pair, list) else (pair, pair)
            keys = {str(old.get('id') or ''), str(incoming.get('_rid') or ''), f'u{index}'}
            if keys & rejected:
                continue
            incoming = _edited(incoming, edits, old.get('id'), f'u{index}')
            images = incoming['images'] if 'images' in incoming else (old.get('images') or [])
            row = {'id': old['id'], 'inner_code': old['inner_code'], 'cs_visible': old['cs_visible'],
                   'status': 'approved', 'supplier': incoming.get('supplier', old.get('supplier', '')),
                   'data': {key: value for key, value in incoming.get('data', {}).items() if key in allowed},
                   'images': images,
                   'source_row': incoming.get('source_row'),
                   'row_fingerprint': incoming.get('row_fingerprint', '')}
            if 'images' in incoming:
                conn.execute('DELETE FROM embedding WHERE product_id=?', (old['id'],))
            outcome = dynamic_catalog.upsert_approved_products(
                conn, approved['key'], [row], source_key=source, source_sheet=sheet_name,
                source_doc=payload.get('doc_id'))
            updated += outcome['updated']
            if images:
                created_rows.append({'id': old['id'], '_category': approved['key'],
                                     '_table': 'product_dynamic', 'images': images,
                                     'image_main': images[0]})
        # Missing Excel rows never authorize delisting, including older pending payloads.
    result = {'created': created, 'updated': updated, 'delisted': delisted,
              'created_rows': created_rows, 'work_dir': payload.get('work_dir')}
    if product_only:
        result.update({'phase': 'products', 'template_doc_id': payload.get('template_doc_id'),
                       'doc_id': payload.get('doc_id')})
        if payload.get('doc_id'):
            conn.execute("UPDATE import_doc SET status='product_approved', stats_json=? WHERE id=?",
                         (json.dumps({'phase': 'products', **{k: result[k] for k in ('created', 'updated', 'delisted')},
                                      'template_doc_id': payload.get('template_doc_id')}, ensure_ascii=False),
                          payload['doc_id']))
    return result


def apply_ticket(conn, payload: dict, decisions: dict | None = None) -> dict:
    phase = payload.get('phase')
    if phase == 'template':
        return _apply_template_ticket(conn, payload, decisions)
    if phase == 'products':
        return _apply_product_payload(conn, payload, decisions)
    return _apply_ticket_payload(conn, payload, decisions)


def _with_product_metadata(builder):
    from functools import wraps
    @wraps(builder)
    def wrapped(*args, **kwargs):
        return _product_metadata(builder(*args, **kwargs))
    return wrapped

build_product_payload = _with_product_metadata(build_product_payload)
build_ticket_payload = _with_product_metadata(build_ticket_payload)
