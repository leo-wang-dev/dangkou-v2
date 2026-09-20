#!/usr/bin/env python3
"""验收 Harness：对运行中的侧车跑 PRD §6 全场景，出桌面报告。
用法：python3 harness/run_all.py [BASE_URL]
前置：侧车已跑、.env 有真 key、源文件存在、本地 DB 路径可读（取审批 token）。"""
import json
import os
import sqlite3
import sys
import time
import urllib.request

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from catalog import config, quote
import openpyxl

BASE = sys.argv[1] if len(sys.argv) > 1 else 'http://127.0.0.1:8890'
TOKEN = config.SERVICE_TOKEN
DB = os.environ.get('CATALOG_V2_DB', os.path.join(
    os.path.dirname(__file__), '..', 'data', 'catalog.db'))
FILES = {
    'razor': os.environ.get('CATALOG_FILES_RAZOR', '/tmp/matrix/input/f2_tixudao.xlsx'),
    'curler': os.environ.get('CATALOG_FILES_CURLER', os.path.expanduser(
        '~/Library/Containers/com.tencent.xinWeChat/Data/Documents/'
        'xwechat_files/wxid_1qtn22qs4hnb22_b73a/msg/file/2026-08/华岳电器有限公司.xlsx')),
}
EXPECT = {cat: os.environ.get('CATALOG_EXPECT_' + cat.upper()) for cat in FILES}  # 按源表数据行，不按型号去重


def api(path, method='GET', body=None):
    req = urllib.request.Request(
        BASE + path, method=method,
        data=json.dumps(body).encode() if body else None,
        headers={'X-Service-Token': TOKEN, 'Content-Type': 'application/json'})
    with urllib.request.urlopen(req, timeout=180) as r:
        return json.loads(r.read())


results = []


def check(name, ok, detail=''):
    results.append((name, bool(ok), str(detail)))
    print(('✅' if ok else '❌'), name, detail, flush=True)


def wait_ticketed(doc_id, minutes=20):
    for _ in range(minutes * 6):
        time.sleep(10)
        s = api(f'/import/{doc_id}')
        if s['status'] != 'parsing':
            return s
    return s


def approve_latest_import(conn, doc_id):
    row = conn.execute(
        "SELECT id, token FROM approval_ticket WHERE status='pending' "
        "AND ticket_type='import' AND json_extract(payload,'$.doc_id')=? ORDER BY id DESC LIMIT 1", (doc_id,)).fetchone()
    api(f"/tickets/{row['id']}/decision", 'POST',
        {'token': row['token'], 'approved': True})
    return row['id']


def verify_products(cat):
    d = api(f'/products/{cat}')
    prods, t = d['products'], d['template']
    expected = int(EXPECT[cat])
    prods = [p for p in prods if p['状态'] != 'delisted']
    check(f'商品数[{cat}]', len(prods) == expected, f'{len(prods)}（源表黄金答案 {expected} 行）')
    key_label = [f['label'] for f in t['fields']
                 if f['col'] == ('model_no' if cat == 'razor' else 'item_no')][0]
    empty = sum(1 for p in prods if not str(p.get(key_label) or '').strip())
    check(f'空型号[{cat}]', empty == 0, f'{empty} 个')
    with_img = sum(1 for p in prods if p.get('主图'))
    rate = with_img * 100 // max(len(prods), 1)
    check(f'图片配对[{cat}]', rate >= 95, f'{rate}%')
    img_base = os.path.join(os.path.dirname(DB), 'images')
    phantom = sum(1 for p in prods if p.get('主图')
                  and not os.path.exists(os.path.join(img_base, str(p['主图']))))
    check(f'无虚构图[{cat}]', phantom == 0, f'{phantom} 个幽灵')
    return prods


def main():
    blocked = [f'缺源表 {cat}: {path}' for cat,path in FILES.items() if not os.path.isfile(path)]
    blocked += [f'缺 CATALOG_EXPECT_{cat.upper()} 源表黄金行数' for cat in FILES if not EXPECT[cat]]
    if not TOKEN:
        blocked.append('缺 CATALOG_V2_SERVICE_TOKEN')
    if not os.path.isfile(quote.TEMPLATE_V2_PATH):
        blocked.append('缺真实报价模板')
    if blocked:
        print('\n'.join('BLOCKED: ' + reason for reason in blocked))
        sys.exit(2)
    conn = sqlite3.connect(DB, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    for cat, path in FILES.items():
        if not os.path.exists(path):
            check(f'导入[{cat}]', False, f'源文件缺失 {path}')
            continue
        doc = api('/import', 'POST', {'path': path, 'category': cat, 'source_key': f'harness-{cat}'})['doc_id']
        s = wait_ticketed(doc)
        check(f'导入[{cat}]完成', s['status'] == 'ticketed',
              json.dumps(s.get('stats'), ensure_ascii=False))
        if s['status'] != 'ticketed':
            check(f'导入[{cat}]失败原因', False, str(s.get('error'))[:120])
            continue
        approve_latest_import(conn, doc)
        verify_products(cat)

    # 四分类重导（剃须刀：改一行报价 + 删一行）
    src = FILES['razor']
    if os.path.exists(src):
        import tempfile
        wb = openpyxl.load_workbook(src)
        ws = wb['Sheet1']
        for r in range(4, min(ws.max_row, 10) + 1):
            if ws.cell(r, 2).value:           # 有型号的行
                ws.cell(r, 11, '999.99')      # 改报价列（K）
                ws.delete_rows(r + 1)         # 删下一行
                break
        fd, tmp = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(tmp)
        doc = api('/import', 'POST', {'path': tmp, 'category': 'razor', 'source_key': 'harness-razor'})['doc_id']
        s = wait_ticketed(doc)
        st = s.get('stats') or {}
        check('重导四分类', s['status'] == 'ticketed'
              and st.get('update', 0) >= 1 and st.get('delist', 0) >= 1,
              json.dumps(st, ensure_ascii=False))
        approve_latest_import(conn, doc)

    # 报价单
    prods = [p for p in api('/products/curler')['products'] if p['状态'] != 'delisted']
    if prods:
        q = api('/quote', 'POST', {'items': [{'category':'curler', 'product_id':p['id'], 'quantity':40} for p in prods[:3]]})
        ok = os.path.exists(q['path'])
        if ok:
            wb2 = openpyxl.load_workbook(q['path'])
            ws = wb2.active
            ok = ws.max_column >= 14 and ws['A18'].value is not None and any(str(c.value).upper()=='TOTAL' for c in ws['A']) and len(ws._images) >= 1
        check('报价单', ok, q.get('path', ''))

    report = ['# dangkou-v2 验收报告', '',
              f'> {time.strftime("%Y-%m-%d %H:%M")} ｜ {BASE}', '',
              '| 场景 | 结果 | 详情 |', '|---|---|---|']
    report += [f'| {"✅" if ok else "❌"} {n} | {"pass" if ok else "FAIL"} | {d} |'
               for n, ok, d in results]
    fails = sum(1 for _, ok, _ in results if not ok)
    report += ['', f'**{len(results) - fails}/{len(results)} 通过**'
               + ('' if not fails else f'，{fails} 项待修')]
    out = os.path.expanduser('~/Desktop/dangkou-v2-验收报告.md')
    try:
        open(out, 'w').write('\n'.join(report) + '\n')
    except FileNotFoundError:
        os.makedirs(os.path.dirname(out), exist_ok=True)
        open(out, 'w').write('\n'.join(report) + '\n')
    print('报告:', out)
    sys.exit(1 if fails else 0)


if __name__ == '__main__':
    main()
