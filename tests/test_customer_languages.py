"""Customer languages exercise live HTTP, SQLite, translation and real XLSX."""
import io
import json
import sqlite3
import pytest
import openpyxl
from fastapi.testclient import TestClient
from catalog import cs_i18n, cs_export, userapp
from tests.test_h5_transactions import h5

CODES = 'zh en fr es pt ru ar de ja ko vi th id'.split()

def test_all_shop_language_choices_are_persisted(h5):
    app, database, _ = h5
    with TestClient(app) as client:
        visitor = client.post('/cs/chat/test-shop/session').json()['visitor']
        for code in CODES:
            r = client.post('/cs/chat/test-shop/lang', json={'visitor':visitor,'lang':code})
            assert r.json()['lang'] == code
            assert client.get('/cs/chat/test-shop/session',params={'visitor':visitor}).json()['lang'] == code
        assert client.post('/cs/chat/test-shop/lang',json={'visitor':visitor,'lang':'English'}).json()['lang'] == 'en'
        assert client.post('/cs/chat/test-shop/lang',json={'visitor':visitor,'lang':'中文'}).json()['lang'] == 'zh'

def test_fixed_errors_and_intent_need_no_provider(h5, monkeypatch):
    app, _, photo = h5
    with TestClient(app) as client:
        for code in CODES:
            headers={'X-Customer-Language':code}
            visitor=client.post('/cs/chat/test-shop/session',headers=headers).json()['visitor']
            reply=client.post('/cs/chat/test-shop/photo',headers=headers,data={'visitor':visitor},files={'file':('p.jpg',photo)}).json()['reply']
            assert reply == cs_i18n.t('photoIntentPrompt',code)
            error=client.post('/cs/chat/test-shop/message',headers=headers,json={'visitor':visitor,'text':''})
            assert error.status_code == 400
            assert error.json()['detail'] == cs_i18n.t('messageRequired',code)

def test_translate_english_prose_protects_repeated_literals_and_rejects_corruption():
    conn=sqlite3.connect(':memory:');conn.row_factory=sqlite3.Row
    source='Order KS-1100 from ACME Ltd: 66.95 USD, phone +86 13800138000 https://example.com/a. KS-1100'
    class Provider:
        calls=0
        def chat_text(self,prompt,messages,**kw):
            self.calls+=1
            rows=json.loads(messages[0]['content'])
            assert 'ACME Ltd' not in rows[0] and 'KS-1100' not in rows[0]
            return json.dumps([x.replace('Order','Commander').replace('from','chez').replace('phone','téléphone') for x in rows])
    provider=Provider()
    result=cs_i18n.translate_texts(conn,provider,'fr',[source],protected=['ACME Ltd'])[0]
    assert provider.calls == 1 and result.startswith('Commander')
    for value in ['ACME Ltd','KS-1100','66.95','+86 13800138000','https://example.com/a']:
        assert result.count(value) == source.count(value)
    class Broken:
        def chat_text(self,*a,**kw):return '["changed values"]'
    result=cs_i18n.translate_texts(conn,Broken(),'ar',[source],protected=['ACME Ltd'])[0]
    assert cs_i18n.t('translationUnavailable','ar') in result and source in result

def test_note_export_headers_and_arabic_direction_preserve_data():
    notes=[{'fields_json':json.dumps({'型号或品名':'KS-1100','单价':'66.95','供应商联系方式':'+86 13800138000'}),'photo':'','batch_fields':{'档口名称':'ACME Ltd','供应商联系方式':'+86 13800138000'}}]
    for code in CODES:
        ws=openpyxl.load_workbook(io.BytesIO(cs_export.render_notes(notes,lang=code))).active
        assert ws.cell(3,1).value == cs_i18n.t('serialHeader',code)
        headers=[cell.value for cell in ws[3]]
        assert cs_i18n.t('modelHeader',code) in headers
        assert ws.cell(4,headers.index(cs_i18n.t('modelHeader',code))+1).value == 'KS-1100'
        assert ws.cell(4,headers.index(cs_i18n.t('priceHeader',code))+1).value == '66.95'
        assert ws.cell(2,2).value == '+86 13800138000'
        assert bool(ws.sheet_view.rightToLeft) == (code=='ar')

def test_central_preference_and_export(tmp_path):
    app=userapp.build_app(str(tmp_path/'tool.db'),str(tmp_path/'photos'),str(tmp_path/'codes'))
    with TestClient(app) as client:
        guest=client.post('/guest').json()['guest']
        assert client.post('/lang',json={'guest':guest,'lang':'fr'}).json()['lang']=='fr'
        error=client.get('/export.xlsx',params={'guest':guest})
        assert error.status_code==409
        assert error.json()['detail']==cs_i18n.t('emptyExport','fr')
    app.state.conn.close()


def test_quotation_language_preserves_numeric_contract(tmp_path):
    from catalog import quote
    from tests.test_quote import _mkdb
    conn, storage = _mkdb(tmp_path,n=1,dims_only=False)
    for code in CODES:
        path=str(tmp_path/f'{code}.xlsx')
        details=[]
        quote.generate_generic(conn,storage,[{'category':'test_cat','product_id':'r0','quantity':50}],3,path,details=details,target_language=code)
        ws=openpyxl.load_workbook(path).active
        assert ws['A1'].value==cs_i18n.t('modelHeader',code)
        assert ws['A3'].value==cs_i18n.t('total',code)
        assert ws['A2'].value=='M0' and ws['E2'].value==10.30 and ws['F2'].value==80
        assert ws['G2'].value==824 and ws['G2'].data_type=='n'
        assert bool(ws.sheet_view.rightToLeft)==(code=='ar')
    conn.close()


def test_language_header_applies_to_first_chat_without_early_writer(h5,monkeypatch):
    from catalog.csbot import CsBot
    from catalog import llm
    import threading
    app,database,_=h5
    entered=threading.Event();release=threading.Event();calls=[]
    monkeypatch.setattr(CsBot,'_on_text',lambda *a,**kw:'Custom response for KS-1100')
    def translate(prompt,messages,**kw):
        if '⟦DKn⟧' not in prompt:return '{"actions":[]}'
        calls.append(1);entered.set();assert release.wait(5)
        return json.dumps([s.replace('Custom response for','Réponse pour') for s in json.loads(messages[0]['content'])])
    monkeypatch.setattr(llm,'chat_text',translate)
    with TestClient(app) as client:
        visitor=client.post('/cs/chat/test-shop/session').json()['visitor']
        result=[]
        thread=threading.Thread(target=lambda:result.append(client.post('/cs/chat/test-shop/message',headers={'X-Customer-Language':'fr'},json={'visitor':visitor,'text':'custom request'})))
        thread.start()
        try:
            assert entered.wait(2),'selected request language never reached dynamic translation'
            with sqlite3.connect(database,timeout=.2) as other:
                other.execute("UPDATE shop_profile SET owner_wechat='parallel-writer'")
        finally:release.set();thread.join(5)
        assert not thread.is_alive()
        assert result[0].json()['reply']=='Réponse pour KS-1100'
        assert len(calls)==1


def test_quote_api_explicit_target_and_download(h5,monkeypatch,tmp_path):
    from tests.test_quote import _mkdb
    from catalog import notify, llm
    app,_,_=h5
    source,_=_mkdb(tmp_path,n=1,dims_only=False)
    for table in ('category_template','product_dynamic'):
        rows=source.execute(f'SELECT * FROM {table}').fetchall()
        for row in rows:
            row=dict(row)
            if 'shop_id' in row: row['shop_id']=app.state.conn.execute('SELECT shop_id FROM shop_profile WHERE id=1').fetchone()[0]
            app.state.conn.execute(f'INSERT INTO {table}({",".join(row.keys())}) VALUES({",".join("?" for _ in row)})',tuple(row.values()))
    app.state.conn.commit();source.close()
    monkeypatch.setattr(notify,'push_file',lambda *a,**kw:None)
    monkeypatch.setattr(llm,'chat_text',lambda *a,**kw:'[]')
    with TestClient(app) as client:
        headers={'X-Service-Token':'test-token'}
        r=client.post('/quote',headers=headers,json={'items':[{'category':'test_cat','product_id':'r0','quantity':50}],'target_language':'ar'})
        assert r.status_code==200
        assert r.json()['target_language']=='ar'
        file=client.get(r.json()['download_url'],headers=headers)
        assert file.status_code==200
        ws=openpyxl.load_workbook(io.BytesIO(file.content)).active
        assert ws.sheet_view.rightToLeft and ws['A1'].value==cs_i18n.t('modelHeader','ar')
        assert ws['F2'].value==80
        assert client.get(r.json()['download_url']).status_code==401


def test_current_guest_language_merges_and_account_restores_on_another_device(tmp_path):
    app=userapp.build_app(str(tmp_path/'tool.db'),str(tmp_path/'photos'),str(tmp_path/'codes'))
    email='buyer@example.test'
    with TestClient(app) as client:
        guest=client.post('/guest',headers={'X-Customer-Language':'ar'}).json()['guest']
        app.state.conn.execute("INSERT INTO auth_codes(email,code_hash,expires_at) VALUES(?,?,datetime('now','+10 minutes'))",(email,userapp._hash(email+':123456')))
        app.state.conn.commit()
        login=client.post('/auth/verify',json={'guest':guest,'email':email,'code':'123456'}).json()
        me=client.get('/me',headers={'Authorization':'Bearer '+login['token'],'X-Customer-Language':'zh'})
        assert me.json()['lang']=='ar'
        assert client.get('/export.xlsx',params={'guest':guest}).status_code==410
    app.state.conn.close()


def test_fixed_chat_fallback_and_empty_confirmation_need_no_translator(h5,monkeypatch):
    from catalog import llm
    app,_,_=h5
    def model(prompt,*args,**kwargs):
        assert '客服消息翻译器' not in prompt,'fixed customer replies must not call the translator'
        return '{"actions":[]}'
    monkeypatch.setattr(llm,'chat_text',model)
    with TestClient(app) as client:
        guest=client.post('/cs/chat/test-shop/session').json()['visitor']
        for code in CODES[1:]:
            headers={'X-Customer-Language':code}
            r=client.post('/cs/chat/test-shop/message',headers=headers,json={'visitor':guest,'text':'Can you explain this unknown topic '+code})
            assert r.status_code==200
            assert r.json()['reply']==cs_i18n.t('unknownAnswer',code)
            r=client.post('/cs/chat/test-shop/message',headers=headers,json={'visitor':guest,'action':'confirm'})
            assert r.json()['reply']==cs_i18n.t('toolListEmpty',code)


def test_translation_cache_cannot_replay_changed_literal_or_duplicate_number():
    conn=sqlite3.connect(':memory:');conn.row_factory=sqlite3.Row
    cs_i18n.ensure_tables(conn)
    conn.execute('INSERT INTO cs_translation VALUES(?,?,?)',('fr','Order KS-1100 at 66.95','Commander KS-9999 à 99.00'));conn.commit()
    class Bad:
        def chat_text(self,prompt,messages,**kw):
            values=json.loads(messages[0]['content'])
            return json.dumps([value.replace('Order','Commander')+' 66.95' for value in values])
    value=cs_i18n.translate_texts(conn,Bad(),'fr',['Order KS-1100 at 66.95'])[0]
    assert value == cs_i18n.t('translationUnavailable','fr')+': Order KS-1100 at 66.95'


def test_receipt_localizes_missing_data_without_changing_prices_or_models():
    value=cs_i18n.receipt([{'型号或品名':'KS-1100','颜色':'未拍到','价格':'66.95（照片识别，待确认）'}],[],'ar')
    assert '未拍到' not in value and '照片识别' not in value
    assert 'KS-1100' in value and '66.95' in value


@pytest.mark.parametrize('cached',[False,True])
def test_translation_rejects_new_monetary_literal(cached):
    conn=sqlite3.connect(':memory:');conn.row_factory=sqlite3.Row
    source='Order KS-1100 at 66.95 USD'
    if cached:
        cs_i18n.ensure_tables(conn)
        conn.execute('INSERT INTO cs_translation VALUES(?,?,?)',('fr',source,'Commander KS-1100 at 66.95 USD 99.00 USD'));conn.commit()
    class Bad:
        def chat_text(self,prompt,messages,**kw):
            return json.dumps([x.replace('Order','Commander')+' 99.00 USD' for x in json.loads(messages[0]['content'])])
    assert cs_i18n.translate_texts(conn,Bad(),'fr',[source])[0]==cs_i18n.t('translationUnavailable','fr')+': '+source
    conn.close()


def test_unchanged_chinese_cache_retries_then_falls_back_honestly():
    conn=sqlite3.connect(':memory:');conn.row_factory=sqlite3.Row
    cs_i18n.ensure_tables(conn);conn.execute('INSERT INTO cs_translation VALUES(?,?,?)',('fr','红色外壳','红色外壳'));conn.commit()
    class Bad:
        calls=0
        def chat_text(self,*args,**kw):self.calls+=1;return '["红色外壳"]'
    model=Bad()
    assert cs_i18n.translate_texts(conn,model,'fr',['红色外壳'])[0]==cs_i18n.t('translationUnavailable','fr')+': 红色外壳'
    assert model.calls==1
    conn.close()


@pytest.mark.parametrize('kind',['central','shop'])
@pytest.mark.parametrize('failure',[False,True])
def test_photo_and_list_prose_projection_preserves_source_outside_writer(h5,tmp_path,monkeypatch,kind,failure):
    from catalog import llm
    fields={'型号或品名':'KS-1100','颜色':'红色','其他':'可折叠便于收纳'}
    app,path,photo=h5
    class Model:
        calls=0
        def chat_vision(self,*args,**kw):return json.dumps([fields])
        def chat_text(self,prompt,messages,**kw):
            self.calls+=1
            with sqlite3.connect(path,timeout=.2) as writer:writer.execute("INSERT INTO display_writes VALUES('during model')")
            if failure:raise RuntimeError('offline')
            return json.dumps([s.replace('红色','Rouge').replace('可折叠便于收纳','Pliable pour le rangement').replace('商品框缺失或无效，暂用整张照片','Photo complète') for s in json.loads(messages[0]['content'])])
    model=Model()
    if kind=='central':
        path=tmp_path/'central-display.db';app=userapp.build_app(str(path),str(tmp_path/'central-photos'),str(tmp_path/'codes'),llm=model)
    else:
        monkeypatch.setattr(llm,'chat_vision',model.chat_vision);monkeypatch.setattr(llm,'chat_text',model.chat_text)
    app.state.conn.execute('CREATE TABLE display_writes(value TEXT)');app.state.conn.commit()
    with TestClient(app) as client:
        headers={'X-Customer-Language':'fr'}
        if kind=='central':
            guest=client.post('/guest').json()['guest'];photo_route='/photo?guest='+guest;data={}
        else:
            guest=client.post('/cs/chat/test-shop/session').json()['visitor']
            client.post('/cs/chat/test-shop/mode',json={'visitor':guest,'mode':'notes'})
            photo_route='/cs/chat/test-shop/photo';data={'visitor':guest}
        response=client.post(photo_route,headers=headers,data=data,files={'file':('p.jpg',photo)})
        assert response.status_code==200,response.text
        expected=(cs_i18n.t('translationUnavailable','fr')+': 可折叠便于收纳') if failure else 'Pliable pour le rangement'
        assert expected in response.json()['reply']
        assert 'KS-1100' in response.json()['reply']
        if kind=='central':r=client.post('/notes?guest='+guest,headers=headers)
        else:
            link=client.get('/cs/chat/test-shop/list-token',params={'visitor':guest}).json()['token']
            r=client.get('/cs/link/'+link,headers=headers)
        assert r.status_code==200,r.text
        note=r.json()['notes'][0]
        assert note['fields']['其他']==fields['其他']
        assert note['display_fields']['其他']==expected
        assert note['display_fields']['型号或品名']=='KS-1100'
        assert note['display_labels']['其他']==cs_i18n.fixed('其他','fr')
        table='notes' if kind=='central' else 'cs_note'
        assert json.loads(app.state.conn.execute(f'SELECT fields_json FROM {table}').fetchone()[0])['其他']==fields['其他']
        assert app.state.conn.execute('SELECT count(*) FROM display_writes').fetchone()[0]>=2
    if kind=='central':app.state.conn.close()


@pytest.mark.parametrize('lang',['fr','ar'])
def test_multi_note_display_batches_finish_before_cache_writer(tmp_path,lang):
    conn=sqlite3.connect(tmp_path/'display.db');conn.row_factory=sqlite3.Row
    conn.execute('CREATE TABLE writer(value TEXT)');conn.commit()
    class Model:
        calls=0
        def chat_text(self,prompt,messages,**kw):
            self.calls+=1
            with sqlite3.connect(tmp_path/'display.db',timeout=.2) as writer:
                writer.execute('INSERT INTO writer VALUES(?)',(lang,))
            assert cs_i18n.LANGUAGES[lang] in prompt
            return json.dumps([x.replace('产品介绍','Description' if lang=='fr' else 'وصف').replace('红色','Rouge' if lang=='fr' else 'أحمر') for x in json.loads(messages[0]['content'])])
    model=Model();notes=[{'id':i,'fields':{'model_id':'KS-1100','产品介绍':f'红色 {i}','数量':i}} for i in range(65)]
    result=cs_i18n.project_notes(conn,model,lang,notes)
    assert model.calls==2 and len(result)==65
    for original,projected in zip(notes,result):
        assert projected['fields']==original['fields']
        assert projected['display_labels']['产品介绍']==('Description' if lang=='fr' else 'وصف')
        assert projected['display_fields']['产品介绍']==('Rouge' if lang=='fr' else 'أحمر')+f" {original['id']}"
        assert projected['display_fields']['model_id']=='KS-1100'
        assert projected['display_fields']['数量']==original['id']
    assert conn.execute('SELECT count(*) FROM writer').fetchone()[0]==2
    conn.rollback();conn.close()


def test_contact_handle_is_protected_and_new_handle_rejected():
    conn=sqlite3.connect(':memory:');conn.row_factory=sqlite3.Row
    source='Contact @buyer'
    class Model:
        def chat_text(self,prompt,messages,**kw):
            value=json.loads(messages[0]['content'])[0]
            assert '@buyer' not in value
            return json.dumps([value.replace('Contact','Contacter')+' @invented'])
    assert cs_i18n.translate_texts(conn,Model(),'fr',[source])[0]==cs_i18n.t('translationUnavailable','fr')+': '+source
    # Cache validation must reject an invented handle even when the source had none.
    cs_i18n.ensure_tables(conn)
    conn.execute('INSERT INTO cs_translation VALUES(?,?,?)',('fr','Hello','Bonjour @invented'));conn.commit()
    assert cs_i18n.translate_texts(conn,None,'fr',['Hello'],max_missing=0)[0]==cs_i18n.t('translationUnavailable','fr')+': Hello'
    conn.close()
