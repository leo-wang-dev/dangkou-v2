from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XImage
from PIL import Image


def blowdryer_fixture(tmp_path: Path) -> Path:
    path = tmp_path / '成本报价单.xlsx'
    photo = tmp_path / 'sample.png'
    Image.new('RGB', (120, 80), 'red').save(photo)
    wb = Workbook()
    ws = wb.active
    ws.title = '吹风机'
    ws.merge_cells('A1:G1')
    ws['A1'] = '吹风机报价单（中规欧规一个价）'
    headers = ['序号', '产品型号', '颜色', '图片', '不含税\n单风嘴成本', '不含税\n五风嘴成本',
               '品质款单嘴成本', '备注', '装箱数量', '体积', '重量']
    ws.append(headers)
    ws.append([1, '戴森款HD15', '大红色', None, 35, 44, 38, '普通盒', '单嘴20PCS', '57*37*42CM', '17KG/箱'])
    ws.append([2, '戴森款HD15', '马卡龙色', None, 37, 46, 40, '普通盒', '单嘴20PCS', '57*37*42CM', '17KG/箱'])
    ws.append([3, '戴森款HD16', '落日玫瑰色', None, 110, 135, None, '精品盒', '单嘴20PCS', '36*14*9CM', '19KG/箱'])
    ws.append([4, None, None, None, None, None, None, None, None, None, None])
    for col in ('C', 'D', 'H', 'I', 'J', 'K'):
        ws.merge_cells(f'{col}5:{col}6')
    ws.add_image(XImage(photo), 'D3')
    ws.add_image(XImage(photo), 'D4')
    ws.add_image(XImage(photo), 'D5')
    ws.add_image(XImage(photo), 'D5')
    wb.save(path)
    return path


def test_discovers_sheet_title_row2_header_fields_and_images(tmp_path):
    from catalog.workbook_templates import discover_workbook

    image_dir = tmp_path / 'images'
    draft = discover_workbook(blowdryer_fixture(tmp_path), image_dir)[0]
    assert draft['name'] == '吹风机'
    assert draft['title'] == '吹风机报价单（中规欧规一个价）'
    assert draft['header_row'] == 2
    assert [field['label'] for field in draft['fields']] == [
        '序号', '产品型号', '颜色', '图片', '不含税 单风嘴成本', '不含税 五风嘴成本',
        '品质款单嘴成本', '备注', '装箱数量', '体积', '重量']
    assert draft['image_count'] == 4
    assert len(draft['rows'][0]['images']) == 1
    assert len(draft['rows'][2]['images']) == 2
    assert not Path(draft['rows'][0]['images'][0]).is_absolute()
    assert (image_dir / draft['rows'][0]['images'][0]).is_file()


def test_each_numbered_row_is_a_product_and_merged_values_inherit(tmp_path):
    from catalog.workbook_templates import discover_workbook

    draft = discover_workbook(blowdryer_fixture(tmp_path), tmp_path / 'images')[0]
    assert len(draft['rows']) == 4
    model_key = next(f['key'] for f in draft['fields'] if f['role'] == 'model')
    color_key = next(f['key'] for f in draft['fields'] if f['label'] == '颜色')
    assert [row['data'][model_key] for row in draft['rows']] == [
        '戴森款HD15', '戴森款HD15', '戴森款HD16', '']
    assert draft['rows'][3]['data'][color_key] == '落日玫瑰色'
    assert len(draft['rows'][3]['images']) == 2
    assert draft['rows'][2]['row_fingerprint'] != draft['rows'][3]['row_fingerprint']


def test_costs_default_internal_and_model_is_optional(tmp_path):
    from catalog.workbook_templates import discover_workbook

    fields = discover_workbook(blowdryer_fixture(tmp_path))[0]['fields']
    costs = [field for field in fields if field['role'] == 'cost']
    assert len(costs) == 3
    assert all(field['visibility'] == 'internal' for field in costs)
    model = next(field for field in fields if field['role'] == 'model')
    assert model['required'] is False and model['visibility'] == 'public'
    sequence = next(field for field in fields if field['role'] == 'sequence')
    assert sequence['visibility'] == 'internal'


def test_price_cost_and_supplier_link_headers_default_internal(tmp_path):
    from openpyxl import Workbook
    from catalog.workbook_templates import discover_workbook

    workbook = Workbook(); sheet = workbook.active; sheet.title = '测试分类'
    sheet.append(['型号', '批发价', 'Cost', '商品链接', '颜色'])
    sheet.append(['SAFE-1', '23', '10', 'https://supplier.example/item', '红色'])
    path = tmp_path / 'sensitive.xlsx'; workbook.save(path)
    fields = {field['label']: field for field in discover_workbook(path)[0]['fields']}
    assert (fields['批发价']['role'], fields['批发价']['visibility']) == ('price', 'internal')
    assert (fields['Cost']['role'], fields['Cost']['visibility']) == ('cost', 'internal')
    assert fields['商品链接']['visibility'] == 'internal'


def test_each_visible_sheet_creates_an_independent_category(tmp_path):
    from catalog.workbook_templates import discover_workbook

    path = blowdryer_fixture(tmp_path)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb.create_sheet('配件')
    ws.append(['配件资料'])
    ws.append(['序号', '产品型号', '材质'])
    ws.append([1, 'PJ-1', 'ABS'])
    hidden = wb.create_sheet('内部说明')
    hidden.sheet_state = 'hidden'
    hidden.append(['不得导入'])
    wb.save(path)
    drafts = discover_workbook(path)
    assert [draft['name'] for draft in drafts] == ['吹风机', '配件']


def test_sheet_names_that_only_differ_by_spaces_do_not_share_a_category_key(tmp_path):
    from openpyxl import Workbook
    from catalog.workbook_templates import discover_workbook

    workbook = Workbook(); first = workbook.active; first.title = 'A B'
    first.append(['型号', '颜色']); first.append(['A', '红'])
    second = workbook.create_sheet('A  B')
    second.append(['型号', '颜色']); second.append(['B', '蓝'])
    path = tmp_path / 'two-sheets.xlsx'; workbook.save(path)
    drafts = discover_workbook(path)
    assert len({draft['key'] for draft in drafts}) == 2
    assert drafts[0]['key'] != drafts[1]['key']


def test_header_hints_keep_a_short_header_above_a_wider_data_row(tmp_path):
    from catalog.workbook_templates import discover_workbook

    workbook = Workbook(); sheet = workbook.active; sheet.title = '吹风机'
    sheet.append(['型号', '颜色', None])
    sheet.append(['HD15', '蓝色', '50'])
    path = tmp_path / 'short-header.xlsx'; workbook.save(path)

    draft = discover_workbook(path)[0]
    assert draft['header_row'] == 1
    assert [field['label'] for field in draft['fields']] == ['型号', '颜色']
    assert len(draft['rows']) == 1
    assert list(draft['rows'][0]['data'].values()) == ['HD15', '蓝色']
