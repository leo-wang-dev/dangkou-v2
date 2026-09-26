"""字段体系（动态分类版）：AI 结构注入 + 后端只校验不翻译。

分工：语义映射是 AI 的活（工具描述里给了字段清单）；
后端 _dynamic_mutation_ticket 只做机械过滤——清单外字段 400 打回，永不 500。
"""
import base64
import sqlite3

import pytest
from fastapi.testclient import TestClient

from catalog import db
from catalog.main import app
from catalog.storage import LocalStorage
from tests.conftest import DYNAMIC_FIELDS, seed_products


@pytest.fixture()
def client(tmp_path):
    conn = sqlite3.connect(':memory:', check_same_thread=False)
    conn.row_factory = sqlite3.Row
    db.init_db(conn)
    app.state.conn = conn
    app.state.token = 'test-service-token'
    app.state.storage = LocalStorage(str(tmp_path))
    app.state.callback = None
    return TestClient(app, headers={'X-Service-Token': 'test-service-token'})


def test_template_fields_exposed_by_products_api(client):
    """模板登记字段 → /products 的 template.fields 全量带出（H5 自动成一列）。"""
    seed_products(client.app.state.conn, [])
    r = client.get('/products/test_cat').json()
    labels = [f['label'] for f in r['template']['fields']]
    assert {'型号', '价格', '规格', '颜色', '电压', '箱规'} <= set(labels)


def test_create_mutate_all_illegal_returns_400_with_field_list(client):
    """B 层：全非法键 → 400 + 无法识别字段清单原文返回，AI 自纠重发。"""
    seed_products(client.app.state.conn, [])
    r = client.post('/products/test_cat', json={'changes': {'随便什么': 'x', '另一个': 'y'}})
    assert r.status_code == 400
    detail = r.json()['detail']
    assert '无法识别字段' in detail and '随便什么' in detail


def test_create_mutate_legal_keys_flow_to_product(client):
    """AI 按模板字段发键（字段名或 key 均可）→ 建单 200 → 审批落库。"""
    seed_products(client.app.state.conn, [])
    r = client.post('/products/test_cat', json={'changes': {
        'model': '8226', '价格': '21.5', '备注转写': '不存在的字段'}})
    assert r.status_code == 400   # 清单外键不允许（动态分类不拼备注，直接打回）
    r = client.post('/products/test_cat', json={'changes': {
        'model': '8226', '价格': '21.5'}})
    assert r.status_code == 200
    tid = r.json()['ticket_id']
    t = [x for x in client.get('/tickets').json()['tickets'] if x['id'] == tid][0]
    client.post(f"/tickets/{t['id']}/decision", json={'token': t['token'], 'approved': True})
    p = client.get('/products/test_cat').json()['products'][0]
    assert p['型号'] == '8226'
    assert p['价格'] == '21.5'


def test_update_changes_and_images_both_applied(tmp_path):
    """update 同带字段+图片，字段不再被 images 处理吞掉。"""
    png = base64.b64decode(
        'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==')
    conn = sqlite3.connect(':memory:', check_same_thread=False)
    conn.row_factory = sqlite3.Row
    db.init_db(conn)
    st = LocalStorage(str(tmp_path))
    rel = st.save('test_cat', 'p1', 'main.png', png)
    seed_products(conn, [{'id': 'p1', 'data': {'model': '8226', 'price': '21.5'},
                          'images': [rel]}])
    from catalog import dynamic_catalog, tickets
    current_row = next(row for row in dynamic_catalog.list_products(conn, 'test_cat')
                       if row['id'] == 'p1')
    tk = tickets.create(conn, 'mutate', 'test_cat',
                        {'kind': 'dynamic_mutate', 'action': 'update', 'product_id': 'p1',
                         'template_version': 1,
                         'changes': {'price': '23'}, 'images': [rel],
                         'before_snapshot': tickets._dynamic_product_snapshot(current_row)})
    row = conn.execute('SELECT token FROM approval_ticket WHERE status=?', ('pending',)).fetchone()
    tid = conn.execute('SELECT id FROM approval_ticket WHERE token=?', (row['token'],)).fetchone()['id']
    r = tickets.decide(conn, tid, row['token'], True)
    assert r['mutated'] == 'update'
    import json as _json
    data = _json.loads(conn.execute(
        "SELECT data_json FROM product_dynamic WHERE id='p1'").fetchone()[0])
    assert data['price'] == '23'


def test_delete_ticket_carries_before_snapshot(client, tmp_path):
    """下架工单带 before（型号/主图），审批页能核对批的是哪个货。"""
    seed_products(client.app.state.conn, [{'id': 'p1', 'data': {'model': '8225'}}])
    r = client.delete('/products/test_cat/p1')
    tid = r.json()['ticket_id']
    payload = client.get(f'/tickets/{tid}').json()['payload']
    assert payload['action'] == 'delete'
    assert payload['before']['型号'] == '8225'


def test_h5_delete_card_renders_before():
    """下架卡片必须用 before 渲染（图+型号），不再裸 ID。"""
    import os as _os
    s = open(_os.path.join(_os.path.dirname(__file__), '..', 'static', 'index.html'),
             encoding='utf-8').read()
    i = s.find("act === 'delete'")
    seg = s[i:i + 900]
    assert i >= 0 and 'p.before' in seg and '主图' in seg


def test_stats_full_mode(client):
    """full=true：全部在售全字段 + total；默认模式不带 products。"""
    seed_products(client.app.state.conn, [
        {'id': f'sr{i}', 'data': {'model': f'M{i}', 'spec': f'描述{i}', 'color': '黑',
                                  'ctn': 'QTY：40 PCS\nN.W.：12 KGS\nG.W.：13 KGS\nMEAS：45*40*40 CM',
                                  'price': str(10 + i)}}
        for i in range(3)])
    r = client.get('/stats').json()
    assert r['total'] == 3 and 'products' not in r
    r = client.get('/stats?full=true&category=test_cat').json()
    ps = r['products']['test_cat']
    assert len(ps) == 3 == r['by_category']['测试品类']
    assert all('型号' in p and 'id' in p and '规格' in p for p in ps)  # 全字段中文label


def test_full_catalog_quote_end_to_end(client):
    """整品类出单链路：stats full 拿 id → 全部拼 quote → 行数=款数。"""
    seed_products(client.app.state.conn, [
        {'id': f'sr{i}', 'data': {'model': f'M{i}', 'spec': f'描述{i}', 'color': '黑',
                                  'ctn': 'QTY：40 PCS\nN.W.：12 KGS\nG.W.：13 KGS\nMEAS：45*40*40 CM',
                                  'price': str(10 + i)}}
        for i in range(3)])
    r = client.get('/stats?full=true&category=test_cat').json()
    items = [{'category': 'test_cat', 'product_id': p['id'], 'quantity': 10}
             for p in r['products']['test_cat']]
    q = client.post('/quote', json={'items': items, 'price_adjustment_pct': 3,
                                    'deposit_pct': 20})
    assert q.status_code == 200
    import openpyxl
    ws = openpyxl.load_workbook(q.json()['path'].replace('\\', '/')).active
    n = len(items)
    assert ws.cell(18 + n - 1, 1).value                       # 最后一行有数据
    T = next(rr for rr in range(18 + n, 18 + n + 12)
             if ws.cell(rr, 1).value == 'TOTAL')
    assert ws.cell(T, 7).value == sum(round((10 + i) * 1.03) * 40 for i in range(n))  # 合计=值（10台→整箱40）
