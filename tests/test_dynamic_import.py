import sqlite3
import time

import pytest

from catalog import db, tickets
from catalog.storage import LocalStorage
from tests.test_workbook_templates import blowdryer_fixture


@pytest.fixture
def conn():
    value = sqlite3.connect(':memory:', check_same_thread=False)
    value.row_factory = sqlite3.Row
    db.init_db(value)
    yield value
    value.close()


def test_first_workbook_builds_one_template_import_section(conn, tmp_path):
    from catalog.dynamic_import import build_ticket_payload

    payload = build_ticket_payload(conn, blowdryer_fixture(tmp_path), tmp_path / 'work',
                                   source_key='vendor-a', doc_id=12)
    assert payload['kind'] == 'template_import'
    assert payload['doc_id'] == 12
    assert len(payload['sheets']) == 1
    sheet = payload['sheets'][0]
    assert sheet['template_action'] == 'create'
    assert sheet['template']['name'] == '吹风机'
    assert sheet['expected_version'] == 0
    assert len(sheet['drafts']['new']) == 4
    assert sheet['drafts']['update'] == [] and sheet['drafts']['delist'] == []


def test_one_approval_atomically_creates_template_and_all_rows(conn, tmp_path):
    from catalog import dynamic_catalog
    from catalog.dynamic_import import build_ticket_payload

    payload = build_ticket_payload(conn, blowdryer_fixture(tmp_path), tmp_path / 'work',
                                   source_key='vendor-a')
    ticket = tickets.create(conn, 'template_import', None, payload)
    result = tickets.decide(conn, ticket['id'], ticket['token'], True)
    template = dynamic_catalog.get_template(conn, payload['sheets'][0]['template']['key'])
    rows = dynamic_catalog.list_products(conn, template['key'])
    assert result['created'] == 4 and result['updated'] == 0
    assert template['name'] == '吹风机' and template['version'] == 1
    assert len(rows) == 4
    model_key = next(field['key'] for field in template['fields'] if field['role'] == 'model')
    assert [row['data'][model_key] for row in rows] == ['戴森款HD15', '戴森款HD15', '戴森款HD16', '']
    assert all(row['cs_visible'] == 1 for row in rows)


def test_reject_writes_neither_template_nor_products(conn, tmp_path):
    from catalog import dynamic_catalog
    from catalog.dynamic_import import build_ticket_payload

    payload = build_ticket_payload(conn, blowdryer_fixture(tmp_path), tmp_path / 'work', source_key='vendor-a')
    ticket = tickets.create(conn, 'template_import', None, payload)
    tickets.decide(conn, ticket['id'], ticket['token'], False)
    with pytest.raises(KeyError):
        dynamic_catalog.get_template(conn, payload['sheets'][0]['template']['key'])
    assert conn.execute('SELECT count(*) FROM product_dynamic').fetchone()[0] == 0


def test_stale_template_approval_cannot_overwrite_new_version(conn, tmp_path):
    from catalog import dynamic_catalog
    from catalog.dynamic_import import build_ticket_payload

    initial = build_ticket_payload(conn, blowdryer_fixture(tmp_path), tmp_path / 'first', source_key='vendor-a')
    first = tickets.create(conn, 'template_import', None, initial)
    tickets.decide(conn, first['id'], first['token'], True)
    pending = build_ticket_payload(conn, blowdryer_fixture(tmp_path), tmp_path / 'second', source_key='vendor-a')
    pending['sheets'][0]['template_action'] = 'update'
    second = tickets.create(conn, 'template_import', None, pending)
    template = dynamic_catalog.get_template(conn, pending['sheets'][0]['template']['key'])
    changed = {**template, 'fields': [*template['fields'], {
        'key': 'power', 'label': '功率', 'type': 'text', 'required': False,
        'visibility': 'public', 'searchable': True, 'role': 'spec'}]}
    dynamic_catalog.approve_template(conn, changed, expected_version=template['version'])
    conn.commit()
    with pytest.raises(tickets.TicketConflict, match='模板已经变化'):
        tickets.decide(conn, second['id'], second['token'], True)


def test_reimport_marks_changed_rows_and_suspected_delists(conn, tmp_path):
    from openpyxl import load_workbook
    from catalog import dynamic_catalog
    from catalog.dynamic_import import build_ticket_payload

    path = blowdryer_fixture(tmp_path)
    first = build_ticket_payload(conn, path, tmp_path / 'first', source_key='vendor-a')
    ticket = tickets.create(conn, 'template_import', None, first)
    tickets.decide(conn, ticket['id'], ticket['token'], True)
    category = first['sheets'][0]['template']['key']
    dynamic_catalog.upsert_approved_products(conn, category, [{
        'id': 'obsolete', 'inner_code': 'OLD', 'data': {'model': '旧款'},
        'images': [], 'source_row': 99, 'row_fingerprint': 'old-fingerprint',
    }], source_key='vendor-a', source_sheet='吹风机')
    conn.commit()
    wb = load_workbook(path)
    ws = wb['吹风机']
    ws['C3'] = '新红色'
    wb.save(path)
    second = build_ticket_payload(conn, path, tmp_path / 'second', source_key='vendor-a')
    drafts = second['sheets'][0]['drafts']
    assert len(drafts['update']) >= 1
    assert len(drafts['delist']) >= 1


def test_reimport_same_model_updates_specs_without_duplicate(conn, tmp_path):
    """同一来源、同一型号重导时更新规格并保留原商品 ID。"""
    from openpyxl import Workbook, load_workbook
    from catalog import dynamic_catalog
    from catalog.dynamic_import import build_ticket_payload

    path = tmp_path / 'same-model.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = '吹风机'
    ws.append(['序号', '产品型号', '颜色', '功率', '备注'])
    ws.append([1, '001', '红色', '1600W', '旧规格'])
    wb.save(path)

    first_payload = build_ticket_payload(conn, path, tmp_path / 'first', source_key='vendor-001')
    first = tickets.create(conn, 'template_import', None, first_payload)
    tickets.decide(conn, first['id'], first['token'], True)
    category = first_payload['sheets'][0]['template']['key']
    before = dynamic_catalog.list_products(conn, category)
    assert len(before) == 1
    product_id = before[0]['id']

    wb = load_workbook(path); ws = wb['吹风机']
    ws['C2'] = '蓝色'; ws['D2'] = '1800W'; ws['E2'] = '新规格'; wb.save(path)
    second_payload = build_ticket_payload(conn, path, tmp_path / 'second', source_key='vendor-001')
    assert len(second_payload['sheets'][0]['drafts']['new']) == 0
    assert len(second_payload['sheets'][0]['drafts']['update']) == 1
    second = tickets.create(conn, 'template_import', None, second_payload)
    result = tickets.decide(conn, second['id'], second['token'], True)
    after = dynamic_catalog.list_products(conn, category)
    assert result['created'] == 0 and result['updated'] == 1
    assert len(after) == 1 and after[0]['id'] == product_id
    fields = dynamic_catalog.get_template(conn, category)['fields']
    by_label = {field['label']: field['key'] for field in fields}
    assert after[0]['data'][by_label['颜色']] == '蓝色'
    assert after[0]['data'][by_label['功率']] == '1800W'
    assert after[0]['data'][by_label['备注']] == '新规格'


def test_pending_reimport_cannot_overwrite_a_newer_product_edit_or_delist(conn, tmp_path):
    from openpyxl import load_workbook
    from catalog import dynamic_catalog
    from catalog.dynamic_import import build_ticket_payload

    path = blowdryer_fixture(tmp_path)
    first_payload = build_ticket_payload(conn, path, tmp_path / 'first', source_key='vendor-a')
    first = tickets.create(conn, 'template_import', None, first_payload)
    tickets.decide(conn, first['id'], first['token'], True)
    workbook = load_workbook(path); workbook['吹风机']['C3'] = '新红色'; workbook.save(path)
    pending_payload = build_ticket_payload(conn, path, tmp_path / 'pending', source_key='vendor-a')
    pending = tickets.create(conn, 'template_import', None, pending_payload)
    category = first_payload['sheets'][0]['template']['key']
    product = dynamic_catalog.list_products(conn, category)[0]
    conn.execute("UPDATE product_dynamic SET status='delisted' WHERE id=?", (product['id'],)); conn.commit()

    with pytest.raises(tickets.TicketConflict, match='商品数据已变化'):
        tickets.decide(conn, pending['id'], pending['token'], True)
    assert conn.execute('SELECT status FROM product_dynamic WHERE id=?', (product['id'],)).fetchone()[0] == 'delisted'


def test_reordered_sheet_rows_keep_existing_product_ids(conn, tmp_path):
    from openpyxl import Workbook, load_workbook
    from catalog import dynamic_catalog
    from catalog.dynamic_import import build_ticket_payload

    path = tmp_path / 'reorder.xlsx'
    workbook = Workbook(); sheet = workbook.active; sheet.title = '配件'
    sheet.append(['序号', '产品型号', '颜色']); sheet.append([1, 'A-1', '红色']); sheet.append([2, 'B-1', '蓝色'])
    workbook.save(path)
    first_payload = build_ticket_payload(conn, path, tmp_path / 'first', source_key='vendor-reorder')
    first = tickets.create(conn, 'template_import', None, first_payload)
    tickets.decide(conn, first['id'], first['token'], True)
    category = first_payload['sheets'][0]['template']['key']
    template = dynamic_catalog.get_template(conn, category)
    model_key = next(field['key'] for field in template['fields'] if field['role'] == 'model')
    before = {row['data'][model_key]: row['id'] for row in dynamic_catalog.list_products(conn, category)}

    workbook = load_workbook(path); sheet = workbook['配件']
    a = [sheet.cell(2, col).value for col in range(1, 4)]
    b = [sheet.cell(3, col).value for col in range(1, 4)]
    for col, value in enumerate(b, 1): sheet.cell(2, col, value)
    for col, value in enumerate(a, 1): sheet.cell(3, col, value)
    workbook.save(path)
    second_payload = build_ticket_payload(conn, path, tmp_path / 'second', source_key='vendor-reorder')
    drafts = second_payload['sheets'][0]['drafts']
    assert drafts == {'new': [], 'update': [], 'delist': []}
    second = tickets.create(conn, 'template_import', None, second_payload)
    tickets.decide(conn, second['id'], second['token'], True)
    after = {row['data'][model_key]: row['id'] for row in dynamic_catalog.list_products(conn, category)}
    assert after == before


def test_async_import_without_category_creates_template_ticket(conn, tmp_path):
    from catalog import ingest

    path = blowdryer_fixture(tmp_path)
    doc_id = ingest.start(conn, LocalStorage(str(tmp_path / 'storage')), str(path), None,
                          source_key='vendor-a')
    for _ in range(100):
        status = ingest.status(conn, doc_id)
        if status['status'] != 'parsing':
            break
        time.sleep(.02)
    assert status['status'] == 'ticketed', status
    ticket = conn.execute('SELECT * FROM approval_ticket WHERE ticket_type=\'template_import\'').fetchone()
    assert ticket is not None and ticket['category'] is None
    assert status['stats']['categories'] == ['吹风机']


def test_category_api_and_blank_template_download(conn, tmp_path):
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from catalog.main import app
    from catalog.dynamic_import import build_ticket_payload

    payload = build_ticket_payload(conn, blowdryer_fixture(tmp_path), tmp_path / 'work', source_key='vendor-a')
    ticket = tickets.create(conn, 'template_import', None, payload)
    tickets.decide(conn, ticket['id'], ticket['token'], True)
    app.state.conn = conn
    app.state.token = 'service'
    app.state.storage = LocalStorage(str(tmp_path / 'storage'))
    with TestClient(app, headers={'X-Service-Token': 'service'}) as client:
        categories = client.get('/categories').json()['categories']
        hairdryer = next(value for value in categories if value['name'] == '吹风机')
        response = client.get(f"/categories/{hairdryer['key']}/template.xlsx")
    assert response.status_code == 200
    import io
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.title == '吹风机'
    assert [cell.value for cell in sheet[2]] == [
        '序号', '产品型号', '颜色', '图片', '不含税 单风嘴成本', '不含税 五风嘴成本',
        '品质款单嘴成本', '备注', '装箱数量', '体积', '重量']
    assert sheet['B3'].value == '示例型号'


def test_dynamic_ticket_preview_and_api_approval_persist_images(conn, tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    from catalog.main import app
    from catalog.dynamic_import import build_ticket_payload

    work = tmp_path / 'work'
    monkeypatch.setattr('catalog.search.embed_image', lambda data: [1.0, 0.0])
    payload = build_ticket_payload(conn, blowdryer_fixture(tmp_path), work, source_key='vendor-a')
    ticket = tickets.create(conn, 'template_import', None, payload)
    storage = LocalStorage(str(tmp_path / 'storage'))
    app.state.conn = conn
    app.state.token = 'service'
    app.state.storage = storage
    with TestClient(app, headers={'X-Service-Token': 'service'}) as client:
        detail = client.get(f"/tickets/{ticket['id']}").json()
        preview = detail['payload']['sheets'][0]['drafts']['new'][0]
        assert len(preview['_imgs']) == 1
        assert client.get(preview['_imgs'][0]).status_code == 200
        response = client.post(f"/tickets/{ticket['id']}/decision", json={
            'token': ticket['token'], 'approved': True})
    assert response.status_code == 200, response.text
    row = conn.execute('SELECT image_main,images_json FROM product_dynamic ORDER BY source_row LIMIT 1').fetchone()
    assert row['image_main'].startswith(payload['sheets'][0]['template']['key'] + '/')
    from pathlib import Path
    assert Path(storage.abs_path(row['image_main'])).is_file()


def test_dynamic_products_support_management_read_edit_and_delist(conn, tmp_path):
    from fastapi.testclient import TestClient
    from catalog.main import app
    from catalog.dynamic_import import build_ticket_payload

    payload = build_ticket_payload(conn, blowdryer_fixture(tmp_path), tmp_path / 'work', source_key='vendor-a')
    ticket = tickets.create(conn, 'template_import', None, payload)
    tickets.decide(conn, ticket['id'], ticket['token'], True)
    category = payload['sheets'][0]['template']['key']
    app.state.conn = conn; app.state.token = 'service'
    app.state.storage = LocalStorage(str(tmp_path / 'storage'))
    with TestClient(app, headers={'X-Service-Token': 'service'}) as client:
        listing = client.get(f'/products/{category}')
        assert listing.status_code == 200
        body = listing.json()
        assert body['template']['name'] == '吹风机' and len(body['products']) == 4
        model = next(field for field in body['template']['fields'] if field['role'] == 'model')
        product = body['products'][0]
        response = client.patch(f"/products/{category}/{product['id']}/direct", json={
            'changes': {model['col']: 'HD15-新版', 'cs_visible': '0'}})
        assert response.status_code == 200, response.text
        updated = client.get(f'/products/{category}').json()['products'][0]
        assert updated[model['label']] == 'HD15-新版' and updated['可观测'] == 0
        assert client.delete(f"/products/{category}/{product['id']}/direct").status_code == 200
        assert client.get(f'/products/{category}').json()['products'][0]['状态'] == 'delisted'


def test_stats_supports_sheet_defined_categories(conn, tmp_path):
    from fastapi.testclient import TestClient
    from catalog.main import app
    from catalog.dynamic_import import build_ticket_payload

    payload = build_ticket_payload(conn, blowdryer_fixture(tmp_path), tmp_path / 'work', source_key='vendor-a')
    ticket = tickets.create(conn, 'template_import', None, payload)
    tickets.decide(conn, ticket['id'], ticket['token'], True)
    category = payload['sheets'][0]['template']['key']
    app.state.conn = conn; app.state.token = 'service'
    app.state.storage = LocalStorage(str(tmp_path / 'storage'))
    with TestClient(app, headers={'X-Service-Token': 'service'}) as client:
        result = client.get(f'/stats?category={category}&full=true')
    assert result.status_code == 200, result.text
    body = result.json()
    assert body['total'] == 4 and body['by_category'] == {'吹风机': 4}
    assert len(body['products'][category]) == 4
    assert body['category_keys']['吹风机'] == category
    assert body['categories'] == [{'key': category, 'name': '吹风机',
                                   'total': 4, 'customer_visible': 4}]
    assert body['customer_visible_total'] == 4


def test_wechat_approval_endpoints_create_update_and_delist_dynamic_products(conn, tmp_path):
    from fastapi.testclient import TestClient
    from catalog.main import app
    from catalog.dynamic_import import build_ticket_payload

    payload = build_ticket_payload(conn, blowdryer_fixture(tmp_path), tmp_path / 'work', source_key='vendor-a')
    imported = tickets.create(conn, 'template_import', None, payload)
    tickets.decide(conn, imported['id'], imported['token'], True)
    category = payload['sheets'][0]['template']['key']
    app.state.conn = conn; app.state.token = 'service'; app.state.storage = LocalStorage(str(tmp_path / 'storage'))
    with TestClient(app, headers={'X-Service-Token': 'service'}) as client:
        template = client.get(f'/products/{category}').json()['template']
        model = next(field for field in template['fields'] if field['role'] == 'model')['col']
        color = next(field for field in template['fields'] if field['label'] == '颜色')['col']
        created = client.post(f'/products/{category}', json={
            'changes': {model: 'WECHAT-HD20', color: '黑色', 'cs_visible': '1'}})
        assert created.status_code == 200, created.text
        approved = client.post(f"/tickets/{created.json()['ticket_id']}/decision", json={
            'token': created.json()['token'], 'approved': True})
        assert approved.status_code == 200, approved.text
        product = next(row for row in client.get(f'/products/{category}').json()['products']
                       if row['产品型号'] == 'WECHAT-HD20')

        updated = client.patch(f"/products/{category}/{product['id']}", json={
            'changes': {color: '银色'}})
        assert updated.status_code == 200, updated.text
        assert client.post(f"/tickets/{updated.json()['ticket_id']}/decision", json={
            'token': updated.json()['token'], 'approved': True}).status_code == 200
        current = next(row for row in client.get(f'/products/{category}').json()['products']
                       if row['id'] == product['id'])
        assert current['颜色'] == '银色'

        deleted = client.delete(f"/products/{category}/{product['id']}")
        assert deleted.status_code == 200, deleted.text
        assert client.post(f"/tickets/{deleted.json()['ticket_id']}/decision", json={
            'token': deleted.json()['token'], 'approved': True}).status_code == 200
        current = next(row for row in client.get(f'/products/{category}').json()['products']
                       if row['id'] == product['id'])
        assert current['状态'] == 'delisted'


def test_dynamic_reimport_update_row_can_be_edited_before_approval(conn, tmp_path):
    from openpyxl import load_workbook
    from fastapi.testclient import TestClient
    from catalog.main import app
    from catalog.dynamic_import import build_ticket_payload

    path = blowdryer_fixture(tmp_path)
    first_payload = build_ticket_payload(conn, path, tmp_path / 'first', source_key='vendor-a')
    first = tickets.create(conn, 'template_import', None, first_payload)
    tickets.decide(conn, first['id'], first['token'], True)
    workbook = load_workbook(path); workbook['吹风机']['C3'] = '表格新红色'; workbook.save(path)
    payload = build_ticket_payload(conn, path, tmp_path / 'second', source_key='vendor-a')
    old, incoming = payload['sheets'][0]['drafts']['update'][0]
    color_key = next(field['key'] for field in payload['sheets'][0]['template']['fields'] if field['label'] == '颜色')
    ticket = tickets.create(conn, 'template_import', None, payload)
    app.state.conn = conn; app.state.token = 'service'; app.state.storage = LocalStorage(str(tmp_path / 'storage'))
    with TestClient(app, headers={'X-Service-Token': 'service'}) as client:
        saved = client.patch(f"/tickets/{ticket['id']}/draft", json={
            'token': ticket['token'], 'row_key': old['id'], 'edits': {color_key: '人工确认红色'}})
        assert saved.status_code == 200, saved.text
        approved = client.post(f"/tickets/{ticket['id']}/decision", json={
            'token': ticket['token'], 'approved': True})
        assert approved.status_code == 200, approved.text
        category = payload['sheets'][0]['template']['key']
        row = next(value for value in client.get(f'/products/{category}').json()['products']
                   if value['id'] == old['id'])
    assert row['颜色'] == '人工确认红色'
