import io
import json
import openpyxl
from tests.test_userapp import client, _new_guest, _upload, _last_code


def test_confirmed_cards_affect_only_subsequent_notes_and_export_groups(client):
    guest = _new_guest(client)
    _upload(client, guest)
    for name in ('A/Shop', 'A:Shop'):
        client.app.state.llm.vision_reply = json.dumps([{'名片':{'档口名称':name,'供应商联系人':name+' contact'}}])
        response = _upload(client, guest).json()
        assert response['pending_batches']
        batch_id = response['pending_batches'][0]['id']
        assert client.post('/batches/confirm', params={'guest':guest}, json={'batch_id':batch_id}).status_code == 200
        client.app.state.llm.vision_reply = json.dumps([{'型号或品名':name+' item'}])
        _upload(client, guest)
    notes = client.post('/notes', params={'guest':guest}).json()['notes']
    assert notes[0]['fields']['档口名称'] == '待补充'
    assert [n['fields']['档口名称'] for n in notes[1:]] == ['A/Shop','A:Shop']
    wb = openpyxl.load_workbook(io.BytesIO(client.get('/export.xlsx',params={'guest':guest}).content))
    assert len(wb.worksheets) == 3
    assert len(set(s.casefold() for s in wb.sheetnames)) == 3
    assert all(len(s)<=31 and not any(c in s for c in '\\/*?:[]') for s in wb.sheetnames)
    assert 'A/Shop contact' in str(list(wb.worksheets[1].values))
    assert 'A:Shop contact' not in str(list(wb.worksheets[1].values))


def test_card_with_product_needs_explicit_association_and_owner_scope(client):
    guest = _new_guest(client)
    client.app.state.llm.vision_reply = json.dumps([{'名片':{'档口名称':'Card'}},{'型号或品名':'Item'}])
    body = _upload(client, guest).json()
    bid = body['pending_batches'][0]['id']
    note = client.post('/notes', params={'guest':guest}).json()['notes'][0]
    assert note['fields']['档口名称'] == '待补充'
    other = _new_guest(client)
    assert client.post('/batches/confirm',params={'guest':other},json={'batch_id':bid}).status_code == 404
    assert client.post('/batches/confirm',params={'guest':guest},json={'batch_id':bid,'note_ids':[note['id']]}).status_code == 200
    assert client.post('/notes',params={'guest':guest}).json()['notes'][0]['fields']['档口名称'] == 'Card'


def test_merge_preserves_batch_cards_and_association(client):
    guest=_new_guest(client)
    client.app.state.llm.vision_reply=json.dumps([{'名片':{'档口名称':'Preserved'}}])
    bid=_upload(client,guest).json()['pending_batches'][0]['id']
    client.post('/batches/confirm',params={'guest':guest},json={'batch_id':bid})
    client.app.state.llm.vision_reply=json.dumps([{'型号或品名':'Goods'}])
    _upload(client,guest)
    client.post('/auth/code',json={'email':'card@example.com'})
    token=client.post('/auth/verify',json={'email':'card@example.com','code':_last_code(client,'card@example.com'),'guest':guest}).json()['token']
    auth={'Authorization':'Bearer '+token}
    data=client.post('/notes',headers=auth).json()
    assert data['notes'][0]['batch_id']==bid
    assert data['notes'][0]['fields']['档口名称']=='Preserved'
    assert data['batches'][0]['owner_kind']=='user'
    assert _upload(client,'',token=token).status_code==200
    assert client.post('/notes',headers=auth).json()['notes'][-1]['batch_id']==bid


def test_shop_card_and_product_do_not_silently_associate_and_manual_edit_detaches(h5, monkeypatch):
    from fastapi.testclient import TestClient
    from catalog import llm
    app,database,photo=h5
    monkeypatch.setattr(llm,'chat_vision',lambda *a,**kw:json.dumps([{'名片':{'档口名称':'A'}},{'型号或品名':'Item','档口名称':'A'}]))
    with TestClient(app) as c:
        guest=c.post('/cs/chat/test-shop/session').json()['visitor']
        c.post('/cs/chat/test-shop/mode',json={'visitor':guest,'mode':'notes'})
        c.post('/cs/chat/test-shop/photo',data={'visitor':guest},files={'file':('x.jpg',photo)})
        state=c.get('/cs/chat/test-shop/session',params={'visitor':guest}).json()
        n=state['notes'][0]
        assert n['fields']['档口名称']=='待补充'
        bid=next(b['id'] for b in state['batches'] if b['state']=='pending')
        c.post('/cs/chat/test-shop/batches/confirm',json={'visitor':guest,'batch_id':bid,'note_ids':[n['id']]})
        link=c.get('/cs/chat/test-shop/list-token',params={'visitor':guest}).json()['token']
        edited=c.patch('/cs/link/'+link+'/note/'+str(n['id']),json={'field':'档口名称','value':'B'})
        assert edited.json()['fields']['档口名称']=='B'
        assert app.state.conn.execute('SELECT fields_json FROM note_batches WHERE id=?',(bid,)).fetchone()[0]=='{"档口名称": "A"}'

from tests.test_h5_transactions import h5


def test_text_notes_share_explicit_current_batch(h5):
    from catalog import note_batches, purchase_notes
    from catalog.csbot import CsBot
    app,_,_=h5
    bot=CsBot(app.state.conn,None)
    cust=bot._ensure_customer({'id':'text-test'})
    bid=note_batches.create(app.state.conn,'guest',cust['id'],{'档口名称':'Text Supplier'},'pending')
    note_batches.confirm(app.state.conn,'guest',cust['id'],bid,[],'cs_note')
    note=purchase_notes.create(bot,cust,{'型号或品名':'Text Goods'})
    assert note['batch_id']==bid


def test_pending_card_blocks_later_central_plain_photos_until_reselection(client):
    guest=_new_guest(client)
    client.post('/batches',params={'guest':guest},json={'fields':{'档口名称':'A'}})
    _upload(client,guest)
    client.app.state.llm.vision_reply=json.dumps([{'名片':{'档口名称':'B'}}])
    _upload(client,guest)
    client.app.state.llm.vision_reply=json.dumps([{'型号或品名':'Unassigned goods'}])
    _upload(client,guest)
    data=client.post('/notes',params={'guest':guest}).json()
    assert data['notes'][0]['fields']['档口名称']=='A'
    assert data['notes'][1]['batch_state']=='unassigned'
    assert data['notes'][1]['fields']['档口名称']=='待补充'
    a=next(b['id'] for b in data['batches'] if b['fields'].get('档口名称')=='A')
    assert client.post('/batches/confirm',params={'guest':guest},json={'batch_id':a}).status_code==200
    _upload(client,guest)
    assert client.post('/notes',params={'guest':guest}).json()['notes'][-1]['fields']['档口名称']=='A'


def test_pending_card_blocks_shop_plain_photos_and_text_until_decline(h5,monkeypatch):
    from catalog import llm
    from fastapi.testclient import TestClient
    app,_,photo=h5
    with TestClient(app) as c:
        guest=c.post('/cs/chat/test-shop/session').json()['visitor']
        c.post('/cs/chat/test-shop/mode',json={'visitor':guest,'mode':'notes'})
        c.post('/cs/chat/test-shop/batches/confirm',json={'visitor':guest,'fields':{'档口名称':'A'}})
        monkeypatch.setattr(llm,'chat_vision',lambda *a,**kw:json.dumps([{'名片':{'档口名称':'B'}}]))
        c.post('/cs/chat/test-shop/photo',data={'visitor':guest},files={'file':('b.jpg',photo)})
        monkeypatch.setattr(llm,'chat_vision',lambda *a,**kw:json.dumps([{'型号或品名':'Plain goods'}]))
        c.post('/cs/chat/test-shop/photo',data={'visitor':guest},files={'file':('goods.jpg',photo)})
        def text_model(system,*a,**kw):
            if '采购记录抽取' in system:
                return json.dumps({'actions':[{'op':'create','fields':{'型号或品名':'Text goods','数量':'2个'}}]})
            return '{"action":"none"}'
        monkeypatch.setattr(llm,'chat_text',text_model)
        assert c.post('/cs/chat/test-shop/message',json={'visitor':guest,'text':'帮我记Text goods 2个'}).status_code==200
        data=c.get('/cs/chat/test-shop/session',params={'visitor':guest}).json()
        assert len(data['notes'])==2
        assert all(n['batch_state']=='unassigned' for n in data['notes'])
        b=next(b['id'] for b in data['batches'] if b['fields'].get('档口名称')=='B')
        assert c.post('/cs/chat/test-shop/batches/confirm',json={'visitor':guest,'batch_id':b,'action':'decline'}).json()=={'declined':True}
        c.post('/cs/chat/test-shop/photo',data={'visitor':guest},files={'file':('goods.jpg',photo)})
        assert c.get('/cs/chat/test-shop/session',params={'visitor':guest}).json()['notes'][-1]['fields']['档口名称']=='A'


def test_export_groups_identical_full_cards_but_not_same_name_conflicts(client):
    guest=_new_guest(client)
    for contact in ('wx-one','wx-one','wx-two'):
        fields={'档口名称':'Same','供应商联系人':'Owner','供应商联系方式':contact,'档口号/地址':'A1'}
        client.app.state.llm.vision_reply=json.dumps([{'名片':fields}])
        bid=_upload(client,guest).json()['pending_batches'][0]['id']
        client.post('/batches/confirm',params={'guest':guest},json={'batch_id':bid})
        client.app.state.llm.vision_reply=json.dumps([{'型号或品名':contact+' goods'}])
        _upload(client,guest)
    wb=openpyxl.load_workbook(io.BytesIO(client.get('/export.xlsx',params={'guest':guest}).content))
    assert len(wb.worksheets)==2
    assert wb.worksheets[0].max_row==7  # four card headers, columns, two goods
    assert 'wx-two' not in str(list(wb.worksheets[0].values))
    assert 'wx-two' in str(list(wb.worksheets[1].values))


def test_ambiguous_or_partial_identical_cards_do_not_merge():
    from catalog.cs_export import render_notes
    for contact in ('', '模糊（待确认）'):
        fields={'档口名称':'Same','供应商联系人':'Owner','供应商联系方式':contact,'档口号/地址':'A1'}
        notes=[{'batch_id':bid,'batch_fields':fields,'fields_json':json.dumps({'型号或品名':str(bid)}),'photo':''} for bid in (1,2)]
        assert len(openpyxl.load_workbook(io.BytesIO(render_notes(notes))).worksheets)==2


def test_supplier_identity_does_not_require_a_named_individual_contact():
    from catalog.cs_export import render_notes
    card={'档口名称':'Known Stall','档口号/地址':'A1','供应商联系方式':'unique-wx'}
    notes=[{'batch_id':bid,'batch_fields':card,'fields_json':json.dumps({'型号或品名':str(bid)}),'photo':''} for bid in (1,2)]
    assert len(openpyxl.load_workbook(io.BytesIO(render_notes(notes))).worksheets)==1
