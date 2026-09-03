"""报价单 v2.2：ELETRO BELEZA 模板全字段 + 动态行数 + 动态定金。

验收口径（用户拍板）：
- 模板=商家真文件（ELETRO BELEZA），不自造；
- 14 列逐属性断言，物流数据来自箱规解析（剃须刀）或结构化字段（卷发棒），没有的留空；
- 行数=商品数（>6插行、<6删空行），收款码/合并单元格/行高跟着搬；
- 合计/DEPOSIT/BALANCE 公式按实际行数重写，定金比例动态（默认30%）。
"""
import base64
import math
import sqlite3

import openpyxl
import pytest

from catalog import db, quote
from catalog.storage import LocalStorage

PNG = base64.b64decode(
    'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==')

# ---- 生产库 24 种真实箱规变体（2026-09-03 从 catalog.db 全量拉取）----
REAL_CTN = [
    ('QTY：40 PCS\nN.W.：16.9 KGS\nG.W.：17.7 KGS\nMEAS：38.5*37.5*42.5 CM', 40, 16.9, 17.7, (38.5, 37.5, 42.5)),
    ('QTY：60 PCS\nN.W.：18 KGS\nG.W.：19 KGS\nMEAS：47*45.5*42.5 CM', 60, 18, 19, (47, 45.5, 42.5)),
    ('QTY:   80 PCS\nN.W.:  14.3  KGS       \nG.W.:  15.1  KGS         \nMEAS:39.5X28X44.5CM', 80, 14.3, 15.1, (39.5, 28, 44.5)),
    ('QTY: 40PCS\nN.W.: 14.8KGS\nG.W.:15.5KGS\nMEAS:37.5*29*40.5CM', 40, 14.8, 15.5, (37.5, 29, 40.5)),
    ('QTY：40 PCS\nN.W.：13.7 KGS\nG.W.：14.5 KGS\nMEAS：38.5*37.5*42.5 CM', 40, 13.7, 14.5, (38.5, 37.5, 42.5)),
    ('QTY：40 PCS\nN.W.：12.4 KGS\nG.W.：13.3 KGS\nMEAS：43*41.5*45  CM', 40, 12.4, 13.3, (43, 41.5, 45)),
    ('QTY：60 PCS\nN.W.：11.9 KGS\nG.W.：12.5 KGS\nMEAS： 42X32X33  CM', 60, 11.9, 12.5, (42, 32, 33)),
    ('QTY：60 PCS\nN.W.：16.2 KGS\nG.W.：17.0 KGS\nMEAS：41*40.5*42cm', 60, 16.2, 17.0, (41, 40.5, 42)),
    ('QTY: 40 PCS\nN.W.:12.2KGS\nG.W.:13.2KGS\nMEAS:48.5X44.5X25.5 CM', 40, 12.2, 13.2, (48.5, 44.5, 25.5)),
    ('QTY: 60 PCS\nN.W.:16.2KGS\nG.W.:17.3KGS\nMEAS:48.5X44.5X37 CM', 60, 16.2, 17.3, (48.5, 44.5, 37)),
    ('QTY：20 PCS\nN.W.：11 KGS\nG.W.：11.7 KGS\nMEAS：37*30*45.5 CM', 20, 11, 11.7, (37, 30, 45.5)),
    ('QTY：60 PCS\nN.W.：17.0 KGS\nG.W.：17.9 KGS\nMEAS：45*39*41.5 CM\n（单机）', 60, 17.0, 17.9, (45, 39, 41.5)),
    ('QTY：40 PCS\nN.W.：11.9 KGS\nG.W.：12..9 KGS\nMEAS：45.5*31*44.5 CM', 40, 11.9, 12.9, (45.5, 31, 44.5)),  # 双点手误
    ('QTY：60 PCS\nN.W.：15 KGS\nG.W.：16.8 KGS\nMEAS：44*39*39 CM', 60, 15, 16.8, (44, 39, 39)),
    ('QTY：40 PCS\nN.W.：24.5 KGS\nG.W.：25.7 KGS\nMEAS：68*43.5*50 CM', 40, 24.5, 25.7, (68, 43.5, 50)),
    ('QTY：24 PCS\nN.W.：13.6 KGS\nG.W.：14.5 KGS\nMEAS：42*38*49 CM\n（套装）', 24, 13.6, 14.5, (42, 38, 49)),
    ('QTY：40 PCS\nN.W.：21.0 KGS\nG.W.：22.2 KGS\nMEAS：66*44.5*44 CM', 40, 21.0, 22.2, (66, 44.5, 44)),
    ('QTY:   80 PCS\nN.W.:  14.4  KGS       \nG.W.:  15.5 KGS         \nMEAS:37.5*32*38CM', 80, 14.4, 15.5, (37.5, 32, 38)),
    ('QTY：20 PCS\nN.W.：10.9 KGS\nG.W.：11.6 KGS\nMEAS：42*31*43.5 CM', 20, 10.9, 11.6, (42, 31, 43.5)),
    ('QTY：60 PCS\nN.W.：14.3 KGS\nG.W.：15.2 KGS\nMEAS：47.5*40.5*40.5CM', 60, 14.3, 15.2, (47.5, 40.5, 40.5)),
    ('QTY：40 PCS\nN.W.：12.7 KGS\nG.W.：13.5 KGS\nMEAS：48.5*32.5*33cm', 40, 12.7, 13.5, (48.5, 32.5, 33)),
    ('QTY：40 PCS\nN.W.：18.1 KGS\nG.W.：19.1 KGS\nMEAS：60*33.5*54 CM', 40, 18.1, 19.1, (60, 33.5, 54)),
    ('QTY：60 PCS\nN.W.：15.3 KGS\nG.W.：16.3 KGS\nMEAS：42*33*42 CM', 60, 15.3, 16.3, (42, 33, 42)),
    ('QTY：12 PCS\nN.W.：15.4 KGS\nG.W.：16.5 KGS\nMEAS：46*44*35 CM', 12, 15.4, 16.5, (46, 44, 35)),
]


@pytest.mark.parametrize('text,pcs,nw,gw,dims', REAL_CTN)
def test_parse_all_real_razor_specs(text, pcs, nw, gw, dims):
    out = quote.parse_ctn_spec(text)
    assert out['pcs'] == pcs
    assert out['nw'] == nw
    assert out['gw'] == gw
    assert out['dims'] == dims
    assert out['meas'] == '*'.join(str(d) for d in dims)


def test_parse_garbage_returns_empty():
    assert quote.parse_ctn_spec('') == {}
    assert quote.parse_ctn_spec('随便写的') == {}
    assert quote.parse_ctn_spec(None) == {}


@pytest.mark.parametrize('text,dims', [
    ('60*40*40', (60, 40, 40)),
    ('50.5*35.5*42.5', (50.5, 35.5, 42.5)),
    ('52*42*46 CM', (52, 42, 46)),
    ('39.5X28X44.5CM', (39.5, 28, 44.5)),
    ('40', None),
    ('', None),
])
def test_parse_curler_dims(text, dims):
    assert quote.parse_dims(text) == dims


# ---------- E2E：逐属性断言 ----------

def _mkdb(tmp_path, n_razor=3, curler=True):
    conn = sqlite3.connect(':memory:', check_same_thread=False)
    conn.row_factory = sqlite3.Row
    db.init_db(conn)
    st = LocalStorage(str(tmp_path))
    for i in range(n_razor):
        rel = st.save('razor', f'r{i}', 'm.png', PNG)
        conn.execute(
            "INSERT INTO product_razor(id, inner_code, model_no, description, color, "
            "ctn_spec, price, image_main) VALUES(?,?,?,?,?,?,?,?)",
            (f'r{i}', f'KS-R{i}', f'M{i}', f'描述{i}', '黑色',
             REAL_CTN[i][0], str(10 + i * 5), rel))
    if curler:
        rel = st.save('curler', 'c0', 'm.png', PNG)
        conn.execute(
            "INSERT INTO product_curler(id, inner_code, item_no, ctn_size, ctn_qty, "
            "price, image_main) VALUES(?,?,?,?,?,?,?)",
            ('c0', 'KS-C0', '8226', '60*40*40', '40', '21.5', rel))
    conn.commit()
    return conn, st


def test_full_fields_three_items(tmp_path):
    """3款（删空行路径）：14列逐格断言 + 公式 + 图片 + 结构搬移。"""
    conn, st = _mkdb(tmp_path, n_razor=2)
    out = str(tmp_path / 'q3.xlsx')
    quote.generate_v2(conn, st, [
        {'category': 'razor', 'product_id': 'r0', 'quantity': 100},   # 40/箱 → 3箱
        {'category': 'razor', 'product_id': 'r1', 'quantity': 60},    # 60/箱 → 1箱
        {'category': 'curler', 'product_id': 'c0', 'quantity': 80},   # 40/箱 → 2箱（无重量）
    ], 3, out, deposit_pct=20)
    ws = openpyxl.load_workbook(out).active
    # 表头不动
    assert ws.cell(17, 1).value == 'ITEM NO.'
    assert ws.cell(17, 14).value == 'T-CBM'
    # 行1：razor r0（40/箱 毛重17.7 净重16.9 38.5*37.5*42.5，+3% 单价 round(10*1.03)=10）
    r = 18
    assert ws.cell(r, 1).value == 'M0'
    assert ws.cell(r, 3).value == '描述0'
    assert ws.cell(r, 4).value == '黑色'
    assert ws.cell(r, 5).value == 10                       # E 价格
    assert ws.cell(r, 6).value == 100                      # F 数量
    assert ws.cell(r, 7).value == '=ROUND(E18*F18,2)'      # G 小计公式
    assert ws.cell(r, 8).value == 40                       # H 每箱
    assert ws.cell(r, 9).value == 3                        # I 箱数 ⌈100/40⌉
    assert ws.cell(r, 10).value == 17.7                    # J 毛重
    assert ws.cell(r, 11).value == 16.9                    # K 净重
    assert ws.cell(r, 12).value == '38.5*37.5*42.5'        # L MEAS
    assert ws.cell(r, 13).value == round(3 * 17.7, 2)      # M 总毛重 53.1
    assert ws.cell(r, 14).value == round(3 * (38.5 * 37.5 * 42.5) / 1e6, 3)  # N 总体积
    # 行2：razor r1（60/箱 → 1箱）
    r = 19
    assert ws.cell(r, 9).value == 1
    assert ws.cell(r, 13).value == round(1 * 19.0, 2)
    # 行3：curler（缺重量留空、描述颜色留空）
    r = 20
    assert ws.cell(r, 3).value in (None, '')               # C 描述空
    assert ws.cell(r, 4).value in (None, '')               # D 颜色空
    assert ws.cell(r, 5).value == 22                       # round(21.5*1.03)
    assert ws.cell(r, 8).value == 40                       # 装箱数
    assert ws.cell(r, 9).value == 2                        # ⌈80/40⌉
    assert ws.cell(r, 10).value is None and ws.cell(r, 11).value is None  # 无重量
    assert ws.cell(r, 12).value == '60*40*40'
    assert ws.cell(r, 13).value is None                    # M 空
    assert ws.cell(r, 14).value == round(2 * (60 * 40 * 40) / 1e6, 3)
    # 合计：3款 → TOTAL 行在 18+3+3=24（删了3个空行，labels 23）
    assert ws.cell(24, 1).value == 'TOTAL'
    assert ws.cell(24, 7).value == '=SUM(G18:G20)'
    assert ws.cell(24, 9).value == '=SUM(I18:I20)'
    assert ws.cell(24, 13).value == '=SUM(M18:M20)'
    assert ws.cell(24, 14).value == '=SUM(N18:N20)'
    # 定金 20%：总额 = 10*100 + round(15*1.03)=15 → 15*60 + 22*80
    total = 10 * 100 + 15 * 60 + 22 * 80
    assert ws.cell(25, 1).value and 'DEPOSIT' in str(ws.cell(25, 1).value)
    assert ws.cell(25, 7).value == f'=ROUND(G24*0.2,2)'
    assert ws.cell(26, 7).value == '=G24-G25'
    # 图片：模板5张（LOGO/门市×2/收款码×2）+ 3张商品图；收款码上移3行（48→45，锚0基44）
    imgs = ws._images
    assert len(imgs) == 8
    qr_rows = sorted(a := [im.anchor._from.row + 1 for im in imgs if im.anchor._from.row > 30])
    assert qr_rows == [45, 45]
    # 备注行合并搬移：数据到20行 → 备注在21
    assert any(m.min_row == 21 and m.max_row == 21 for m in ws.merged_cells.ranges)


def test_twenty_items_insert_path(tmp_path):
    """20款（插行路径）：行数对、合计范围对、收款码下移14行。"""
    conn, st = _mkdb(tmp_path, n_razor=20, curler=False)
    items = [{'category': 'razor', 'product_id': f'r{i}', 'quantity': 10 + i}
             for i in range(20)]
    out = str(tmp_path / 'q20.xlsx')
    quote.generate_v2(conn, st, items, 0, out)
    ws = openpyxl.load_workbook(out).active
    assert ws.cell(18, 1).value == 'M0' and ws.cell(37, 1).value == 'M19'
    T = 41                                                   # 18+20 → remark38 spacer39 labels40 TOTAL41
    assert ws.cell(T, 1).value == 'TOTAL'
    assert ws.cell(T, 7).value == '=SUM(G18:G37)'
    assert ws.cell(T + 2, 7).value == '=G41-G42'             # BALANCE
    qr = [im.anchor._from.row + 1 for im in ws._images if im.anchor._from.row > 40]
    assert qr == [62, 62]                                    # 48+14
    assert len(ws._images) == 25                             # 5模板 + 20商品


def test_deposit_default_and_edges(tmp_path):
    conn, st = _mkdb(tmp_path, n_razor=1, curler=False)
    out = str(tmp_path / 'qd.xlsx')
    quote.generate_v2(conn, st, [{'category': 'razor', 'product_id': 'r0', 'quantity': 100}], 0, out)
    ws = openpyxl.load_workbook(out).active
    T = next(r for r in range(19, 30) if ws.cell(r, 1).value == 'TOTAL')
    assert ws.cell(T + 1, 7).value == '=ROUND(G%d*0.3,2)' % T     # 默认30%
    # 定金100%（尾款0）与0%（定金0）不炸
    for pct in (0, 100):
        quote.generate_v2(conn, st, [{'category': 'razor', 'product_id': 'r0', 'quantity': 7}],
                          0, str(tmp_path / f'q{pct}.xlsx'), deposit_pct=pct)
    ws2 = openpyxl.load_workbook(str(tmp_path / 'q100.xlsx')).active
    T2 = next(r for r in range(19, 30) if ws2.cell(r, 1).value == 'TOTAL')
    assert ws2.cell(T2 + 1, 7).value == f'=ROUND(G{T2}*1.0,2)'


def test_ceil_edges():
    assert quote._ceil_div(80, 40) == 2      # 整除
    assert quote._ceil_div(81, 40) == 3      # 差1进位
    assert quote._ceil_div(1, 40) == 1
    assert quote._ceil_div(100, 40) == 3
