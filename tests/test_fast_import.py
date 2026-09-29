from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from catalog.fast_import import parse_structured
from catalog.fast_import import parse_grouped


def test_picture_column_with_skus_is_model_and_uses_structured_parser(tmp_path, monkeypatch):
    from catalog import dynamic_import

    photo = tmp_path / 'photo.png'
    Image.new('RGB', (12, 12), 'red').save(photo)
    book = Workbook()
    sheet = book.active
    sheet.title = 'SDQ'
    sheet.append(['PICTURE', 'DESCRIPTION', 'PACKING DETAILS', None, 'PRICE'])
    sheet.append(['SDQ-101', 'Hair dryer', 'Product size', '20cm', 20])
    sheet.append([None, '1200W', 'QTY:', '40pcs'])
    sheet.append(['Photo is injection color'])
    sheet.append(['SDQ-102', 'Hair dryer', 'Product size', '25cm', 25])
    sheet.append([None, '1500W', 'QTY:', '50pcs'])
    sheet.append(['Remark'])
    sheet.append(['1. Payment term: deposit'])
    sheet.add_image(ExcelImage(photo), 'A2')
    sheet.add_image(ExcelImage(photo), 'A5')
    path = tmp_path / 'sdq.xlsx'
    book.save(path)
    monkeypatch.setattr(dynamic_import, '_qwen_template_sheets', lambda path: [{
        'key': 'sdq', 'name': 'SDQ', 'source_sheet': 'SDQ', 'title': 'SDQ',
        'header_row': 1, 'fields': [
            {'key': 'image', 'label': 'PICTURE', 'type': 'image', 'role': 'image',
             'visibility': 'public', 'searchable': False, 'source_column': 1},
            {'key': 'description', 'label': 'DESCRIPTION', 'type': 'text', 'role': 'spec',
             'visibility': 'public', 'searchable': False, 'source_column': 2},
            {'key': 'packing', 'label': 'PACKING DETAILS', 'type': 'text', 'role': 'spec',
             'visibility': 'public', 'searchable': False, 'source_column': 3},
            {'key': 'price', 'label': 'PRICE', 'type': 'money', 'role': 'price',
             'visibility': 'internal', 'searchable': False, 'source_column': 5}],
    }])
    monkeypatch.setattr(dynamic_import.ai_extract, 'guess_supplier', lambda *args: '')
    import sqlite3
    from catalog import db
    conn = sqlite3.connect(':memory:')
    conn.row_factory = sqlite3.Row
    db.init_db(conn)
    try:
        payload = dynamic_import.build_template_payload(
            conn, path, tmp_path / 'work', source_key='sdq')
        fields = payload['sheets'][0]['template']['fields']
        assert fields[0]['role'] == 'model'
        assert fields[0]['type'] == 'text'
        assert fields[0]['label'] == '型号'
        from catalog import dynamic_catalog
        approved = dynamic_catalog.approve_template(conn, payload['sheets'][0]['template'],
                                                    expected_version=0)
        assert approved['fields'][0]['label'] == '型号'
        monkeypatch.setattr('catalog.agent.parse_dynamic', lambda *a, **kw:
                            (_ for _ in ()).throw(AssertionError('slow agent called')))
        rows = dynamic_import._agent_rows(approved, str(path), tmp_path / 'products', sheet='SDQ')
        assert [row['data'][fields[0]['key']] for row in rows] == ['SDQ-101', 'SDQ-102']
        assert all(len(row['images']) == 1 for row in rows)
        assert 'QTY: 40pcs' in rows[0]['data']['packing']
        assert any('Photo is injection color' in failure['reason'] for failure in rows.failures)
        assert any(8 in failure['source_rows'] for failure in rows.failures)
    finally:
        conn.close()


def test_picture_formula_references_are_not_promoted_to_models(tmp_path):
    from catalog.dynamic_import import _recover_picture_models

    book = Workbook()
    sheet = book.active
    sheet.title = 'Products'
    sheet.append(['PICTURE', 'PRICE'])
    sheet.append(['#NAME?', 20])
    sheet.append(['=DISPIMG("id1")', 25])
    path = tmp_path / 'formulas.xlsx'
    book.save(path)
    discovered = [{'source_sheet': 'Products', 'header_row': 1, 'fields': [
        {'key': 'image', 'label': 'PICTURE', 'role': 'image', 'type': 'image',
         'source_column': 1}]}]
    _recover_picture_models(discovered, path)
    assert discovered[0]['fields'][0]['role'] == 'image'


def test_exact_header_groups_continuations_and_images(tmp_path):
    photo = tmp_path / 'photo.png'
    Image.new('RGB', (12, 12), 'red').save(photo)
    book = Workbook()
    sheet = book.active
    sheet.title = 'Sheet1'
    sheet.append(['产品型号', '产品名称', '产品图片', '产品价格', '装箱数量'])
    sheet.append(['A1', '一号', None, '82元', '50台/箱'])
    sheet.append([None, None, None, None, '30台/箱'])
    sheet.append([None, '二号', '#NAME?', '90元', '20台/箱'])
    sheet.add_image(ExcelImage(photo), 'C2')
    sheet.add_image(ExcelImage(photo), 'C3')
    path = tmp_path / 'products.xlsx'
    book.save(path)
    fields = [
        {'key': 'model', 'label': '产品型号', 'role': 'model'},
        {'key': 'name', 'label': '产品名称', 'role': 'spec'},
        {'key': 'image', 'label': '产品图片', 'role': 'image'},
        {'key': 'price', 'label': '产品价格', 'role': 'price'},
        {'key': 'stock', 'label': '装箱数量', 'role': 'stock'},
    ]
    result = parse_structured({'fields': fields}, path, tmp_path / 'out')
    assert [item['model'] for item in result['products']] == ['A1', 'A1']
    assert result['products'][0]['stock'] == '50台/箱\n30台/箱'
    assert result['products'][0]['source_rows'] == [2, 3]
    assert len(result['products'][0]['images']) == 2
    assert result['products'][1]['source_rows'] == [4]
    assert result['failures'][0]['source_rows'] == [4]


def test_mismatched_header_falls_back(tmp_path):
    book = Workbook()
    book.active.append(['型号', '品名'])
    book.active.append(['A1', '一号'])
    path = tmp_path / 'different.xlsx'
    book.save(path)
    assert parse_structured({'fields': [
        {'key': 'model', 'label': '产品型号', 'role': 'model'},
        {'key': 'name', 'label': '产品名称', 'role': 'spec'},
    ]}, path, tmp_path / 'out') is None


def test_product_import_uses_fast_path_and_keeps_coverage(tmp_path, monkeypatch):
    from catalog.dynamic_import import _agent_rows

    book = Workbook()
    sheet = book.active
    sheet.title = 'Sheet1'
    sheet.append(['产品型号', '产品名称', '备注'])
    sheet.append(['A1', '一号', '首行'])
    sheet.append([None, None, '补充'])
    sheet.append(['A2', '二号', '末行'])
    path = tmp_path / 'products.xlsx'
    book.save(path)
    monkeypatch.setattr('catalog.agent.parse_dynamic', lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError('model called')))
    rows = _agent_rows({'key': 'test', 'fields': [
        {'key': 'model', 'label': '产品型号', 'role': 'model'},
        {'key': 'name', 'label': '产品名称', 'role': 'spec'},
        {'key': 'note', 'label': '备注', 'role': 'note'},
    ]}, str(path), tmp_path / 'work', sheet='Sheet1')
    assert len(rows) == 2
    assert rows[0]['data']['note'] == '首行\n补充'
    assert rows.coverage['uncertain'] is False
    assert rows.coverage['logical_products'] == 2


def test_model_per_row_with_unlabelled_last_column_is_reviewable(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Sheet1'
    sheet.append(['型号', '图片', '价格', '产品规格', '起订量', None])
    sheet.append(['KZ-228', None, '28.5元/台', '1300w', 1000, '24台装'])
    sheet.append(['KZ-3366', None, '35元/台', '2000w', 1000, '30台装'])
    book.create_sheet('Sheet2')
    path = tmp_path / 'kesen.xlsx'
    book.save(path)
    result = parse_structured({'fields': [
        {'key': 'model', 'label': '型号', 'role': 'model'},
        {'key': 'image', 'label': '图片', 'role': 'image'},
        {'key': 'price', 'label': '价格', 'role': 'price'},
        {'key': 'spec', 'label': '产品规格', 'role': 'spec'},
        {'key': 'stock', 'label': '起订量', 'role': 'stock'},
        {'key': 'note', 'label': '备注', 'role': 'spec'},
    ]}, path, tmp_path / 'out', sheet='Sheet1')
    assert [item['model'] for item in result['products']] == ['KZ-228', 'KZ-3366']
    assert result['products'][0]['note'] == '24台装'
    assert result['failures'][0]['source_rows'] == [1]


def test_model_only_sheet_with_unidentified_data_row_uses_agent(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.append(['型号', '价格'])
    sheet.append(['A1', '10元'])
    sheet.append([None, '20元'])
    path = tmp_path / 'ambiguous.xlsx'
    book.save(path)
    assert parse_structured({'fields': [
        {'key': 'model', 'label': '型号', 'role': 'model'},
        {'key': 'price', 'label': '价格', 'role': 'price'},
    ]}, path, tmp_path / 'out') is None


def test_grouped_repeated_models_keep_images_and_flag_conflicts(tmp_path):
    photo = tmp_path / 'photo.png'
    Image.new('RGB', (12, 12), 'red').save(photo)
    book = Workbook()
    sheet = book.active
    sheet.title = 'Sheet1'
    sheet['A1'] = 'ITEM.NO 型号'
    sheet['G2'] = '装箱尺寸'
    sheet['K2'] = '价格'
    for row, model, price in [(3, 'A', 21), (4, 'A', 21), (5, 'B', 22), (6, 'B', 18.52)]:
        sheet.cell(row, 1, model)
        sheet.cell(row, 7, '57*36*46')
        sheet.cell(row, 11, price)
        sheet.add_image(ExcelImage(photo), f'B{row}')
    sheet['C3'] = '电直梳'
    sheet['E3'] = '7两'
    path = tmp_path / 'grouped.xlsx'
    book.save(path)
    template = {'fields': [
        {'key': 'model', 'label': 'ITEM.NO 型号', 'role': 'model'},
        {'key': 'image', 'label': '图片', 'role': 'image'},
        {'key': 'kind', 'label': '电吹风', 'role': 'spec'},
        {'key': 'carton', 'label': '装箱尺寸', 'role': 'spec'},
        {'key': 'price', 'label': '价格', 'role': 'price'},
    ]}
    columns = {'model': 1, 'image': 2, 'kind': 3, 'carton': 7, 'price': 11}
    result = parse_grouped(template, path, tmp_path / 'grouped-out',
                           columns=columns, header_row=2, sheet='Sheet1')
    assert [p['model'] for p in result['products']] == ['A']
    assert result['products'][0]['source_rows'] == [3, 4]
    assert result['products'][0]['kind'] == '电直梳'
    assert len(result['products'][0]['images']) == 2
    assert any('E' in f['reason'] and f['source_rows'] == [3]
               for f in result['failures'])
    assert any('价格' in f['reason'] and f['source_rows'] == [5, 6]
               for f in result['failures'])


def test_product_import_uses_reviewable_grouped_path_before_slow_agent(tmp_path, monkeypatch):
    from catalog.dynamic_import import _agent_rows

    book = Workbook()
    sheet = book.active
    sheet.title = 'Sheet1'
    sheet['A1'] = '型号'
    sheet['K2'] = '价格'
    for row, model in [(3, 'A'), (4, 'A'), (5, 'B')]:
        sheet.cell(row, 1, model)
        sheet.cell(row, 11, 21)
    path = tmp_path / 'grouped.xlsx'
    book.save(path)
    monkeypatch.setattr('catalog.dynamic_import.ai_extract.map_approved_fields',
                        lambda path, fields, sheet: {'header_row': 2,
                                                     'columns': {'model': 1, 'price': 11}})
    monkeypatch.setattr('catalog.agent.parse_dynamic', lambda *a, **kw: (_ for _ in ()).throw(AssertionError('slow agent called')))
    rows = _agent_rows({'key': 'test', 'fields': [
        {'key': 'model', 'label': '型号', 'role': 'model'},
        {'key': 'price', 'label': '价格', 'role': 'price'},
    ]}, str(path), tmp_path / 'work', sheet='Sheet1')
    assert [r['data']['model'] for r in rows] == ['A', 'B']
    assert rows.coverage['uncertain'] is False


def test_grouped_sheet_with_unavailable_column_map_uses_semantic_agent(tmp_path, monkeypatch):
    from catalog.dynamic_import import _agent_rows
    photo = tmp_path / 'photo.png'
    Image.new('RGB', (12, 12), 'red').save(photo)
    book = Workbook(); ws = book.active; ws.title = 'Sheet1'
    ws['A1'] = '型号'; ws['K2'] = '价格'
    for row in (3, 4):
        ws.cell(row, 1, 'A'); ws.cell(row, 11, 21)
        ws.add_image(ExcelImage(photo), f'B{row}')
    path = tmp_path / 'grouped.xlsx'; book.save(path)
    monkeypatch.setattr('catalog.dynamic_import.ai_extract.map_approved_fields',
                        lambda *a: None)
    calls = []
    def semantic_parse(*args, **kwargs):
        calls.append(kwargs)
        return {'products': [{'source_sheet': 'Sheet1', 'source_rows': [3, 4],
                              'model': 'A', 'price': '21', 'images': [], 'image_count': 0}],
                'failures': [], 'vendor': None}
    monkeypatch.setattr('catalog.agent.parse_dynamic', semantic_parse)
    rows = _agent_rows({'key': 'test', 'fields': [
        {'key': 'model', 'label': '型号', 'role': 'model'},
        {'key': 'image', 'label': '图片', 'role': 'image'},
        {'key': 'price', 'label': '价格', 'role': 'price'},
    ]}, str(path), tmp_path / 'work', sheet='Sheet1')
    assert [r['data']['model'] for r in rows] == ['A']
    assert len(calls) == 1


def test_approved_column_mapping_recovers_blank_image_header(tmp_path, monkeypatch):
    from catalog import ai_extract
    photo = tmp_path / 'photo.png'
    Image.new('RGB', (12, 12), 'red').save(photo)
    book = Workbook()
    ws = book.active
    ws.title = 'Sheet1'
    ws['A1'] = '型号'
    ws['K2'] = '价格'
    ws['A3'] = 'A'
    ws['C3'] = '电直梳'
    ws['K3'] = 21
    ws.add_image(ExcelImage(photo), 'B3')
    path = tmp_path / 'blank-header.xlsx'
    book.save(path)
    monkeypatch.setattr(ai_extract.config, 'BAILIAN_API_KEY', 'test-key')
    monkeypatch.setattr(ai_extract, '_enabled', lambda: True)
    monkeypatch.setattr(ai_extract.llm, 'chat_text', lambda *a, **kw:
                        '{"header_row":2,"columns":{"型号":1,"图片":null,"电吹风":3,"价格":11}}')
    fields = [{'key': key, 'label': label, 'role': role} for key, label, role in [
        ('model', '型号', 'model'), ('image', '图片', 'image'),
        ('kind', '电吹风', 'spec'), ('price', '价格', 'price')]]
    assert ai_extract.map_approved_fields(str(path), fields, 'Sheet1') == {
        'header_row': 2, 'columns': {'model': 1, 'image': 2, 'kind': 3, 'price': 11}}
