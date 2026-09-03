import os
import sqlite3

import openpyxl
import pytest

from catalog import db, quote
from catalog.storage import LocalStorage


import base64
PNG = base64.b64decode(
    'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==')


@pytest.fixture()
def seeded(tmp_path):
    conn = sqlite3.connect(':memory:', check_same_thread=False)
    conn.row_factory = sqlite3.Row
    db.init_db(conn)
    st = LocalStorage(str(tmp_path))
    rel = st.save('curler', 'p1', 'main.png', PNG)
    conn.execute(
        "INSERT INTO product_curler(id, inner_code, item_no, ctn_qty, price, "
        "voltage, power, material, image_main) VALUES(?,?,?,?,?,?,?,?,?)",
        ('p1', 'KS-AAAAAAAA', '8226', '40', '21.5', '110-240', '44w', 'PBT', rel))
    conn.commit()
    return conn, st, str(tmp_path)


def test_quote_five_cols_and_rows(seeded):
    conn, st, base = seeded
    out = base + '/q.xlsx'
    quote.generate(conn, st, 'curler', ['p1'], out)
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert headers[:5] == ['型号', '图片', '价格', '产品规格', '起订量']
    assert ws['A2'].value == '8226' and ws['C2'].value == '21.5'
    assert ws['E2'].value == '40'                       # 卷发棒：起订量=装箱数量
    assert '110-240' in ws['D2'].value                  # 规格拼串含电压
    assert len(ws._images) >= 1                         # 图片已嵌入


def test_quote_razor_moq_from_ctn_spec(seeded):
    conn, st, base = seeded
    conn.execute(
        "INSERT INTO product_razor(id, inner_code, model_no, ctn_spec, price, "
        "size_mm, color, image_main) VALUES(?,?,?,?,?,?,?,?)",
        ('r1', 'KS-BBBBBBBB', 'S1', '40pcs/箱', '9.9', '120x60', '黑', ''))
    conn.commit()
    out = base + '/q2.xlsx'
    quote.generate(conn, st, 'razor', ['r1'], out)
    ws = openpyxl.load_workbook(out).active
    assert ws['E2'].value == '40'                       # 剃须刀：从箱规文本取数量
    assert '120x60' in ws['D2'].value


def test_quote_async_job_with_file_push(tmp_path, monkeypatch):
    """报价单异步化：job+估时 → 后台生成 → 完成推文件（notify.push_file 被调）。"""
    import sqlite3
    from catalog import db
    conn = sqlite3.connect(':memory:', check_same_thread=False)
    conn.row_factory = sqlite3.Row
    db.init_db(conn)
    from catalog.storage import LocalStorage
    st = LocalStorage(str(tmp_path))
    import base64
    png = base64.b64decode(
        'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==')
    rel = st.save('curler', 'p1', 'main.png', png)
    conn.execute("INSERT INTO product_curler(id,inner_code,item_no,ctn_qty,price,"
                 "voltage,image_main) VALUES('p1','KS-AAAAAAAA','8226','40','21.5','110-240',?)",
                 (rel,))
    conn.commit()
    pushed = []
    import catalog.notify as notify_mod
    monkeypatch.setattr(notify_mod, 'push_file',
                        lambda text, fp: pushed.append((text, fp)))
    from catalog import quote as quote_mod
    import time
    job = quote_mod.start_job(conn, st, 'curler', ['p1'], str(tmp_path))
    assert job['est_sec'] >= 10
    for _ in range(100):
        time.sleep(0.05)
        s = quote_mod.job_status(job['job_id'])
        if s['status'] != 'building':
            break
    assert s['status'] == 'done', s
    assert os.path.exists(s['path'])
    assert pushed and pushed[0][1] == s['path']     # 完成即推文件
    # 显示修复：行高≥图高（px→pt 换算后不叠行）
    wb = openpyxl.load_workbook(s['path'])
    assert wb.active.row_dimensions[2].height >= 40


def test_multi_item_quote_with_adjustment(seeded):
    """多商品报价（v2.2 ELETRO BELEZA 模板）：各自数量 + 百分比调整 + 小计公式 + 合计。"""
    import openpyxl as _oxl
    conn, st, base = seeded
    conn.execute("INSERT INTO product_curler(id,inner_code,item_no,ctn_qty,ctn_size,"
                 "price,voltage,image_main) VALUES('p2','KS-BBBBBBBB','8227','50',"
                 "'60*40*40','30','220v',NULL)")
    conn.commit()
    from catalog import quote as q
    out = base + '/multi.xlsx'
    q.generate_v2(conn, st, [
        {'category': 'curler', 'product_id': 'p1', 'quantity': 100},
        {'category': 'curler', 'product_id': 'p2', 'quantity': 500},
    ], price_adjustment_pct=3, out_path=out)
    ws = _oxl.load_workbook(out).active
    # 表头（17行：F=QUANTITY G=TOTAL AMOUNT）
    assert ws.cell(17, 6).value == 'QUANTITY'
    assert ws.cell(17, 7).value == 'TOTAL AMOUNT'
    # 数据
    assert ws.cell(18, 1).value == '8226'          # ITEM NO.
    assert ws.cell(18, 5).value == 22              # 21.5*1.03=22.145 round=22
    assert ws.cell(18, 6).value == 100             # QUANTITY
    assert ws.cell(18, 7).value == '=ROUND(E18*F18,2)'
    assert ws.cell(19, 1).value == '8227'
    assert ws.cell(19, 5).value == 31              # 30*1.03=30.9 round=31
    assert ws.cell(19, 6).value == 500
    assert ws.cell(19, 7).value == '=ROUND(E19*F19,2)'
    # 合计（k=2 删4空行 → TOTAL 在 18+2+3=23）
    assert ws.cell(23, 1).value == 'TOTAL'
    assert ws.cell(23, 7).value == '=SUM(G18:G19)'
    assert ws.cell(24, 7).value == '=ROUND(G23*0.3,2)'    # 默认定金30%
    assert ws.cell(25, 7).value == '=G23-G24'
