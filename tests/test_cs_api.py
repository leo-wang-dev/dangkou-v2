"""C端 API：红线读取/修改（微信对话→工具→审批→生效）+ 清单链接页数据 + Excel 导出。"""
import io
import sqlite3

import pytest
from fastapi.testclient import TestClient

from catalog import db
from catalog.main import app


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setattr('catalog.config.DB_PATH', str(tmp_path / 't.db'))
    monkeypatch.setattr('catalog.config.IMG_DIR', str(tmp_path / 'img'))
    conn = sqlite3.connect(':memory:', check_same_thread=False)
    conn.row_factory = sqlite3.Row
    db.init_db(conn)
    app.state.conn = conn
    app.state.token = 'test-service-token'
    app.state.storage = None
    app.state.callback = None
    return TestClient(app, headers={'X-Service-Token': 'test-service-token'})


def _seed_customer_note(conn):
    conn.execute("INSERT INTO cs_customer(id, tg_id, tg_name) VALUES('c1','100','买家')")
    conn.execute("INSERT INTO cs_note(customer_id, photo, fields_json, status) VALUES("
                 "'c1','', '{\"价格\":\"80R\",\"颜色\":\"黑色\"}', 'confirmed')")
    conn.commit()


# ---------- 红线 ----------

def test_get_redline_default(client):
    r = client.get('/cs/redline')
    assert r.status_code == 200
    assert r.json()['text_raw'] == ''                # 平台不预置红线


def test_set_redline_via_approval(client):
    """微信 AI 工具调用 → 建审批工单 → 批准 → 生效（与改价格同款管道）。"""
    r = client.post('/cs/redline', json={'text_raw': '数量少于50的转人工'})
    assert r.status_code == 200
    tid = r.json()['ticket_id']
    t = [x for x in client.get('/tickets').json()['tickets'] if x['id'] == tid][0]
    payload = client.get(f'/tickets/{tid}').json()['payload']
    assert payload['kind'] == 'redline'
    assert payload['text_raw'] == '数量少于50的转人工'
    old = client.get('/cs/redline').json()['text_raw']
    client.post(f"/tickets/{t['id']}/decision", json={'token': t['token'], 'approved': True})
    new = client.get('/cs/redline').json()['text_raw']
    assert old != new and new == '数量少于50的转人工'   # 批准即生效


def test_reject_redline_no_change(client):
    r = client.post('/cs/redline', json={'text_raw': 'X'})
    tid = r.json()['ticket_id']
    t = [x for x in client.get('/tickets').json()['tickets'] if x['id'] == tid][0]
    before = client.get('/cs/redline').json()['text_raw']
    client.post(f"/tickets/{t['id']}/decision", json={'token': t['token'], 'approved': False})
    assert client.get('/cs/redline').json()['text_raw'] == before


def test_set_product_redline(client):
    r = client.post('/cs/redline', json={'product_id': 'p1', 'text_raw': '这款50个起'})
    tid = r.json()['ticket_id']
    t = [x for x in client.get('/tickets').json()['tickets'] if x['id'] == tid][0]
    client.post(f"/tickets/{t['id']}/decision", json={'token': t['token'], 'approved': True})
    assert client.get('/cs/redline?product_id=p1').json()['text_raw'] == '这款50个起'
    assert client.get('/cs/redline?product_id=p2').json()['text_raw'] == ''      # 未设置则为空


# ---------- 清单链接 ----------

def test_link_list_and_edit(client):
    conn = app.state.conn
    _seed_customer_note(conn)
    token = 'tok123'
    conn.execute("INSERT INTO cs_link(token, customer_id) VALUES(?, 'c1')", (token,))
    conn.commit()
    r = client.get(f'/cs/link/{token}')
    assert r.status_code == 200
    notes = r.json()['notes']
    assert notes[0]['fields']['价格'] == '80R'
    nid = notes[0]['id']
    r = client.patch(f'/cs/link/{token}/note/{nid}', json={'field': '价格', 'value': '2.5'})
    assert r.status_code == 200
    assert client.get(f'/cs/link/{token}').json()['notes'][0]['fields']['价格'] == '2.5'


def test_link_bad_token_404(client):
    assert client.get('/cs/link/nope').status_code == 404


def test_link_export_excel(client):
    conn = app.state.conn
    _seed_customer_note(conn)
    token = 'tok456'
    conn.execute("INSERT INTO cs_link(token, customer_id) VALUES(?, 'c1')", (token,))
    conn.commit()
    r = client.get(f'/cs/link/{token}/export.xlsx')
    assert r.status_code == 200
    assert 'spreadsheet' in r.headers['content-type']
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    heads = [c.value for c in ws[1]]
    assert '价格' in heads and '颜色' in heads          # 动态字段列
    assert ws.cell(2, heads.index('价格') + 1).value == '80R'


def test_draft_list_photo_edit_and_export_without_confirmation(client, tmp_path):
    import json
    import openpyxl
    from PIL import Image
    conn = app.state.conn
    _seed_customer_note(conn)
    photo = tmp_path / 'draft.jpg'
    Image.new('RGB', (64, 64), 'blue').save(photo)
    conn.execute("UPDATE cs_note SET status='draft',photo=?", (str(photo),))
    conn.execute("INSERT INTO cs_link(token,customer_id) VALUES('draft-link','c1')")
    conn.execute("INSERT INTO cs_customer(id,tg_id) VALUES('c2','200')")
    conn.execute("INSERT INTO cs_note(customer_id,fields_json,status) VALUES('c2','{\"其他客户\":\"secret\"}','draft')")
    conn.execute("INSERT INTO cs_note(customer_id,fields_json,status) VALUES('c1','{\"已删除\":\"deleted\"}','discarded')")
    conn.commit()
    notes = client.get('/cs/link/draft-link').json()['notes']
    assert len(notes) == 1 and notes[0]['status'] == 'draft'
    assert client.get(notes[0]['photo']).content == photo.read_bytes()
    assert client.patch(f"/cs/link/draft-link/note/{notes[0]['id']}",json={'field':'颜色','value':'蓝色'}).status_code == 200
    response = client.get('/cs/link/draft-link/export.xlsx')
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    assert sheet.max_row == 2 and len(sheet._images) == 1
    values = str(list(sheet.values))
    assert '蓝色' in values and '待确认' not in values   # 确认状态列已按需求移除
    assert 'secret' not in values and 'deleted' not in values
    assert conn.execute('SELECT status FROM cs_note WHERE id=?',(notes[0]['id'],)).fetchone()[0] == 'draft'


def test_empty_export_returns_actionable_error(client):
    conn = app.state.conn
    conn.execute("INSERT INTO cs_customer(id,tg_id) VALUES('empty','300')")
    conn.execute("INSERT INTO cs_link(token,customer_id) VALUES('empty-link','empty')")
    conn.commit()
    response = client.get('/cs/link/empty-link/export.xlsx')
    assert response.status_code == 409 and '暂无可导出' in response.json()['detail']
