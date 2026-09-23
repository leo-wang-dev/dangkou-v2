import json
import sqlite3
import time

import pytest
from openpyxl import load_workbook

from catalog import db, ingest, tickets
from catalog.storage import LocalStorage
from tests.test_workbook_templates import blowdryer_fixture


@pytest.fixture
def conn():
    value = sqlite3.connect(':memory:', check_same_thread=False)
    value.row_factory = sqlite3.Row
    db.init_db(value)
    yield value
    value.close()


def _wait(conn, doc_id, *, terminal=('ticketed', 'failed', 'template_approved', 'product_approved'), timeout=10):
    for _ in range(timeout * 20):
        value = ingest.status(conn, doc_id)
        if value['status'] in terminal:
            return value
        time.sleep(.05)
    raise AssertionError(f'import {doc_id} did not finish: {ingest.status(conn, doc_id)}')


def test_template_then_products_requires_second_upload_and_keeps_template_ticket_small(conn, tmp_path):
    path = blowdryer_fixture(tmp_path)
    storage = LocalStorage(str(tmp_path / 'storage'))

    template_doc = ingest.start(conn, storage, str(path), None, source_key='vendor-a',
                                mode='new', phase='template')
    assert _wait(conn, template_doc)['status'] == 'ticketed'
    template_ticket = conn.execute(
        "SELECT * FROM approval_ticket WHERE json_extract(payload,'$.doc_id')=?", (template_doc,)
    ).fetchone()
    template_payload = json.loads(template_ticket['payload'])
    assert template_payload['phase'] == 'template'
    assert template_payload['sheets'][0]['drafts'] == {'new': [], 'update': [], 'delist': []}
    assert template_payload['sheets'][0]['image_count'] == 0
    assert conn.execute('SELECT COUNT(*) FROM product_dynamic').fetchone()[0] == 0

    approved = tickets.decide(conn, template_ticket['id'], template_ticket['token'], True)
    assert approved['phase'] == 'template'
    assert ingest.status(conn, template_doc)['status'] == 'template_approved'

    product_doc = ingest.start(conn, storage, str(path), None, source_key='vendor-a',
                               mode='new', phase='products', template_doc_id=template_doc)
    assert product_doc != template_doc
    assert _wait(conn, product_doc)['status'] == 'ticketed'
    product_ticket = conn.execute(
        "SELECT * FROM approval_ticket WHERE json_extract(payload,'$.doc_id')=?", (product_doc,)
    ).fetchone()
    product_payload = json.loads(product_ticket['payload'])
    assert product_payload['phase'] == 'products'
    assert len(product_payload['sheets'][0]['drafts']['new']) == 4
    assert conn.execute('SELECT COUNT(*) FROM product_dynamic').fetchone()[0] == 0

    tickets.decide(conn, product_ticket['id'], product_ticket['token'], True)
    assert ingest.status(conn, product_doc)['status'] == 'product_approved'
    assert conn.execute('SELECT COUNT(*) FROM product_dynamic').fetchone()[0] == 4


def test_products_phase_cannot_bypass_template_approval(conn, tmp_path):
    path = blowdryer_fixture(tmp_path)
    storage = LocalStorage(str(tmp_path / 'storage'))
    with pytest.raises(ValueError, match='template_doc_id'):
        ingest.start(conn, storage, str(path), None, source_key='vendor-a',
                     mode='new', phase='products')

    template_doc = ingest.start(conn, storage, str(path), None, source_key='vendor-a',
                                mode='new', phase='template')
    _wait(conn, template_doc)
    blocked = ingest.start(conn, storage, str(path), None, source_key='vendor-a',
                           mode='new', phase='products', template_doc_id=template_doc)
    blocked_status = _wait(conn, blocked)
    assert blocked_status['status'] == 'failed'
    assert '尚未审批' in blocked_status['error']


def test_each_phase_has_its_own_idempotency_key(conn, tmp_path):
    path = blowdryer_fixture(tmp_path)
    storage = LocalStorage(str(tmp_path / 'storage'))
    first = ingest.start(conn, storage, str(path), None, source_key='vendor-a',
                         mode='new', phase='template')
    again = ingest.start(conn, storage, str(path), None, source_key='vendor-a',
                         mode='new', phase='template')
    assert again == first
    _wait(conn, first)
    ticket = conn.execute("SELECT * FROM approval_ticket WHERE json_extract(payload,'$.doc_id')=?", (first,)).fetchone()
    tickets.decide(conn, ticket['id'], ticket['token'], True)
    product = ingest.start(conn, storage, str(path), None, source_key='vendor-a',
                           mode='new', phase='products', template_doc_id=first)
    product_again = ingest.start(conn, storage, str(path), None, source_key='vendor-a',
                                 mode='new', phase='products', template_doc_id=first)
    assert product_again == product
    assert product != first
    _wait(conn, product)


def test_omitted_mode_round_trips_between_phases(conn, tmp_path):
    """API callers that omit the optional mode keep the same phase contract."""
    path = blowdryer_fixture(tmp_path)
    storage = LocalStorage(str(tmp_path / 'storage'))
    template_doc = ingest.start(conn, storage, str(path), None,
                                source_key='vendor-a', phase='template')
    _wait(conn, template_doc)
    template_ticket = conn.execute(
        "SELECT * FROM approval_ticket WHERE json_extract(payload,'$.doc_id')=?",
        (template_doc,)).fetchone()
    tickets.decide(conn, template_ticket['id'], template_ticket['token'], True)

    product_doc = ingest.start(conn, storage, str(path), None,
                               source_key='vendor-a', phase='products',
                               template_doc_id=template_doc)
    assert _wait(conn, product_doc)['status'] == 'ticketed'


def test_fixed_category_also_requires_template_then_products(conn, monkeypatch, tmp_path):
    path = tmp_path / 'razor.xlsx'
    path.write_bytes(b'xlsx-placeholder')
    storage = LocalStorage(str(tmp_path / 'storage'))
    monkeypatch.setattr(ingest.agent, 'parse',
                        lambda cat, p, wd: {'vendor': '厂',
                                            'products': [{'model_no': 'R-1', 'price': '21'}]})

    template_doc = ingest.start(conn, storage, str(path), 'razor', source_key='vendor-r',
                                mode='new', phase='template')
    _wait(conn, template_doc)
    template_ticket = conn.execute(
        "SELECT * FROM approval_ticket WHERE json_extract(payload,'$.doc_id')=?", (template_doc,)
    ).fetchone()
    assert json.loads(template_ticket['payload'])['fixed_category'] == 'razor'
    tickets.decide(conn, template_ticket['id'], template_ticket['token'], True)

    product_doc = ingest.start(conn, storage, str(path), 'razor', source_key='vendor-r',
                               mode='new', phase='products', template_doc_id=template_doc)
    _wait(conn, product_doc)
    product_ticket = conn.execute(
        "SELECT * FROM approval_ticket WHERE json_extract(payload,'$.doc_id')=?", (product_doc,)
    ).fetchone()
    assert json.loads(product_ticket['payload'])['phase'] == 'products'
    tickets.decide(conn, product_ticket['id'], product_ticket['token'], True)
    assert conn.execute('SELECT COUNT(*) FROM product_razor').fetchone()[0] == 1


def test_existing_mode_maps_all_reuploaded_sheets(conn, tmp_path):
    path = blowdryer_fixture(tmp_path)
    workbook = load_workbook(path)
    duplicate = workbook.copy_worksheet(workbook[workbook.sheetnames[0]])
    duplicate.title = '供应商补充'
    workbook.save(path)
    storage = LocalStorage(str(tmp_path / 'storage'))

    first = ingest.start(conn, storage, str(path), None, source_key='vendor-a',
                         mode='new', phase='template')
    _wait(conn, first)
    first_ticket = conn.execute(
        "SELECT * FROM approval_ticket WHERE json_extract(payload,'$.doc_id')=?",
        (first,)).fetchone()
    tickets.decide(conn, first_ticket['id'], first_ticket['token'], True)
    first_mapping = json.loads(conn.execute(
        'SELECT template_keys_json FROM import_doc WHERE id=?', (first,)).fetchone()[0])
    category_key = first_mapping[0]['category_key']

    existing_template = ingest.start(conn, storage, str(path), None, source_key='vendor-b',
                                     mode='existing', category_key=category_key,
                                     phase='template')
    _wait(conn, existing_template)
    template_ticket = conn.execute(
        "SELECT * FROM approval_ticket WHERE json_extract(payload,'$.doc_id')=?",
        (existing_template,)).fetchone()
    tickets.decide(conn, template_ticket['id'], template_ticket['token'], True)

    products = ingest.start(conn, storage, str(path), None, source_key='vendor-b',
                            mode='existing', category_key=category_key,
                            phase='products', template_doc_id=existing_template)
    _wait(conn, products)
    product_ticket = conn.execute(
        "SELECT * FROM approval_ticket WHERE json_extract(payload,'$.doc_id')=?",
        (products,)).fetchone()
    payload = json.loads(product_ticket['payload'])
    assert len(payload['sheets']) == 1
    assert len(payload['sheets'][0]['drafts']['new']) == 8
