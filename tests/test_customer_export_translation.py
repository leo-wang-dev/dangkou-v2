"""Export display translation must preserve stored evidence and writer availability."""
import io
import json
import sqlite3
import threading

import openpyxl
import pytest
from fastapi.testclient import TestClient

from catalog import cs_export, cs_i18n, userapp, guest_sessions, llm
from tests.test_h5_transactions import h5

FIELDS = {
    '型号或品名':'KS-1100',
    '档口名称':'ACME Ltd',
    '供应商联系方式':'+86 13800138000',
    '单价':66.95,
    '数量':80,
    '颜色':'红色',
    '产品介绍':'红色外壳，型号 KS-1100，供应商 ACME Ltd，电话 +86 13800138000，价格 66.95 USD，https://example.test/KS-1100',
}

class FrenchProvider:
    def __init__(self):self.calls=[]
    def chat_text(self,prompt,messages,**kwargs):
        values=json.loads(messages[0]['content']);self.calls.append(values)
        # Identities inside prose must be occurrence placeholders in the actual request.
        assert all('KS-1100' not in x and 'ACME Ltd' not in x and '+86 13800138000' not in x for x in values)
        return json.dumps([x.replace('产品介绍','Présentation du produit').replace('红色外壳','Coque rouge').replace('红色','Rouge').replace('型号','modèle').replace('供应商','fournisseur').replace('电话','téléphone').replace('价格','prix') for x in values])


def cells(ws):return [cell.value for row in ws for cell in row]


def assert_translated(content):
    ws=openpyxl.load_workbook(io.BytesIO(content)).active
    values=cells(ws)
    assert 'Présentation du produit' in values
    assert 'Rouge' in values
    assert 'KS-1100' in values and 'ACME Ltd' in values and '+86 13800138000' in values
    assert any(c.value==66.95 and c.data_type=='n' for row in ws for c in row)
    assert any(c.value==80 and c.data_type=='n' for row in ws for c in row)
    prose=next(v for v in values if isinstance(v,str) and v.startswith('Coque rouge'))
    for value in ('KS-1100','ACME Ltd','+86 13800138000','66.95 USD','https://example.test/KS-1100'):
        assert prose.count(value)==FIELDS['产品介绍'].count(value)
    return ws


def test_notes_export_translates_prose_and_custom_headers_not_evidence(tmp_path):
    path=tmp_path/'cache.db';conn=sqlite3.connect(path);conn.row_factory=sqlite3.Row
    source=json.dumps(FIELDS,ensure_ascii=False)
    conn.execute('CREATE TABLE evidence(fields_json TEXT)');conn.execute('INSERT INTO evidence VALUES(?)',(source,));conn.commit()
    provider=FrenchProvider()
    note={'fields_json':source,'photo':'','batch_fields':{'档口名称':'ACME Ltd','供应商联系方式':'+86 13800138000'}}
    content=cs_export.render_notes([note],lang='fr',conn=conn,llm=provider)
    assert_translated(content)
    assert provider.calls and conn.execute('SELECT fields_json FROM evidence').fetchone()[0]==source
    assert note['fields_json']==source
    conn.close()


def test_failed_translation_and_existing_writer_use_visible_source_fallback(tmp_path):
    conn=sqlite3.connect(tmp_path/'cache.db');conn.row_factory=sqlite3.Row
    conn.execute('CREATE TABLE evidence(value TEXT)');conn.commit()
    class Broken:
        calls=0
        def chat_text(self,*args,**kwargs):self.calls+=1;raise RuntimeError('offline')
    provider=Broken();note={'fields_json':json.dumps(FIELDS),'photo':''}
    values=cells(openpyxl.load_workbook(io.BytesIO(cs_export.render_notes([note],lang='fr',conn=conn,llm=provider))).active)
    assert provider.calls
    assert cs_i18n.t('translationUnavailable','fr')+': '+FIELDS['产品介绍'] in values
    assert cs_i18n.t('translationUnavailable','fr')+': 产品介绍' in values
    assert 'KS-1100' in values
    conn.execute('INSERT INTO evidence VALUES(?)',('uncommitted',))
    provider.calls=0
    content=cs_export.render_notes([note],lang='fr',conn=conn,llm=provider)
    assert provider.calls==0 and conn.in_transaction
    assert cs_i18n.t('translationUnavailable','fr')+': '+FIELDS['产品介绍'] in cells(openpyxl.load_workbook(io.BytesIO(content)).active)
    conn.rollback();conn.close()


@pytest.mark.parametrize('kind',['central','shop'])
def test_real_http_export_translation_allows_independent_writer_and_keeps_db(h5,tmp_path,monkeypatch,kind):
    shop,shop_path,_=h5
    provider=FrenchProvider();entered=threading.Event();release=threading.Event()
    original=provider.chat_text
    def blocked(*args,**kwargs):
        entered.set();assert release.wait(5)
        return original(*args,**kwargs)
    provider.chat_text=blocked
    source=json.dumps(FIELDS,ensure_ascii=False)
    if kind=='central':
        path=tmp_path/'central.db';app=userapp.build_app(str(path),str(tmp_path/'photos'),str(tmp_path/'codes'),llm=provider)
        capability=guest_sessions.issue(app.state.conn);owner=guest_sessions.validate(app.state.conn,capability)['owner_id']
        app.state.conn.execute('INSERT INTO notes(owner_kind,owner_id,fields_json) VALUES(?,?,?)',('guest',owner,source))
        route='/export.xlsx?guest='+capability;table='notes'
    else:
        app=shop;path=shop_path;table='cs_note'
        app.state.conn.execute("INSERT INTO cs_customer(id,tg_id,lang) VALUES('export-customer','export-owner','fr')")
        app.state.conn.execute("INSERT INTO cs_note(customer_id,fields_json,status) VALUES('export-customer',?,'confirmed')",(source,))
        app.state.conn.execute("INSERT INTO cs_link(token,customer_id) VALUES('export-link','export-customer')")
        route='/cs/link/export-link/export.xlsx'
        monkeypatch.setattr(llm,'chat_text',blocked)
    app.state.conn.execute('CREATE TABLE export_concurrency(value TEXT)');app.state.conn.commit()
    with TestClient(app) as client:
        result=[]
        thread=threading.Thread(target=lambda:result.append(client.get(route,headers={'X-Customer-Language':'fr'})));thread.start()
        try:
            assert entered.wait(2),'export prose never reached translation provider'
            with sqlite3.connect(path,timeout=.2) as independent:
                independent.execute("INSERT INTO export_concurrency VALUES('while translating')")
        finally:release.set();thread.join(5)
        assert not thread.is_alive()
        assert result[0].status_code==200
        assert_translated(result[0].content)
    assert app.state.conn.execute(f'SELECT fields_json FROM {table}').fetchone()[0]==source
    assert app.state.conn.execute('SELECT COUNT(*) FROM export_concurrency').fetchone()[0]==1
    if kind=='central':app.state.conn.close()


def test_multiple_sheets_and_provider_batches_finish_before_cache_writer(tmp_path):
    conn=sqlite3.connect(tmp_path/'cache.db');conn.row_factory=sqlite3.Row
    provider=FrenchProvider();original=provider.chat_text
    def outside_writer(*args,**kwargs):
        assert not conn.in_transaction
        return original(*args,**kwargs)
    provider.chat_text=outside_writer
    notes=[{'batch_id':str(i%2), 'fields_json':json.dumps({**FIELDS,
            '产品介绍':FIELDS['产品介绍']+f' 编号 {i}'}), 'photo':''} for i in range(65)]
    content=cs_export.render_notes(notes,lang='fr',conn=conn,llm=provider)
    wb=openpyxl.load_workbook(io.BytesIO(content))
    assert len(wb.worksheets)==2 and len(provider.calls)==2
    assert sum(isinstance(v,str) and v.startswith('Coque rouge')
               for ws in wb for v in cells(ws))==65
    assert conn.in_transaction  # cache writes happen only after both calls
    conn.rollback();conn.close()


def test_custom_underscore_identity_fields_remain_literal(tmp_path):
    conn=sqlite3.connect(tmp_path/'cache.db');conn.row_factory=sqlite3.Row
    class Provider:
        def chat_text(self,prompt,messages,**kwargs):
            values=json.loads(messages[0]['content'])
            assert all('MARVEL' not in value and 'ACME' not in value for value in values)
            return json.dumps([value.replace('产品介绍','Présentation').replace('红色外壳','Coque rouge') for value in values])
    note={'fields_json':json.dumps({'model_id':'MARVEL','supplier_name':'ACME',
          '产品介绍':'红色外壳 MARVEL ACME'}),'photo':''}
    values=cells(openpyxl.load_workbook(io.BytesIO(cs_export.render_notes(
        [note],lang='fr',conn=conn,llm=Provider()))).active)
    assert 'MARVEL' in values and 'ACME' in values
    assert 'Coque rouge MARVEL ACME' in values
    conn.close()
