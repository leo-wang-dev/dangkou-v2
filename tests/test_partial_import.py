import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from catalog import agent, db, dynamic_import, workbook_templates
from tests.conftest import seed_category

# Capture before the autouse offline fake: these regressions exercise real validation.
_REAL_PARSE_DYNAMIC = agent.parse_dynamic


def book(tmp_path):
    path = tmp_path / 'rows.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = '测试品类'
    ws.append(['型号', '价格']); ws.append(['A', 2]); ws.append(['B', 3])
    wb.save(path)
    return path


def test_style_only_tail_not_data_extent(tmp_path):
    path = book(tmp_path)
    from openpyxl import load_workbook
    wb = load_workbook(path); wb.active.cell(3, 257).fill = PatternFill('solid', fgColor='FFFF00'); wb.save(path)
    assert len(workbook_templates.discover_workbook(path)[0]['rows']) == 2


def test_substantive_far_column_still_rejected(tmp_path):
    path = book(tmp_path)
    from openpyxl import load_workbook
    wb = load_workbook(path); wb.active.cell(3, 257, 'actual data'); wb.save(path)
    with pytest.raises(ValueError, match='200'):
        workbook_templates.discover_workbook(path)


def test_partial_results_keep_valid_and_expose_missing_source(tmp_path, monkeypatch):
    conn = db.connect(':memory:'); db.init_db(conn)
    template = seed_category(conn)
    monkeypatch.setattr(agent, 'parse_dynamic', lambda *a, **kw: {'products': [
        {'model': 'A', 'price': '2', 'source_sheet': '测试品类', 'source_rows': [2]},
        {'model': '', 'source_sheet': '测试品类', 'source_rows': [3]}]})
    result = dynamic_import.build_product_payload(conn, book(tmp_path), tmp_path / 'work',
        source_key='source', template_doc_id=None, category_key=template['key'])
    section = result['sheets'][0]
    assert len(section['drafts']['new']) == 1
    assert section['failures'] and section['failures'][0]['source_rows'] == [3]
    assert section['coverage']['candidate_rows'] == 2
    conn.close()


def test_no_valid_results_raise_actionable_error(tmp_path, monkeypatch):
    conn = db.connect(':memory:'); db.init_db(conn); template = seed_category(conn)
    monkeypatch.setattr(agent, 'parse_dynamic', lambda *a, **kw: {'products': [{'model': ''}]})
    with pytest.raises(ValueError, match='有效商品|解析服务'):
        dynamic_import.build_product_payload(conn, book(tmp_path), tmp_path / 'work',
            source_key='source', template_doc_id=None, category_key=template['key'])
    conn.close()


def test_retry_and_second_approval_never_duplicate_good_products(tmp_path, monkeypatch):
    from catalog import tickets, dynamic_catalog
    conn = db.connect(':memory:'); db.init_db(conn); template = seed_category(conn)
    path = book(tmp_path)
    products = [{'model': 'A', 'price': '2', 'source_rows': [2]}, {'model': '', 'source_rows': [3]}]
    monkeypatch.setattr(agent, 'parse_dynamic', lambda *a, **kw: {'products': products})
    def approve(work):
        payload = dynamic_import.build_product_payload(conn, path, tmp_path / work,
            source_key='source', template_doc_id=None, category_key=template['key'])
        ticket = tickets.create(conn, 'product_import', None, payload)
        tickets.decide(conn, ticket['id'], ticket['token'], True)
        with pytest.raises(tickets.TicketError):
            tickets.decide(conn, ticket['id'], ticket['token'], True)
    approve('first')
    products[1] = {'model': 'B', 'price': '3', 'source_rows': [3]}
    approve('retry')
    assert sorted(p['data']['model'] for p in dynamic_catalog.list_products(conn, template['key'])) == ['A', 'B']
    conn.close()


def test_added_template_image_field_imports_verified_image(tmp_path, monkeypatch):
    from catalog import dynamic_catalog, tickets
    from PIL import Image
    conn = db.connect(':memory:'); db.init_db(conn); template = seed_category(conn)
    fields = [*template['fields'], {'key': 'photo', 'label': '图片', 'type': 'image', 'role': 'image',
        'required': False, 'visibility': 'public', 'searchable': False}]
    template = dynamic_catalog.approve_template(conn, {**template, 'fields': fields}, expected_version=template['version']); conn.commit()
    def parsed(template, source, work, **kw):
        Image.new('RGB', (12, 13), 'red').save(work / 'real.png')
        return {'products': [{'model': 'A', 'source_rows': [2], 'images': ['real.png'], 'image_count': 1}]}
    monkeypatch.setattr(agent, 'parse_dynamic', parsed)
    payload = dynamic_import.build_product_payload(conn, book(tmp_path), tmp_path / 'work',
        source_key='source', template_doc_id=None, category_key=template['key'])
    draft = payload['sheets'][0]['drafts']['new'][0]
    with Image.open(tmp_path / 'work' / draft['image_main']) as image:
        assert image.size == (12, 13)
    tk = tickets.create(conn, 'product_import', None, payload); tickets.decide(conn, tk['id'], tk['token'], True)
    assert dynamic_catalog.list_products(conn, template['key'])[0]['images'] == draft['images']
    conn.close()


def test_missing_image_is_row_failure_with_good_row_preserved(tmp_path, monkeypatch):
    conn = db.connect(':memory:'); db.init_db(conn); template = seed_category(conn)
    monkeypatch.setattr(agent, 'parse_dynamic', lambda *a, **kw: {'products': [
        {'model': 'A', 'source_rows': [2]}, {'model': 'B', 'source_rows': [3], 'images': ['missing.png']}]})
    payload = dynamic_import.build_product_payload(conn, book(tmp_path), tmp_path / 'work',
        source_key='source', template_doc_id=None, category_key=template['key'])
    assert len(payload['sheets'][0]['drafts']['new']) == 1
    assert any('图片文件缺失' in f['reason'] for f in payload['sheets'][0]['failures'])
    conn.close()


def test_failed_sheet_does_not_erase_completed_sheet(tmp_path, monkeypatch):
    import json
    conn = db.connect(':memory:'); db.init_db(conn)
    a = seed_category(conn, key='a', name='A'); b = seed_category(conn, key='b', name='B')
    path = tmp_path / 'two.xlsx'; wb = Workbook(); wb.active.title = 'A'
    for ws in [wb.active, wb.create_sheet('B')]:
        ws.append(['型号', '价格']); ws.append([ws.title, 2])
    wb.save(path)
    doc = conn.execute("INSERT INTO import_doc(filename,category,status,phase,template_keys_json) VALUES('two.xlsx','auto','template_approved','template',?)", (json.dumps([
        {'source_sheet': t['name'], 'category_key': t['key'], 'version': t['version']} for t in [a,b]]),)).lastrowid
    conn.commit()
    def parse(template, *args, sheet='', **kw):
        if sheet == 'B':
            raise RuntimeError('B model failure')
        return {'products': [{'model': 'A', 'source_sheet': 'A', 'source_rows': [2]}]}
    monkeypatch.setattr(agent, 'parse_dynamic', parse)
    payload = dynamic_import.build_product_payload(conn, path, tmp_path / 'work', source_key='source', template_doc_id=doc)
    assert len(payload['sheets'][0]['drafts']['new']) == 1
    assert 'B model failure' in payload['sheets'][1]['failures'][0]['reason']
    conn.close()


def test_legacy_existing_import_accumulates_all_sheets(tmp_path):
    conn = db.connect(':memory:'); db.init_db(conn); template = seed_category(conn)
    path = tmp_path / 'legacy-two.xlsx'; wb = Workbook(); wb.active.title = 'A'
    for ws in [wb.active, wb.create_sheet('B')]:
        ws.append(['型号', '价格']); ws.append([ws.title, 2])
    wb.save(path)
    payload = dynamic_import.build_ticket_payload(conn, path, tmp_path / 'work', source_key='source',
        mode='existing', category_key=template['key'])
    assert {r['data']['model'] for r in payload['sheets'][0]['drafts']['new']} == {'A', 'B'}
    conn.close()


def test_image_anchor_beyond_real_column_cap_rejected(tmp_path):
    from PIL import Image
    from openpyxl.drawing.image import Image as ExcelImage
    path = book(tmp_path)
    from openpyxl import load_workbook
    image = tmp_path / 'far.png'; Image.new('RGB', (5, 5)).save(image)
    wb = load_workbook(path); wb.active.add_image(ExcelImage(image), 'IW2'); wb.save(path)
    with pytest.raises(ValueError, match='200'):
        workbook_templates.discover_workbook(path)


def test_reclaimed_attempt_copies_checkpoint_images_to_own_directory(tmp_path, monkeypatch):
    from PIL import Image
    from pathlib import Path
    conn = db.connect(':memory:'); db.init_db(conn); template = seed_category(conn)
    path = book(tmp_path)
    def parse(template, source, work, **kw):
        Image.new('RGB', (7, 8), 'blue').save(Path(work) / 'a.png')
        return {'products': [{'model': 'A', 'source_rows': [2], 'images': ['a.png']}]}
    monkeypatch.setattr(agent, 'parse_dynamic', parse)
    first = tmp_path / 'job' / 'attempt-first'
    rows = dynamic_import._agent_rows(template, path, first)
    monkeypatch.setattr(agent, 'parse_dynamic', lambda *a, **kw: (_ for _ in ()).throw(AssertionError('must reuse checkpoint')))
    second = tmp_path / 'job' / 'attempt-second'
    recovered = dynamic_import._agent_rows(template, path, second)
    assert len(recovered) == 1
    assert (second / recovered[0]['image_main']).read_bytes() == (first / rows[0]['image_main']).read_bytes()
    (first / rows[0]['image_main']).write_bytes(b'stale worker damaged its own file')
    with Image.open(second / recovered[0]['image_main']) as image:
        assert image.size == (7, 8)
    conn.close()


def test_malformed_failure_coordinates_remain_renderable(tmp_path, monkeypatch):
    conn = db.connect(':memory:'); db.init_db(conn); template = seed_category(conn)
    monkeypatch.setattr(agent, 'parse_dynamic', lambda *a, **kw: {'products': [{'model': 'A', 'source_rows': [2]}],
        'failures': [{'source_sheet': '测试品类', 'source_rows': '3', 'reason': '未识别'}]})
    rows = dynamic_import._agent_rows(template, book(tmp_path), tmp_path / 'work')
    assert rows.failures[0]['source_rows'] == []
    assert '位置' in rows.failures[0]['reason']
    conn.close()


def test_product_spanning_image_only_row_survives_real_agent_boundary(tmp_path, monkeypatch):
    from PIL import Image
    from openpyxl.drawing.image import Image as ExcelImage
    from pathlib import Path
    conn = db.connect(':memory:'); db.init_db(conn); template = seed_category(conn)
    path = tmp_path / 'image-support.xlsx'; wb = Workbook(); ws = wb.active
    ws.title = '测试品类'; ws.append(['型号', '价格']); ws.append(['A', 2])
    photo = tmp_path / 'source.png'; Image.new('RGB', (8, 9), 'red').save(photo)
    ws.add_image(ExcelImage(photo), 'B3'); wb.save(path)
    def container(prompt, source, work):
        (Path(work) / 'a.png').write_bytes(photo.read_bytes())
        return {'products': [{'model': 'A', 'source_sheet': '测试品类', 'source_rows': [2, 3], 'images': ['a.png']}]}
    monkeypatch.setattr(agent, 'parse_dynamic', _REAL_PARSE_DYNAMIC)
    monkeypatch.setattr(agent, '_run_container', container)
    rows = dynamic_import._agent_rows(template, path, tmp_path / 'work')
    assert len(rows) == 1
    assert rows[0]['source_rows'] == [2, 3] and rows[0]['images']
    assert rows.coverage['uncertain'] is False
    conn.close()


def test_mixed_malformed_product_keeps_valid_sibling_at_real_agent_boundary(tmp_path, monkeypatch):
    conn = db.connect(':memory:'); db.init_db(conn); template = seed_category(conn)
    monkeypatch.setattr(agent, 'parse_dynamic', _REAL_PARSE_DYNAMIC)
    monkeypatch.setattr(agent, '_run_container', lambda *a, **kw: {
        'products': [{'model': 'A', 'source_rows': [2]}, None]})
    rows = dynamic_import._agent_rows(template, book(tmp_path), tmp_path / 'work')
    assert [row['data']['model'] for row in rows] == ['A']
    assert any('第 2 个结果不是商品对象' in f['reason'] for f in rows.failures)
    conn.close()


@pytest.mark.parametrize('discovery', ['qwen', 'agent', 'fallback'])
def test_template_preflight_precedes_every_discovery_path(tmp_path, monkeypatch, discovery):
    from catalog import ai_extract
    from openpyxl import load_workbook
    path = book(tmp_path); wb = load_workbook(path); wb.active.cell(3, 257, 'forbidden value'); wb.save(path)
    conn = db.connect(':memory:'); db.init_db(conn)
    response = [{'title': '测试品类', 'header_row': 1, 'columns': [
        {'col': 1, 'label': '型号', 'role': 'model'}, {'col': 2, 'label': '价格', 'role': 'price'}]}]
    calls = []
    monkeypatch.setattr(ai_extract, 'discover_headers', lambda *a: calls.append('qwen') or (response if discovery == 'qwen' else None))
    monkeypatch.setattr(agent, 'parse_dynamic_template', lambda *a: calls.append('agent') or {'sheets': response if discovery == 'agent' else []})
    with pytest.raises(ValueError, match='200'):
        dynamic_import.build_template_payload(conn, path, tmp_path / 'work', source_key='source')
    assert calls == []
    conn.close()


def test_header_evidence_preflights_before_openpyxl(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    path = book(tmp_path); wb = load_workbook(path); wb.active.cell(3, 257, 'forbidden value'); wb.save(path)
    loads = []
    original = workbook_templates.load_workbook
    monkeypatch.setattr(workbook_templates, 'load_workbook', lambda *a, **kw: loads.append(1) or original(*a, **kw))
    with pytest.raises(ValueError, match='200'):
        workbook_templates.extract_header_evidence(path)
    assert loads == []


@pytest.mark.parametrize('discovery', ['qwen', 'agent'])
def test_template_preflight_keeps_style_only_tail_on_successful_model_path(tmp_path, monkeypatch, discovery):
    from catalog import ai_extract
    from openpyxl import load_workbook
    path = book(tmp_path); wb = load_workbook(path)
    wb.active.cell(3, 257).fill = PatternFill('solid', fgColor='FFFF00'); wb.save(path)
    conn = db.connect(':memory:'); db.init_db(conn)
    response = [{'title': '测试品类', 'header_row': 1, 'columns': [
        {'col': 1, 'label': '型号', 'role': 'model'}, {'col': 2, 'label': '价格', 'role': 'price'}]}]
    monkeypatch.setattr(ai_extract, 'discover_headers', lambda *a: response if discovery == 'qwen' else None)
    monkeypatch.setattr(agent, 'parse_dynamic_template', lambda *a: {'sheets': response})
    payload = dynamic_import.build_template_payload(conn, path, tmp_path / 'work', source_key='source')
    assert payload['sheets'][0]['template']['name'] == '测试品类'
    conn.close()
