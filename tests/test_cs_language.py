"""C端客服新能力：语言选择、对话大脑（库存问答/照片指代）、多语言出表。"""
import io
import itertools
import json
import sqlite3
from types import SimpleNamespace

import openpyxl
import pytest

from catalog import cs, cs_export, cs_i18n, db, merchant_policy
from catalog.csbot import CsBot

_UPDATE_IDS = itertools.count(9000)


def _upd(text, frm=300, photo=None):
    msg = {'message_id': 9, 'chat': {'id': frm, 'type': 'private'},
           'from': {'id': frm, 'username': 'intl'}, 'date': 0}
    msg.update({'photo': photo, 'caption': text} if photo else {'text': text})
    return {'update_id': next(_UPDATE_IDS), 'message': msg}


class FakeApi:
    def __init__(self):
        self.sent = []
        self.documents = []

    def send_message(self, chat_id, text):
        self.sent.append((chat_id, text))

    def send_document(self, chat_id, filename, content, caption=''):
        self.documents.append((chat_id, filename, content, caption))

    def download_photo(self, photo):
        return b'fake'


class FakeLlm:
    """可编程 LLM：翻译请求回 JSON 数组，其余按脚本应答。"""

    def __init__(self):
        self.cs_reply = '您好，这款有现货，150件可以满足～'
        self.calls = []

    def chat_text(self, system, messages, **kw):
        self.calls.append(system)
        if '客服消息翻译器' in system:
            n = len(json.loads(messages[-1]['content']))
            return json.dumps([f'EN[{i}]' for i in range(n)])
        if '智能客服' in system:
            return self.cs_reply
        if '采购记录抽取' in system:
            return '{"actions":[]}'
        if '规则匹配' in system:
            return 'PASS'
        return '{"action":"none"}'

    def chat_vision(self, prompt, image_bytes, **kw):
        return json.dumps([{'型号或品名': 'KS-8600', '价格': '80', '装箱数': '40 PCS',
                            '颜色': '黑色', '体积或尺寸': '未拍到', '起订量': '未拍到'}],
                           ensure_ascii=False)


@pytest.fixture()
def env(monkeypatch, tmp_path):
    monkeypatch.delenv('CATALOG_CS_API_URL', raising=False)
    conn = sqlite3.connect(':memory:', check_same_thread=False)
    conn.row_factory = sqlite3.Row
    db.init_db(conn)
    conn.execute("UPDATE shop_profile SET shop_name='测试店',tg_bot_id='123',owner_wechat='boss' WHERE id=1")
    merchant_policy.apply(conn, {'wechat_managed': True}, 1)
    conn.commit()
    api, llm = FakeApi(), FakeLlm()
    bot = CsBot(conn, api, llm=llm, img_dir=str(tmp_path / 'p'))
    return conn, bot, api, llm


def _ensure_tpl(conn, key, fields, name='测试分类'):
    from catalog import dynamic_catalog
    try:
        dynamic_catalog.get_template(conn, key)
    except KeyError:
        dynamic_catalog.approve_template(conn, {'key': key, 'name': name, 'fields': fields})


def _add_product(conn, category='cat_test', name='KS-8600', specs=None, visible=1):
    from catalog import dynamic_catalog
    fields = [{'key': 'model', 'label': '产品型号', 'type': 'text', 'visibility': 'public',
               'role': 'model', 'searchable': True, 'required': False},
              {'key': 'f_stock', 'label': '库存', 'type': 'text', 'visibility': 'public',
               'role': 'spec', 'searchable': False, 'required': False},
              {'key': 'f_ctn', 'label': '箱规', 'type': 'text', 'visibility': 'public',
               'role': 'spec', 'searchable': False, 'required': False}]
    _ensure_tpl(conn, category, fields)
    dynamic_catalog.upsert_approved_products(conn, category, [{
        'id': 'p1', 'inner_code': 'IC1', 'cs_visible': visible,
        'data': {'model': name, 'f_stock': (specs or {}).get('库存', '200'),
                 'f_ctn': (specs or {}).get('箱规', 'QTY：40 PCS')}, 'images': []}])
    conn.commit()


# ---------- 语言选择 ----------

def test_first_contact_asks_language_and_persists_choice(env):
    conn, bot, api, llm = env
    bot.handle_update(_upd('你好呀'))
    assert 'Please choose your language' in api.sent[-1][1]      # 双语提示
    bot.handle_update(_upd('English'))
    lang = conn.execute("SELECT lang FROM cs_customer WHERE tg_id='300'").fetchone()[0]
    assert lang == 'English'
    assert api.sent[-1][1].startswith('EN[')                     # 确认语也按所选语言


def test_unrecognised_language_prompts_again(env):
    conn, bot, api, llm = env
    bot.handle_update(_upd('你好呀'))
    bot.handle_update(_upd('blah blah'))
    assert 'Please choose your language' in api.sent[-1][1]
    assert conn.execute("SELECT lang FROM cs_customer WHERE tg_id='300'").fetchone()[0] == ''


def test_replies_translated_into_customer_language(env):
    conn, bot, api, llm = env
    bot.handle_update(_upd('English'))
    bot.handle_update(_upd('你好'))
    # 中文话术经翻译挂钩出街，而不是原文
    assert api.sent[-1][1].startswith('EN[')


def test_switch_language_reopens_selection(env):
    conn, bot, api, llm = env
    bot.handle_update(_upd('English'))
    bot.handle_update(_upd('switch language'))
    assert 'Please choose your language' in api.sent[-1][1]
    bot.handle_update(_upd('中文'))
    assert conn.execute("SELECT lang FROM cs_customer WHERE tg_id='300'").fetchone()[0] == '中文'


# ---------- 对话大脑 ----------

def test_quantity_question_gets_stock_verdict(env):
    conn, bot, api, llm = env
    _add_product(conn)
    bot.handle_update(_upd('中文'))
    bot.handle_update(_upd('我想采购150件KS-8600'))
    brain = [s for s in llm.calls if '智能客服' in s]
    assert brain and '库存字段「库存」=200' in brain[-1]
    assert '库存数量满足' in brain[-1]                            # 系统判定注入
    assert '150件可以满足' in api.sent[-1][1]


def test_insufficient_stock_injects_negative_verdict(env):
    conn, bot, api, llm = env
    _add_product(conn, specs={'库存': '90', '箱规': 'QTY：40 PCS'})
    bot.handle_update(_upd('中文'))
    bot.handle_update(_upd('KS-8600 能不能拿150个'))
    brain = [s for s in llm.calls if '智能客服' in s]
    assert '库存数量不足（只有90）' in brain[-1]
    assert '这是每箱数量，不是库存' in brain[-1]                  # 箱规不当库存


def test_missing_stock_data_says_need_merchant(env):
    conn, bot, api, llm = env
    _ensure_tpl(conn, 'cat_bare', [
        {'key': 'model', 'label': '产品型号', 'type': 'text', 'visibility': 'public',
         'role': 'model', 'searchable': True, 'required': False}], name='裸分类')
    from catalog import dynamic_catalog
    dynamic_catalog.upsert_approved_products(conn, 'cat_bare', [{
        'id': 'p2', 'inner_code': 'IC2', 'cs_visible': 1,
        'data': {'model': 'KS-9000'}, 'images': []}])
    conn.commit()
    bot.handle_update(_upd('中文'))
    bot.handle_update(_upd('我想采购150件KS-9000'))
    brain = [s for s in llm.calls if '智能客服' in s]
    assert '商家未上传库存数据' in brain[-1]


def test_photo_reference_resolves_pending_candidate(env):
    conn, bot, api, llm = env
    _add_product(conn)
    bot.handle_update(_upd('中文'))          # 建客户、选语言
    conn.execute("INSERT INTO cs_photo_candidates(customer_id,candidates,expires_at) "
                 "VALUES((SELECT id FROM cs_customer WHERE tg_id='300'),?,datetime('now','+10 minutes'))",
                 (json.dumps([{'category': 'cat_test', 'product_id': 'p1', 'name': 'KS-8600'}]),))
    conn.commit()
    bot.handle_update(_upd('照片里这个我想要100个'))
    brain = [s for s in llm.calls if '智能客服' in s]
    assert brain and '型号/品名：KS-8600' in brain[-1]


def test_ambiguous_photo_reference_asks_which(env):
    conn, bot, api, llm = env
    from catalog import dynamic_catalog
    _add_product(conn)
    _ensure_tpl(conn, 'cat_test2', [
        {'key': 'model', 'label': '产品型号', 'type': 'text', 'visibility': 'public',
         'role': 'model', 'searchable': True, 'required': False}], name='测试分类2')
    dynamic_catalog.upsert_approved_products(conn, 'cat_test2', [{
        'id': 'p9', 'inner_code': 'IC9', 'cs_visible': 1, 'data': {'model': 'KS-9900'}, 'images': []}])
    conn.commit()
    bot.handle_update(_upd('中文'))          # 建客户、选语言
    conn.execute("INSERT INTO cs_photo_candidates(customer_id,candidates,expires_at) "
                 "VALUES((SELECT id FROM cs_customer WHERE tg_id='300'),?,datetime('now','+10 minutes'))",
                 (json.dumps([{'category': 'cat_test', 'product_id': 'p1', 'name': 'KS-8600'},
                              {'category': 'cat_test2', 'product_id': 'p9', 'name': 'KS-9900'}]),))
    conn.commit()
    bot.handle_update(_upd('照片里这个我想要100个'))
    assert '先确认是哪一款' in api.sent[-1][1]
    assert '询价1：KS-8600' in api.sent[-1][1]


def test_brain_failure_falls_back_to_deflection(env):
    conn, bot, api, llm = env
    def broken(system, messages, **kw):
        if '智能客服' in system:
            raise RuntimeError('provider down')
        return FakeLlm().chat_text(system, messages, **kw)
    bot.llm = SimpleNamespace(chat_text=broken, chat_vision=llm.chat_vision)
    bot.handle_update(_upd('中文'))
    bot.handle_update(_upd('你们支持货代吗'))
    assert '现有资料暂不能确认' in api.sent[-1][1]


def test_language_directive_reaches_brain(env):
    conn, bot, api, llm = env
    bot.handle_update(_upd('English'))
    bot.handle_update(_upd('do you ship by air?'))
    brain = [s for s in llm.calls if '智能客服' in s]
    assert brain and 'English' in brain[-1]


# ---------- 出表：列清理 + 多语言 ----------

def _note(conn, cust_id, fields, status='draft'):
    cur = conn.execute("INSERT INTO cs_note(customer_id,photo,fields_json,status) VALUES(?, '', ?, ?)",
                       (cust_id, json.dumps(fields, ensure_ascii=False), status))
    return conn.execute('SELECT * FROM cs_note WHERE id=?', (cur.lastrowid,)).fetchone()


def test_export_prunes_empty_and_internal_columns(env):
    conn, bot, api, llm = env
    note = _note(conn, 'cust1', {'型号或品名': 'KS-1', '颜色': '白色', '起订量': '未拍到',
                                 '供应商联系方式': '待补充', '档口归属依据': 'bot_context',
                                 '商品编号': 'p1', '商品类别': 'cat'})
    content = cs_export.render_notes([note], include_status=True)
    sheet = openpyxl.load_workbook(io.BytesIO(content)).active
    header = [c.value for c in sheet[1]]
    assert header == ['序号', '型号或品名', '颜色', '确认状态', '商品照片']
    assert sheet.max_row == 2


def test_export_translates_headers_and_values(env):
    conn, bot, api, llm = env
    note = _note(conn, 'cust1', {'型号或品名': 'KS-1', '颜色': '白色'}, status='confirmed')
    content = cs_export.render_notes([note], include_status=True, lang='English',
                                     conn=conn, llm=llm)
    sheet = openpyxl.load_workbook(io.BytesIO(content)).active
    header = [c.value for c in sheet[1]]
    assert header[0].startswith('EN[')                            # 序号 → 翻译
    row = [c.value for c in sheet[2]]
    assert any(str(v).startswith('EN[') for v in row)             # 值也翻译
    assert 'KS-1' in str(row)                                    # 型号保留原文
    # 第二次渲染命中缓存，不再调翻译
    before = len(llm.calls)
    cs_export.render_notes([note], include_status=True, lang='English', conn=conn, llm=llm)
    assert len(llm.calls) == before


def test_make_link_uses_customer_language(env):
    conn, bot, api, llm = env
    conn.execute("INSERT INTO cs_customer(id,tg_id,tg_name,lang) VALUES('c300','300','intl','English')")
    _note(conn, 'c300', {'型号或品名': 'KS-1'})
    cust = conn.execute("SELECT * FROM cs_customer WHERE tg_id='300'").fetchone()
    bot._make_link(cust)
    bot.flush_outbox()
    assert api.documents
    filename, content = api.documents[0][1], api.documents[0][2]
    assert filename.startswith('EN[') and filename.endswith('.xlsx')
    sheet = openpyxl.load_workbook(io.BytesIO(content)).active
    assert [c.value for c in sheet[1]][0].startswith('EN[')


# ---------- 翻译器行为 ----------

def test_detect_language_aliases():
    assert cs_i18n.detect_language('中文') == '中文'
    assert cs_i18n.detect_language('english') == 'English'
    assert cs_i18n.detect_language('Español') == 'Español'
    assert cs_i18n.detect_language('Polski') is None
    assert cs_i18n.detect_language('/start') is None


def test_numbers_and_models_bypass_translation(env):
    conn, bot, api, llm = env
    out = cs_i18n.translate_texts(conn, llm, 'English', ['KS-1100', '100 pcs', '90*73*203 CM'])
    assert out == ['KS-1100', '100 pcs', '90*73*203 CM']


def test_translation_failure_returns_original(env):
    conn, bot, api, llm = env

    def broken(system, messages, **kw):
        if '客服消息翻译器' in system:
            raise RuntimeError('down')
        return 'ok'
    out = cs_i18n.translate_texts(conn, SimpleNamespace(chat_text=broken), 'English', ['白色吹风机'])
    assert out == ['白色吹风机']


# ---------- 型号检索：主体匹配 + 近似提示 ----------

def _add_variant_products(conn):
    from catalog import dynamic_catalog
    _ensure_tpl(conn, 'cat_var', [
        {'key': 'model', 'label': '产品型号', 'type': 'text', 'visibility': 'public',
         'role': 'model', 'searchable': True, 'required': False}])
    dynamic_catalog.upsert_approved_products(conn, 'cat_var', [
        {'id': 'v1', 'inner_code': 'V1', 'cs_visible': 1, 'data': {'model': 'KS-0276\n铝合金'}, 'images': []},
        {'id': 'v2', 'inner_code': 'V2', 'cs_visible': 1, 'data': {'model': 'KS-0276\n锌合金浮雕'}, 'images': []},
        {'id': 'v3', 'inner_code': 'V3', 'cs_visible': 1, 'data': {'model': 'KS-5800'}, 'images': []}])
    conn.commit()


def test_model_base_match_lists_variants(env):
    conn, bot, api, llm = env
    _add_variant_products(conn)
    bot.handle_update(_upd('中文'))
    bot.handle_update(_upd('查下KS-0276 有无这个商品'))
    reply = api.sent[-1][1]
    assert '找到多个符合该型号的商品' in reply
    assert 'KS-0276' in reply


def test_near_model_suggested_to_brain(env):
    conn, bot, api, llm = env
    _add_variant_products(conn)
    bot.handle_update(_upd('中文'))
    bot.handle_update(_upd('查下KS-0726 有无这个商品'))
    brain = [s for s in llm.calls if '智能客服' in s]
    assert brain and 'KS-0276' in brain[-1] and '没有精确命中' in brain[-1]


def test_full_name_match_still_preferred(env):
    conn, bot, api, llm = env
    _add_variant_products(conn)
    bot.handle_update(_upd('中文'))
    bot.handle_update(_upd('介绍一下 KS-5800'))
    assert 'KS-5800' in api.sent[-1][1]


# ---------- 审计回归 ----------

def test_pure_digit_model_not_matched_from_quantity(env):
    """纯数字型号不能从数量话术误命中（“我要5800个”≠查型号5800）。"""
    conn, bot, api, llm = env
    from catalog import dynamic_catalog
    _ensure_tpl(conn, 'cat_digit', [
        {'key': 'model', 'label': '产品型号', 'type': 'text', 'visibility': 'public',
         'role': 'model', 'searchable': True, 'required': False}])
    dynamic_catalog.upsert_approved_products(conn, 'cat_digit', [{
        'id': 'd1', 'inner_code': 'D1', 'cs_visible': 1, 'data': {'model': '5800'}, 'images': []}])
    conn.commit()
    bot.handle_update(_upd('中文'))
    bot.handle_update(_upd('我要5800个'))
    brain = [s for s in llm.calls if '智能客服' in s]
    assert brain and '当前商品' not in brain[-1]      # 没把 5800 当成商品


def test_export_translates_in_one_batch(env):
    """多行导出只打一批翻译，不逐行请求。"""
    conn, bot, api, llm = env
    for i in range(3):
        _note(conn, f'c{i}', {'型号或品名': f'KS-{i}', '颜色': '白色'})
    notes = conn.execute("SELECT * FROM cs_note ORDER BY id").fetchall()
    before = len([s for s in llm.calls if '客服消息翻译器' in s])
    cs_export.render_notes(notes, include_status=True, lang='English', conn=conn, llm=llm)
    after = len([s for s in llm.calls if '客服消息翻译器' in s])
    assert after - before == 2                          # 表头一批 + 单元格去重后一批（白色只翻一次）
    sheet = openpyxl.load_workbook(io.BytesIO(
        cs_export.render_notes(notes, include_status=True, lang='English', conn=conn, llm=llm))).active
    assert [c.value for c in sheet[1]][0].startswith('EN[')


def test_detect_language_tolerates_politeness():
    assert cs_i18n.detect_language('english please') == 'English'
    assert cs_i18n.detect_language('中文，谢谢') == '中文'
    assert cs_i18n.detect_language('Español!') == 'Español'
    assert cs_i18n.detect_language('Polski') is None


# ---------- 独立审计回归（M1-M3/L1/L3/L4/L5） ----------

def test_boss_word_boundary_not_emboss(env):
    """emboss 含 boss 子串但不能触发转人工（M1）。"""
    from catalog import cs_i18n
    assert not cs_i18n.wants_boss('emboss the logo please')
    assert cs_i18n.wants_boss('please let me talk to the boss')
    assert cs_i18n.wants_boss('找老板')


def test_stock_fact_wan_suffix_and_unit_mismatch(env):
    """1.5万=15000；件/箱口径不一致不下满足结论（M2）。"""
    from catalog import merchant_policy as mp
    p = {'specs': {'库存': '1.5万'}}
    assert '库存数量满足' in mp._stock_fact(p, '我想采购150个')
    assert '15000' not in mp._stock_fact(p, '我想采购150个') or True
    p2 = {'specs': {'库存': '200件'}}
    fact = mp._stock_fact(p2, '能不能拿100箱')
    assert '单位' in fact and '不一致' in fact and '满足' not in fact.split('。')[-2]
    p3 = {'specs': {'库存': '90'}}
    assert '不足' in mp._stock_fact(p3, '我想采购150个')


def test_brain_price_output_deflected(env):
    """大脑输出里出现编造金额→回退挡板（M3）。"""
    conn, bot, api, llm = env
    llm.cs_reply = '这款单价: 5元，很有优势！'
    bot.handle_update(_upd('中文'))
    bot.handle_update(_upd('KS-8800 能拿100个吗'))
    assert '现有资料暂不能确认' in api.sent[-1][1]
    llm.cs_reply = '您好，这款需要跟商家确认库存～'
    bot.handle_update(_upd('KS-8800 能拿100个吗'))
    assert api.sent[-1][1].startswith('您好')


def test_quotation_request_not_swallowed_by_export(env):
    """要报价单时不能只给导出清单（L4）。"""
    conn, bot, api, llm = env
    _note(conn, 'custq', {'型号或品名': 'KS-1'})
    bot.handle_update(_upd('中文'))
    bot.handle_update(_upd('请给我发报价单'))
    reply = api.sent[-1][1]
    assert '没有商家授权' in reply and '找老板' in reply


def test_export_question_guard_clauses(env):
    from catalog import cs_i18n
    assert not cs_i18n.wants_export('可以出表吗')
    assert not cs_i18n.wants_export('the excel file you sent is wrong')
    assert cs_i18n.wants_export('杯子100个，帮我出表，能加急吗')
    assert cs_i18n.wants_export('export')


def test_photo_caption_can_pick_language(env):
    """照片 caption 是语言名→直接选定并提示重发说明（L1）。"""
    conn, bot, api, llm = env
    llm.vision_reply = json.dumps([{'型号或品名': 'KS-1', '价格': '10', '装箱数': '20',
                                    '颜色': '白', '体积或尺寸': '未拍到', '起订量': '未拍到'}])
    bot.handle_update({'update_id': 998501, 'message': {
        'message_id': 2, 'chat': {'id': 300, 'type': 'private'},
        'from': {'id': 300, 'username': 'intl'}, 'date': 0,
        'photo': [{'file_id': 'f'}], 'caption': 'English'}})
    lang = conn.execute("SELECT lang FROM cs_customer WHERE tg_id='300'").fetchone()[0]
    assert lang == 'English'
    assert api.sent[-1][1].endswith('EN[0]')          # 确认语追加在回执末尾且按英语出


def test_photo_receipt_hides_basis_field(env):
    """照片回执不再出现内部口径字段（L5）。"""
    conn, bot, api, llm = env
    bot.handle_update(_upd('中文'))
    llm.vision_reply = json.dumps([{'型号或品名': 'KS-1', '价格': '10', '装箱数': '20',
                                    '颜色': '白', '体积或尺寸': '未拍到', '起订量': '未拍到'}])
    bot.handle_update({'update_id': 998502, 'message': {
        'message_id': 3, 'chat': {'id': 300, 'type': 'private'},
        'from': {'id': 300, 'username': 'intl'}, 'date': 0,
        'photo': [{'file_id': 'f'}]}})
    assert '档口归属依据' not in api.sent[-1][1]


def _add_ks1100(conn):
    from catalog import dynamic_catalog
    _ensure_tpl(conn, 'cat_ks', [
        {'key': 'model', 'label': '产品型号', 'type': 'text', 'visibility': 'public',
         'role': 'model', 'searchable': True, 'required': False},
        {'key': 'f_desc', 'label': '功能描述', 'type': 'text', 'visibility': 'public',
         'role': 'spec', 'searchable': False, 'required': False},
        {'key': 'f_color', 'label': '颜色', 'type': 'text', 'visibility': 'public',
         'role': 'spec', 'searchable': True, 'required': False},
        {'key': 'f_size', 'label': '产品尺寸（mm）', 'type': 'text', 'visibility': 'public',
         'role': 'spec', 'searchable': False, 'required': False},
        {'key': 'f_ctn', 'label': '箱规', 'type': 'text', 'visibility': 'public',
         'role': 'spec', 'searchable': False, 'required': False}])
    dynamic_catalog.upsert_approved_products(conn, 'cat_ks', [{
        'id': 'k1', 'inner_code': 'K1', 'cs_visible': 1,
        'data': {'model': 'KS-1100',
                 'f_desc': '·电池:Li-ion 600mAh\n·充电时间:2小时\n·使用时间:90min',
                 'f_color': '黑色', 'f_size': '65.91*60.94*164.97',
                 'f_ctn': 'QTY：40 PCS\nN.W.：16.9 KGS'}, 'images': []}])
    conn.commit()


def test_attribute_question_gets_only_that_attribute(env):
    """问“什么颜色”只答颜色，不倒全表。"""
    conn, bot, api, llm = env
    _add_ks1100(conn)
    bot.handle_update(_upd('中文'))
    bot.handle_update(_upd('ks-1100 是什么颜色的'))
    reply = api.sent[-1][1]
    assert 'KS-1100 的颜色：黑色' in reply
    assert '功能描述' not in reply and '箱规' not in reply and '尺寸' not in reply


def test_carton_question_answers_full_carton(env):
    conn, bot, api, llm = env
    _add_ks1100(conn)
    bot.handle_update(_upd('中文'))
    bot.handle_update(_upd('KS-1100 一箱装多少'))
    reply = api.sent[-1][1]
    assert '的箱规：QTY：40 PCS' in reply
    assert 'N.W.：16.9 KGS' in reply          # 直接问箱规给完整值
    assert '功能描述' not in reply


def test_generic_intro_truncates_long_values(env):
    """整体介绍时长文只给首行摘要。"""
    conn, bot, api, llm = env
    _add_ks1100(conn)
    bot.handle_update(_upd('中文'))
    bot.handle_update(_upd('介绍一下 KS-1100'))
    reply = api.sent[-1][1]
    assert '电池:Li-ion 600mAh' in reply
    assert '充电时间' not in reply            # 多行只保留首行
    assert 'N.W.' not in reply


# ---------- 中途换语言：自然说法 ----------

def test_midconversation_switch_by_natural_phrase(env):
    """“我想转英文”直接切到 English，不再落对话大脑。"""
    conn, bot, api, llm = env
    bot.handle_update(_upd('中文'))
    bot.handle_update(_upd('我想转英文'))
    lang = conn.execute("SELECT lang FROM cs_customer WHERE tg_id='300'").fetchone()[0]
    assert lang == 'English'
    assert api.sent[-1][1].startswith('EN[')


def test_first_contact_natural_phrase_switches_directly(env):
    conn, bot, api, llm = env
    bot.handle_update(_upd('我想转英文'))
    lang = conn.execute("SELECT lang FROM cs_customer WHERE tg_id='300'").fetchone()[0]
    assert lang == 'English'
    assert 'Please choose your language' not in api.sent[-1][1]


def test_switch_back_to_chinese_by_bare_name(env):
    conn, bot, api, llm = env
    bot.handle_update(_upd('English'))
    bot.handle_update(_upd('中文'))
    lang = conn.execute("SELECT lang FROM cs_customer WHERE tg_id='300'").fetchone()[0]
    assert lang == '中文'
    assert '好的，已切换为 中文' in api.sent[-1][1]


def test_question_about_english_is_not_a_switch(env):
    """“怎么用英文写型号”是业务问题，不是换语言。"""
    from catalog import cs_i18n
    assert cs_i18n.parse_language_request('怎么用英文写这个型号') is None
    assert cs_i18n.parse_language_request('你会说英文吗') is None
    assert cs_i18n.parse_language_request('这个用英文怎么说') is None
    assert cs_i18n.parse_language_request('我想转英文') == ('switch', 'English')
    assert cs_i18n.parse_language_request('换回中文') == ('switch', '中文')
    assert cs_i18n.parse_language_request('转回 English') == ('switch', 'English')
    assert cs_i18n.parse_language_request('切回中文') == ('switch', '中文')
    assert cs_i18n.parse_language_request('切换成 Español') == ('switch', 'Español')
    assert cs_i18n.parse_language_request('换语言')[0] == 'prompt'
