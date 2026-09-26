"""新G 平台中央用户工具：游客拍照→清单→导出→验证码登录→游客记录合并。

独立 app（catalog/userapp.build_app）：独立 SQLite、不连档口商品库；
抽取链 FakeLlm 替身（不打网络），验证码从桩日志（codes_log）读取。
"""
import io
import json
import re

import openpyxl
import pytest
from fastapi.testclient import TestClient

from catalog import userapp


class FakeLlm:
    """与 tests/test_csbot 同款替身：可编程视觉应答。"""
    def __init__(self, vision_reply=None):
        self.vision_reply = vision_reply or json.dumps([{
            '型号或品名': '直发夹板', '价格': '80R', '装箱数': '40 PCS',
            '颜色': '黑色', '体积或尺寸': '未拍到', '其他': '未拍到'}],
            ensure_ascii=False)
        self.calls = []

    def chat_text(self, system, messages, **kw):
        return '<<PASS>>'

    def chat_vision(self, prompt, image_bytes, **kw):
        self.calls.append(prompt)
        return self.vision_reply


@pytest.fixture()
def client(tmp_path):
    llm = FakeLlm()
    app = userapp.build_app(db_path=str(tmp_path / 'u.db'),
                            photo_dir=str(tmp_path / 'ph'),
                            codes_log=str(tmp_path / 'codes.log'),
                            llm=llm)
    app.state.fake_llm = llm
    return TestClient(app)


def _new_guest(client):
    r = client.post('/guest')
    assert r.status_code == 200
    guest = r.json()['guest']
    assert re.fullmatch(r'guest-[0-9a-f]+', guest)
    return guest


def _last_code(client, email):
    """从桩日志拿该邮箱最新明文码（真渠道接入前的测试通道）。"""
    for line in reversed(open(client.app.state.codes_log, encoding='utf-8').read().splitlines()):
        mail, _, code = line.partition('\t')
        if mail == email:
            return code
    raise AssertionError('桩日志里没有该邮箱的验证码')


def _upload(client, guest, data=b'fake-jpeg-bytes', token=''):
    headers = {'Authorization': f'Bearer {token}'} if token else {}
    return client.post('/photo', data={'owner': guest}, files={'file': ('a.jpg', data, 'image/jpeg')},
                       headers=headers)


# ---------- 游客拍照 → 清单 ----------

def test_guest_photo_creates_notes_and_receipt(client):
    guest = _new_guest(client)
    r = _upload(client, guest)
    assert r.status_code == 200
    body = r.json()
    assert body['added'] == 1
    assert '【1】' in body['reply'] and '直发夹板' in body['reply']
    assert '已加入清单' in body['reply']
    rows = client.app.state.conn.execute(
        "SELECT * FROM notes WHERE owner_kind='guest' AND owner_id=?", (guest,)).fetchall()
    assert len(rows) == 1
    fields = json.loads(rows[0]['fields_json'])
    assert fields['型号或品名'] == '直发夹板'
    assert '待确认' in fields['价格']                    # 照片价仅记录口径
    assert rows[0]['photo_path']                          # 照片已落盘


def test_photo_without_identity_rejected(client):
    r = client.post('/photo', files={'file': ('a.jpg', b'x', 'image/jpeg')})
    assert r.status_code == 401


def test_business_card_routed_out_of_notes(client):
    """名片分流：名片不建清单条目，只进回执。"""
    client.app.state.llm.vision_reply = json.dumps([
        {'型号或品名': '直发夹板', '价格': '80R'},
        {'名片': {'档口名称': 'A区5号', '供应商联系人': '王先生'}}], ensure_ascii=False)
    guest = _new_guest(client)
    r = _upload(client, guest)
    assert r.json()['added'] == 1
    assert '名片' in r.json()['reply'] and 'A区5号' in r.json()['reply']
    assert client.app.state.conn.execute('SELECT COUNT(*) FROM notes').fetchone()[0] == 1


def test_unreadable_photo_no_note(client):
    client.app.state.llm.vision_reply = json.dumps([{'名片': {}}], ensure_ascii=False)
    guest = _new_guest(client)
    r = _upload(client, guest)
    assert r.json()['added'] == 0
    assert '没能认出' in r.json()['reply']
    assert client.app.state.conn.execute('SELECT COUNT(*) FROM notes').fetchone()[0] == 0


# ---------- 清单 / 导出 ----------

def test_notes_listing_by_guest(client):
    guest = _new_guest(client)
    _upload(client, guest)
    r = client.post('/notes', params={'guest': guest})
    assert r.status_code == 200
    notes = r.json()['notes']
    assert len(notes) == 1
    assert notes[0]['fields']['型号或品名'] == '直发夹板'
    assert notes[0]['photo'].startswith('/notes/')


def test_export_xlsx_structure(client):
    """列契约与 cs_export.render_notes 一致：序号+动态字段+商品照片，无确认状态列。"""
    guest = _new_guest(client)
    _upload(client, guest)
    r = client.get('/export.xlsx', params={'guest': guest})
    assert r.status_code == 200
    assert 'spreadsheetml' in r.headers['content-type']
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    headers = [c.value for c in ws[1]]
    assert headers[0] == '序号' and '型号或品名' in headers and headers[-1] == '商品照片'
    assert '确认状态' not in headers
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=2, column=headers.index('型号或品名') + 1).value == '直发夹板'


def test_export_empty_409(client):
    guest = _new_guest(client)
    r = client.get('/export.xlsx', params={'guest': guest})
    assert r.status_code == 409


def test_owner_isolation_between_guests(client):
    g1, g2 = _new_guest(client), _new_guest(client)
    _upload(client, g1)
    assert client.post('/notes', params={'guest': g2}).json()['notes'] == []


# ---------- 验证码登录 + 游客合并 ----------

def test_auth_code_and_verify_merges_guest_notes(client):
    guest = _new_guest(client)
    _upload(client, guest)
    email = 'buyer@example.com'
    r = client.post('/auth/code', json={'email': email})
    assert r.status_code == 200 and r.json()['sent'] is True
    assert 'code' not in r.json()                        # 响应统一“已发送”，不回显
    code = _last_code(client, email)
    assert re.fullmatch(r'\d{6}', code)

    wrong = client.post('/auth/verify', json={'email': email, 'code': '000000'})
    assert wrong.status_code == 400

    r = client.post('/auth/verify', json={'email': email, 'code': code, 'guest': guest})
    assert r.status_code == 200
    token = r.json()['token']

    me = client.get('/me', headers={'Authorization': f'Bearer {token}'})
    assert me.status_code == 200 and me.json()['email'] == email

    conn = client.app.state.conn
    assert conn.execute("SELECT COUNT(*) FROM notes WHERE owner_kind='user' AND owner_id=?",
                        (email,)).fetchone()[0] == 1     # 迁到账号
    assert conn.execute("SELECT COUNT(*) FROM notes WHERE owner_kind='guest' AND owner_id=?",
                        (guest,)).fetchone()[0] == 0     # guest 行清空
    assert client.post('/notes', params={'guest': guest}).json()['notes'] == []
    assert conn.execute("SELECT used_at FROM auth_codes WHERE code_hash=?",
                        (userapp._hash(f'{email}:{code}'),)).fetchone()[0]

    again = client.post('/auth/verify', json={'email': email, 'code': code, 'guest': guest})
    assert again.status_code == 400                      # 验证码一次性

    r = _upload(client, guest='', token=token)           # 登录后拍照直接挂账号
    assert r.status_code == 200
    assert conn.execute("SELECT COUNT(*) FROM notes WHERE owner_kind='user' AND owner_id=?",
                        (email,)).fetchone()[0] == 2


def test_expired_code_rejected(client, tmp_path):
    email = 'old@example.com'
    r = client.post('/auth/code', json={'email': email})
    assert r.status_code == 200
    conn = client.app.state.conn
    conn.execute("UPDATE auth_codes SET expires_at=datetime('now','-1 minute') WHERE email=?",
                 (email,))
    conn.commit()
    code = _last_code(client, email)
    r = client.post('/auth/verify', json={'email': email, 'code': code})
    assert r.status_code == 400
    assert client.post('/auth/verify', json={'email': email, 'code': '123456'}) .status_code == 400


def test_bad_email_and_bad_token(client):
    assert client.post('/auth/code', json={'email': 'not-an-email'}).status_code == 400
    assert client.get('/me', headers={'Authorization': 'Bearer nope'}).status_code == 401
    assert client.get('/me').status_code == 401


# ---------- 页面 ----------

def test_tool_page_served(client):
    r = client.get('/tool/')
    assert r.status_code == 200
    assert '拍照清单' in r.text and 'guest' in r.text      # 原生前端 + 游客模式
