"""动态分类报价映射：自动推荐、显式配置、改名，以及 generate_v2 动态分类出单。"""
import base64
import sqlite3

import openpyxl
import pytest
from fastapi.testclient import TestClient

from catalog import db, dynamic_catalog, quote
from catalog.main import app
from catalog.storage import LocalStorage

PNG = base64.b64decode(
    'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==')

CTN = 'QTY：40 PCS\nN.W.：13.7 KGS\nG.W.：14.5 KGS\nMEAS：38.5*37.5*42.5 CM'

FIELDS = [
    {'key': 'model', 'label': '产品型号', 'type': 'text', 'visibility': 'public', 'role': 'model'},
    {'key': 'color', 'label': '颜色', 'type': 'text', 'visibility': 'public', 'role': 'spec'},
    {'key': 'ctn', 'label': '箱规', 'type': 'text', 'visibility': 'public', 'role': 'spec'},
    {'key': 'price', 'label': '报价（不含税不含运）', 'type': 'money', 'visibility': 'public', 'role': 'price'},
    {'key': 'desc', 'label': '功能描述', 'type': 'text', 'visibility': 'public', 'role': 'spec'},
]


@pytest.fixture(autouse=True)
def conn():
    value = sqlite3.connect(':memory:', check_same_thread=False)
    value.row_factory = sqlite3.Row
    db.init_db(value)
    yield value
    value.close()


def _make_category(conn, fields=FIELDS, key='cat_razor_live', name='剃须刀现货'):
    dynamic_catalog.approve_template(conn, {
        'key': key, 'name': name, 'source_sheet': 'Sheet1',
        'fields': fields}, expected_version=0)
    return dynamic_catalog.get_template(conn, key)


def _add_product(conn, storage, category, pid='p1', price='69', model='KS-0273'):
    rel = storage.save(category, pid, 'main.png', PNG)
    dynamic_catalog.upsert_approved_products(conn, category, [{
        'id': pid, 'inner_code': 'KS-TEST', 'cs_visible': 1,
        'data': {'model': model, 'price': price, 'ctn': CTN, 'color': '华为灰',
                 'desc': '锂电池 800mAh'},
        'images': [rel]}])
    conn.commit()
    return rel


def test_suggest_quote_map_binds_unique_price_column(conn):
    mapping = dynamic_catalog.suggest_quote_map(FIELDS)
    assert mapping['price_field'] == 'price'
    assert mapping['model_field'] == 'model'
    assert mapping['ctn_field'] == 'ctn'
    assert mapping['color_field'] == 'color'


def test_suggest_quote_map_refuses_ambiguous_price(conn):
    fields = FIELDS + [{'key': 'price2', 'label': '内销价', 'type': 'money',
                        'visibility': 'public', 'role': 'price'}]
    mapping = dynamic_catalog.suggest_quote_map(fields)
    assert 'price_field' not in mapping


def test_approve_template_autosuggests_and_marks_quotable(conn):
    template = _make_category(conn)
    assert dynamic_catalog.quotable(template) is True
    ambiguous = _make_category(conn, key='cat_multi', name='多价格',
                               fields=FIELDS + [{'key': 'price2', 'label': '内销价',
                                                  'type': 'money', 'visibility': 'public', 'role': 'price'}])
    assert dynamic_catalog.quotable(ambiguous) is False


def test_set_quote_map_accepts_labels_and_keys(conn):
    _make_category(conn, key='cat_multi', name='多价格',
                   fields=FIELDS + [{'key': 'price2', 'label': '内销价',
                                     'type': 'money', 'visibility': 'public', 'role': 'price'}])
    merged = dynamic_catalog.set_quote_map(conn, 'cat_multi', {
        'price_field': '报价（不含税不含运）', 'model_field': 'model'})
    assert merged['price_field'] == 'price'
    assert dynamic_catalog.quotable(dynamic_catalog.get_template(conn, 'cat_multi'))


def test_set_quote_map_rejects_unknown_field_and_preset(conn):
    _make_category(conn)
    with pytest.raises(ValueError, match='不在分类'):
        dynamic_catalog.set_quote_map(conn, 'cat_razor_live', {
            'price_field': '不存在的列', 'model_field': 'model'})
    with pytest.raises(ValueError, match='预置分类'):
        dynamic_catalog.set_quote_map(conn, 'razor', {
            'price_field': 'price', 'model_field': 'model_no'})


def test_rename_template(conn):
    _make_category(conn)
    renamed = dynamic_catalog.rename_template(conn, 'cat_razor_live', '剃须刀现货表')
    assert renamed['name'] == '剃须刀现货表'
    with pytest.raises(ValueError, match='预置分类'):
        dynamic_catalog.rename_template(conn, 'razor', '新名字')


def test_generate_v2_dynamic_category_full_row(tmp_path, conn):
    _make_category(conn)
    storage = LocalStorage(str(tmp_path))
    _add_product(conn, storage, 'cat_razor_live')
    out = str(tmp_path / 'q.xlsx')
    quote.generate_v2(conn, storage,
                      [{'category': 'cat_razor_live', 'product_id': 'p1', 'quantity': 60}],
                      0, out, deposit_pct=30)
    ws = openpyxl.load_workbook(out).active
    assert ws.cell(18, 1).value == 'KS-0273'          # A 型号来自映射列
    assert ws.cell(18, 3).value == '锂电池 800mAh'      # C 描述来自剩余 spec 列
    assert ws.cell(18, 4).value == '华为灰'              # D 颜色映射列
    assert ws.cell(18, 5).value == 69                   # E 报价映射列
    assert ws.cell(18, 6).value == 80                   # F 整箱 40×2
    assert ws.cell(18, 8).value == 40                   # H PCS/CTN
    assert ws.cell(18, 9).value == 2                    # I 箱数
    assert ws.cell(18, 12).value == '38.5*37.5*42.5'    # L MEAS
    assert ws.cell(18, 13).value == 29.0                # M 总毛重 14.5×2


def test_generate_v2_dynamic_without_mapping_refuses(tmp_path, conn):
    _make_category(conn, key='cat_multi', name='多价格',
                   fields=FIELDS + [{'key': 'price2', 'label': '内销价',
                                     'type': 'money', 'visibility': 'public', 'role': 'price'}])
    storage = LocalStorage(str(tmp_path))
    _add_product(conn, storage, 'cat_multi')
    with pytest.raises(ValueError, match='未配置报价字段映射'):
        quote.generate_v2(conn, storage,
                          [{'category': 'cat_multi', 'product_id': 'p1', 'quantity': 10}],
                          0, str(tmp_path / 'q.xlsx'))


def test_quote_map_and_rename_endpoints(tmp_path, conn):
    _make_category(conn, key='cat_multi', name='多价格',
                   fields=FIELDS + [{'key': 'price2', 'label': '内销价',
                                     'type': 'money', 'visibility': 'public', 'role': 'price'}])
    app.state.conn = conn
    app.state.token = 'T0KEN'
    app.state.storage = LocalStorage(str(tmp_path))
    app.state.callback = None
    client = TestClient(app)
    headers = {'X-Service-Token': 'T0KEN'}

    r = client.get('/categories/cat_multi/quote-map', headers=headers)
    assert r.status_code == 200
    body = r.json()
    assert body['quotable'] is False
    assert body['suggestion'].get('model_field') == 'model'
    assert any(f['label'] == '内销价' for f in body['fields'])

    r = client.put('/categories/cat_multi/quote-map', headers=headers,
                   json={'price_field': '报价（不含税不含运）', 'model_field': '产品型号',
                         'ctn_field': '箱规', 'color_field': '颜色'})
    assert r.status_code == 200 and r.json()['quotable'] is True

    r = client.patch('/categories/cat_multi', headers=headers, json={'name': '吹风机现货'})
    assert r.status_code == 200 and r.json()['name'] == '吹风机现货'

    assert client.put('/categories/nope/quote-map', headers=headers,
                      json={'price_field': 'x', 'model_field': 'y'}).status_code == 404
    assert client.patch('/categories/razor', headers=headers,
                        json={'name': 'x'}).status_code == 400


def test_category_visibility_bulk_toggle(tmp_path, conn):
    from catalog.main import app
    from catalog.storage import LocalStorage
    _make_category(conn, key='cat_bulk', name='批量可见性')
    dynamic_catalog.upsert_approved_products(conn, 'cat_bulk', [
        {'id': f'p{i}', 'data': {'model': f'M{i}', 'price': '10'}, 'cs_visible': 1}
        for i in range(3)])
    conn.commit()
    app.state.conn = conn
    app.state.token = 'T0KEN'
    app.state.storage = LocalStorage(str(tmp_path))
    app.state.callback = None
    client = TestClient(app)
    headers = {'X-Service-Token': 'T0KEN'}

    r = client.patch('/categories/cat_bulk/visibility', headers=headers,
                     json={'visible': False})
    assert r.status_code == 200 and r.json()['updated'] == 3
    rows = dynamic_catalog.list_products(conn, 'cat_bulk')
    assert all(row['cs_visible'] == 0 for row in rows)
    assert dynamic_catalog.list_products(conn, 'cat_bulk', public_only=True) == []

    r = client.patch('/categories/cat_bulk/visibility', headers=headers,
                     json={'visible': True})
    assert r.status_code == 200
    assert len(dynamic_catalog.list_products(conn, 'cat_bulk', public_only=True)) == 3

    assert client.patch('/categories/nope/visibility', headers=headers,
                        json={'visible': True}).status_code == 404
